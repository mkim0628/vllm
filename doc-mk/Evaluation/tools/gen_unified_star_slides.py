#!/usr/bin/env python3
"""Update the all-DP overview deck with the unified star criteria (v1).

    python doc-mk/Evaluation/tools/gen_unified_star_slides.py IN.pptx OUT.pptx

Edits (text only, cell formats cloned from the deck itself):
  * DP2 slide: fills the Tradeoff stars/values (from DP2 results 2026-10-06, same rules as DP1/DP4).
  * DP3 slide: marks the Tradeoff cells "미평가 (TBD)" (no evaluation data exists).
  * DP1/DP4 slides: QA4 line shows (토큰 비용, M/M, 모듈 수) with unified basis (mid tier USD, man-month); stars unchanged.
  * inserts the unified-criteria appendix section after the "Appendix" divider; banners on superseded per-DP criteria slides.
Stars and sums come from tools/unified_stars.py (no hand-typed stars).
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import unified_stars as U  # noqa: E402

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

FONT = "맑은 고딕"
NAVY, GREY, RED = "2F5597", "404040", "C00000"
FILL = {3: "E2EFDA", 2: "FFF2CC", 1: "FBE5D6"}
rgb = RGBColor.from_string


def S(n):
    return "★" * n + "☆" * (3 - n)


# ------------------------------------------------------------------ cell helpers on existing slides
def _runs_fill(p, star, text, color=None):
    """Rewrite paragraph p to: star run + ' ' + text run (clone first run's formatting)."""
    runs = p.findall(qn("a:r"))
    for r in runs[1:]:
        p.remove(r)
    r0 = runs[0]
    r0.find(qn("a:t")).text = star
    rsp = copy.deepcopy(r0)
    rsp.find(qn("a:t")).text = " "
    rt = copy.deepcopy(r0)
    rt.find(qn("a:t")).text = text
    for r in (rsp, rt):
        sf = r.find(qn("a:rPr")).find(qn("a:solidFill"))
        if sf is not None and sf.find(qn("a:srgbClr")) is not None:
            sf.find(qn("a:srgbClr")).set("val", "1A1A1A")
    if color:
        for r in (r0, rt):
            sf = r.find(qn("a:rPr")).find(qn("a:solidFill"))
            sf.find(qn("a:srgbClr")).set("val", color)
    p.insert(list(p).index(r0) + 1, rsp)
    p.insert(list(p).index(rsp) + 1, rt)


def fill_value_cell(cell, template_p, lines):
    """lines: [(star_text, value_text, color or None)]"""
    txb = cell._tc.txBody
    for p in txb.findall(qn("a:p")):
        txb.remove(p)
    for star, text, color in lines:
        p = copy.deepcopy(template_p)
        _runs_fill(p, star, text, color)
        txb.append(p)


def tradeoff_table(slide):
    for sh in slide.shapes:
        if sh.has_table:
            return sh.table
    raise KeyError


def set_run_text(para, idx, text):
    para.runs[idx].text = text


# ------------------------------------------------------------------ new slides (style cloned from the deck's own appendix slides)
class Builder:
    def __init__(self, prs, layout):
        self.prs, self.layout = prs, layout

    def tb(self, s, x, y, w, h, paras, size=9, anchor=MSO_ANCHOR.TOP, name=None):
        """paras: list of (text, dict(bold,color,size)) or str"""
        b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            b.name = name
        tf = b.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(73152)
        tf.margin_top = tf.margin_bottom = Emu(36576)
        tf.vertical_anchor = anchor
        for i, pr in enumerate(paras):
            text, o = (pr, {}) if isinstance(pr, str) else pr
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(2.5)
            r = p.add_run()
            r.text = text
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.name = FONT
            r.font.color.rgb = rgb(o.get("color", "000000"))
            ea = etree.SubElement(r._r.get_or_add_rPr(), qn("a:ea"))
            ea.set("typeface", FONT)
        return b

    def frame(self, title, subtitle, footnote=None, src=None):
        s = self.prs.slides.add_slide(self.layout)
        for sh in list(s.placeholders):
            sh._element.getparent().remove(sh._element)
        self.tb(s, 0.4, 0.22, 12.5, 0.5, [(title, dict(bold=True, color=NAVY, size=22))], name="title", anchor=MSO_ANCHOR.MIDDLE)
        self.tb(s, 0.4, 0.74, 12.5, 0.35, [(subtitle, dict(color=GREY, size=11))], name="subtitle")
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(1.1), Inches(12.53), Inches(0.03))
        r.name = "rule"
        r.fill.solid(); r.fill.fore_color.rgb = rgb(NAVY); r.line.fill.background(); r.shadow.inherit = False
        if footnote:
            self.tb(s, 0.4, 7.0, 12.5, 0.3, [(footnote, dict(color=GREY, size=8.5))], name="footnote")
        if src:
            self.tb(s, 0.4, 7.2, 12.5, 0.25, [(src, dict(color=GREY, size=7.5))], name="src")
        return s

    def table(self, s, x, y, widths, rows, header=True, size=9.5, row_h=0.3, fills=None, bold_first_col=False, align=None, head_h=0.34):
        n, m = len(rows), len(widths)
        gf = s.shapes.add_table(n, m, Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * n))
        gf.name = "tbl"
        t = gf.table
        tblPr = t._tbl.tblPr
        for k in ("firstRow", "bandRow"):
            tblPr.set(k, "0")
        for j, w in enumerate(widths):
            t.columns[j].width = Inches(w)
        for i in range(n):
            t.rows[i].height = Inches(head_h if (header and i == 0) else row_h)
            for j in range(m):
                cell = t.cell(i, j)
                txt = rows[i][j]
                cell.margin_left = cell.margin_right = Emu(54864)
                cell.margin_top = cell.margin_bottom = Emu(22860)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                lines = txt if isinstance(txt, list) else [txt]
                for k, ln in enumerate(lines):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    p.alignment = PP_ALIGN.CENTER if (header and i == 0) or (align and align[j] == "c") else PP_ALIGN.LEFT
                    r = p.add_run()
                    r.text = ln
                    r.font.size = Pt(size)
                    r.font.name = FONT
                    r.font.bold = bool(header and i == 0) or (bold_first_col and j == 0)
                    r.font.color.rgb = rgb("FFFFFF" if header and i == 0 else "000000")
                    ea = etree.SubElement(r._r.get_or_add_rPr(), qn("a:ea"))
                    ea.set("typeface", FONT)
                color = NAVY if (header and i == 0) else (fills[i][j] if fills and fills.get(i) and fills[i][j] else "FFFFFF")
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(color)
                tcPr = cell._tc.get_or_add_tcPr()
                for q, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                    ln = etree.Element(qn(tag), w="6350", cap="flat", cmpd="sng", algn="ctr")
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr")).set("val", "BFBFBF")
                    etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
                    tcPr.insert(q, ln)
        return t

    def stars_table(self, s, rows, widths=(1.0, 2.6, 6.6, 2.33), y=1.4, size=10):
        data = [["별", "성능 기준 (정량)", "왜 이 값인가 (정량 근거)", "출처 · 확인 수준"]] + rows
        fills = {i + 1: [FILL[3 - i]] * 4 for i in range(3)}
        return self.table(s, 0.4, y, list(widths), data, size=size, row_h=0.72, fills=fills, align=["c", "l", "l", "l"])

    def blocks(self, s, left, right, y, h, lt="지표와 이전 기준", rt="정직한 주의", lw=6.12, size=10):
        self.tb(s, 0.4, y, lw, h, [(lt, dict(bold=True, color=NAVY, size=size + 1))] + [("• " + t, dict(size=size)) for t in left], name="l")
        self.tb(s, 6.78, y, 6.12, h, [(rt, dict(bold=True, color=RED, size=size + 1))] + [("• " + t, dict(size=size)) for t in right], name="r")


def build_slides(prs, res, inp):
    layout = [l for l in prs.slide_layouts if l.name == "빈 화면"][0]
    B = Builder(prs, layout)
    new = []

    # ---- overview
    s = B.frame("공통 별 기준 v1 — 모든 DP에 같은 QA는 같은 기준", "DP1·DP4의 기존 기준을 비교해 QA마다 하나로 통일 · 기존 별과 후보별 별 합계는 바뀌지 않음 · proposal (소유자 확정 전)",
                "※ 근거 문서: doc-mk/Evaluation/qa-star-criteria-unified.md · 별 계산: tools/unified_stars.py (값과 출처는 tools/unified_inputs.json)")
    rows = [["QA", "지표 (같은 DP의 Baseline 대비 또는 절대 SLO)", "★★★", "★★", "★", "적용 DP"],
            ["QA1 처리량", "Max SLO goodput ÷ Baseline (쌍별 기하평균)", "≥ 1.30", "0.97 ~ 1.30", "< 0.97", "DP1 · DP2 · DP4"],
            ["QA2 지연", "TTFT P99, TPOT P99 (두 지표 중 낮은 등급)", "≤ 2 s 그리고 ≤ 50 ms", "≤ 4 s 그리고 ≤ 100 ms", "그 외", "DP1~DP4"],
            ["QA3 자원", "HBM 사용량 감소 배수 = Baseline ÷ 후보 (시간 평균 GiB)", "≥ 1.25", "0.95 ~ 1.25", "< 0.95", "DP1 · DP3"],
            ["QA4 변경 용이성", "토큰 비용(mid tier $) 50% · 공수(MM) 30% · module 수 20%", "비용 ≤ $0.5, 공수 ≤ 0.5, module ≤ 2", "비용 ≤ $2, 공수 ≤ 1.0, module < 6", "그 위", "DP1 · DP2 · DP4"],
            ["QA5 기능 정확성", "F1 상대 하락 (압축 전 + full recompute 대비)", "≤ 1%", "≤ 3% (제안)", "> 3%", "DP3"],
            ["QA6 확장성", "SE(N_max=16) = Goodput(N) ÷ ((N/N0) × Goodput(N0)), N0=2", "≥ 90%", "70 ~ 90%", "< 70%", "DP4"]]
    fills = {i: ["EAF0FA", None, "E2EFDA", "FFF2CC", "FBE5D6", None] for i in range(1, 7)}
    B.table(s, 0.4, 1.35, [1.5, 4.5, 2.0, 2.0, 1.1, 1.43], rows, size=9, row_h=0.42, fills=fills, bold_first_col=True)
    B.blocks(s, ["같은 QA는 DP와 무관하게 같은 metric·집계·경계를 쓴다. DP마다 다른 것은 후보와 시나리오(공통 + DP 전용)뿐이다.",
                 "경계마다 근거를 달았다: 측정(잡음 폭), 환산(자원·시간), 표준·문헌, 정책 선택을 구분해 적었다.",
                 "별 합계가 높은 후보를 선택하고 동점일 때만 QA 우선순위를 쓴다(모든 DP 공통).",
                 "QA 번호는 이 PPT 1장 기준(QA5 정확성, QA6 확장성)."],
             ["이 기준은 기존 DP1·DP4 결과를 보존한다는 제약 아래 정했다. 모든 경계에 독립 근거를 달았지만 선택 자체는 결과를 본 뒤의 정의다(defined_after_first_look).",
              "QA5의 두 번째 경계(3%)는 데이터로 도출하지 못한 제안값이다.",
              "시뮬레이션 [B+C]와 Baseline 실측 + 시뮬레이션 [A+C]를 같은 표에 쓴다. 근거 수준은 별과 분리해 표기한다."],
             y=4.5, h=2.4, lt="원칙", rt="정직한 주의", size=10)
    new.append(s)

    def crit(title, subtitle, rows, left, right, foot, src, lt="지표와 이전 기준"):
        s = B.frame(title, subtitle, foot, src)
        B.stars_table(s, rows)
        B.blocks(s, left, right, 4.2, 2.8, lt=lt, size=10.5)
        new.append(s)

    crit("QA1 처리량 — 공통 별 기준", "Max SLO goodput(output tok/s) ÷ 같은 DP Baseline · 시나리오 쌍별 비율의 기하평균 · SLO = TTFT P99 ≤ 2 s, TPOT P99 ≤ 50 ms",
         [["★★★", "≥ 1.30 × Baseline", "같은 수요를 77% 이하의 자원으로 처리(자원 23% 절감, 8-GPU 노드 기준 약 1.85 GPU). 추가 메모리 HW와 migration·라우팅 복잡도를 정당화할 크기라는 판단. 보고된 모든 값에서 별이 같은 구간은 1.298 < 상한 ≤ 1.423이고 1.30은 그 안에 있다", "환산 [C], 정책 선택. 결과를 본 뒤 정함"],
          ["★★", "0.97 ~ 1.30 × Baseline", "하한 0.97 = Baseline끼리 seed 묶음을 바꿔 비교한 20회 집계 범위 0.968~1.033의 아래 끝. 그 안은 잡음이라 열세로 보지 않는다(보고된 최저 값은 1.04)", "측정 [B+C], DP1 star_basis.py"],
          ["★", "< 0.97 × Baseline", "잡음 대역 밖의 열세: Baseline보다 느리다", "측정 [B+C]"]],
         ["이전: DP1·DP2 0.97/1.30, DP4 0.90/1.10(공통 룰 §4.3). 통합: 0.97/1.30.",
          "공통 룰 0.90/1.10을 쓰면 DP1 C1 QA1이 ★★ → ★★★로 바뀌어 DP1 합계가 10 → 11이 된다. 0.97/1.30에서는 DP1·DP2·DP4 모두 그대로다.",
          "결과: DP1 C1 ×1.298(★★) · C2 ×1.423(★★★), DP2 ×1.66(둘 다 ★★★), DP4 ×1.05 · ×1.04(둘 다 ★★)."],
         ["★★★ 경계 1.30은 DP1 결과를 본 뒤 정했다. C1이 경계에서 0.2% 아래이고 1.25로 하면 C1도 ★★★이다.",
          "DP4는 Baseline이 실측이지만 재측정 잡음 폭을 아직 재지 않았다. 0.97은 시뮬레이터의 잡음 폭이므로 실측 후 같은 방법으로 확인해야 한다.",
          "DP2의 ×1.66은 Baseline이 SLO를 거의 못 지키는 시나리오가 끌어올린다(DP2 결과 문서 §0.2)."],
         "※ 공통 별점(0.90/1.10)은 필요하면 병기한다. 상한을 바꿔도 되는 구간과 다른 기준 계열의 영향은 qa-star-criteria-unified.md §7.",
         "출처: DP1/qa-criteria-dp1.md §A·§J · DP1/sim/star_basis.py · 공통 QA 문서 §4.3 · tools/unified_stars.py")

    crit("QA2 지연 — 공통 별 기준", "TTFT P99와 TPOT P99(시나리오 쌍별 값의 기하평균) · 두 지표 중 낮은 등급 적용 · TTFT와 TPOT는 별도 행으로 보고",
         [["★★★", "TTFT P99 ≤ 2 s 그리고 TPOT P99 ≤ 50 ms", "공통 기본 SLO. TTFT 2 s는 8K prefill 이론 약 0.32 s(H100×8, MFU 0.45)의 약 6배(큐잉 여유 포함)이고 MLPerf Llama-2-70B 표준 시나리오와 같다. TPOT 50 ms = 20 tok/s는 성인 묵독(238 wpm ≈ 5~6 tok/s, Brysbaert 2019)의 약 3.5배이고 MLPerf interactive 40 ms와 표준 200 ms 사이", "표준·문헌 [B], MLPerf 값은 2차 자료"],
          ["★★", "TTFT P99 ≤ 4 s 그리고 TPOT P99 ≤ 100 ms", "기본 SLO의 2배 이내 = 체감은 느리지만 대화형 사용은 가능한 한계", "정책 환산 [C]"],
          ["★", "TTFT P99 > 4 s 또는 TPOT P99 > 100 ms", "대화형 사용 한계 초과", "정책 환산 [C]"]],
         ["이전: DP4 = 위 기준(공통 룰 §5). DP1·DP2 = 6개 지표(TTFT·TPOT × P50/P95/P99)의 개선 배수 geomean 1.25/0.95. 통합: 절대 SLO.",
          "개선 배수를 DP4에 쓰면 후보 2가 ★★★ → ★★로 바뀌어 합계 9 → 8, 선택이 뒤집힌다. 절대 SLO에서는 DP1·DP2·DP4 별이 모두 그대로다.",
          "개선 배수는 별이 아니라 진단으로 병기하고, Baseline보다 나빠진 쌍 수와 최악 쌍을 함께 보고한다."],
         ["Baseline이 이미 SLO를 지키는 DP1에서는 두 후보가 모두 ★★★이라 이 QA가 후보를 가르지 못한다(개선 배수 기준도 같았다). 가르는 것은 진단 값이다: DP1 C1은 TTFT P99가 Baseline보다 나쁘다(×1.04, 21쌍 중 6쌍, 최악 ×4.32).",
          "기하평균 P99는 쌍별 꼬리를 평균해 한 쌍이 2 s를 넘는 경우를 가린다. 쌍별 위반 수를 병기해 보완한다.",
          "DP4 후보 간 별 차이는 TTFT 2 s 경계를 사이에 둔 약 6%(2,060 ms 대 1,940 ms)에서 갈린다. ±3~4%에서 뒤집힌다."],
         "※ 출처 링크는 PPT DP4 28장·34장과 같다. MLPerf 표준 값은 2차 자료 확인이 필요하다.",
         "출처: 공통 QA 문서 §5 · MLPerf Inference v5.1 요약(marktechpost, it-online) [2차] · Brysbaert 2019 · DP4 평가 시스템 장 · tools/unified_stars.py")

    crit("QA3 자원 (HBM 사용량) — 공통 별 기준", "HBM 사용량(GiB, 시간 평균 점유)의 Baseline 대비 감소 배수 = Baseline ÷ 후보 · 성능은 섞지 않는다 · DP1·DP3 적용",
         [["★★★", "감소 배수 ≥ 1.25", "HBM 점유 20% 감소 = DP1 평균 146.6 GiB 중 약 29 GiB(8K KV object 약 1.6개분). QA2와 같은 '눈에 띄는 개선' 크기로 둔 정책 선택", "환산 [C]"],
          ["★★", "0.95 ~ 1.25", "하한 0.95: Baseline끼리 비교한 점유 잡음 범위 0.984~1.016보다 느슨하게 둠(5% 이내 증가는 열세로 보지 않음)", "측정 [B+C]"],
          ["★", "< 0.95", "HBM을 5% 넘게 더 쓴다(잡음 밖)", "측정 [B+C]"]],
         ["이전: DP1이 이 기준(v5). 공통 룰 §6은 활용률 65%/85%인데 이동이 풀 총량을 바꾸지 않아 DP1에서 처리량과 상관이 0.99라 별에 쓰지 않았다.",
          "통합: DP1 기준을 공통으로 승격한다. DP3의 압축·재사용도 HBM 점유 감소가 목적이라 그대로 쓴다. DP2·DP4는 선정 QA가 아니며 진단 값으로 같은 조건에서 잰다.",
          "결과: DP1 C1 142.7 GiB(감소 1.03, ★★) · C2 178.0 GiB(감소 0.82, ★)."],
         ["DP1은 QA3 정의를 결과를 본 뒤 여러 번 바꿨다(활용률 → 풀 U → 성능÷비용 → HBM 사용량). 마지막 변경이 선택을 C2에서 C1로 바꿨다.",
          "'비우되 성능이 나빠지는 정책'이 유리하므로 QA1·QA2와 함께 읽어야 한다. 비용 가중 점유(ASSUMED 가격)는 보조로 병기한다.",
          "공통으로 승격하면 DP3부터는 사전 등록이 된다."],
         "※ 공통 룰의 활용률(65%/85%)은 필요하면 병기한다.",
         "출처: DP1/qa-criteria-dp1.md §A·§I·§J · DP1/sim/cost_model.py · 공통 QA 문서 §6 · tools/unified_stars.py")

    # ---- QA4 (two tables)
    s = B.frame("QA4 변경 용이성 — 공통 별 기준", "세 지표를 모두 보고 비중 토큰 비용 50% > 공수 30% > module 수 20%의 가중 하한 중앙값으로 별 결정 · 변경 시나리오 단순 평균에 경계 적용",
                "※ 공수·비용은 구현 측정(module·LOC)에 가정 상수를 곱한 값이며 실제 에이전트 세션 측정이 아니다 [B+C]. 1 MM = 21 man-day, 비용은 mid tier($3/$15 per MTok).",
                "출처: DP1/qa4-preregistration.md · DP4 평가 장(토큰 비용 경계) · 공통 QA 문서 §7 · tools/unified_stars.py")
    t1 = [["별", "① 토큰 비용 ($/변경, mid tier)", "② 개발 공수 (MM/변경)", "③ 변경 module 수"],
          ["★★★", "≤ $0.5", "≤ 0.5", "≤ 2 (주요 interface 유지)"],
          ["★★", "≤ $2", "≤ 1.0", "3 ~ 5 (2 초과 6 미만)"],
          ["★", "> $2", "> 1.0", "≥ 6 또는 major interface 변경"]]
    B.table(s, 0.4, 1.3, [1.0, 4.0, 3.5, 4.03], t1, size=9.5, row_h=0.3, fills={1: [FILL[3]] * 4, 2: [FILL[2]] * 4, 3: [FILL[1]] * 4}, align=["c", "l", "l", "l"])
    rows = [["DP · 후보", "① 비용 (mid)", "② 공수 (MM)", "③ module", "sub-star ①/②/③", "QA4 별"]]
    names = {"DP1": {"C1": "DP1 C1", "C2": "DP1 C2"}, "DP2": {"C1": "DP2 C1", "C2": "DP2 C2"}, "DP4": {"C1": "DP4 후보1", "C2": "DP4 후보2"}}
    for dp in ("DP1", "DP2", "DP4"):
        for c in ("C1", "C2"):
            v = inp[dp]["values"][c]
            q = v.get("qa4") or v["qa4_equal"]
            a, b_, m = U.qa4_subs(q["usd_t2"], q["mm"], q["modules"])
            rows.append([names[dp][c], f"${q['usd_t2']:.2f}", f"{q['mm']:.2f}", f"{q['modules']:.2f}", f"{S(a)} / {S(b_)} / {S(m)}", S(res[dp][c]["QA4"])])
    B.table(s, 0.4, 2.65, [2.0, 1.6, 1.6, 1.6, 3.2, 1.2], rows, size=9, row_h=0.27, align=["l", "c", "c", "c", "c", "c"])
    B.blocks(s, ["별 s는 'sub-star가 s 이상인 지표의 비중 합이 50%를 넘을 때' 얻는다. 가장 큰 비중(토큰 비용 50%)만으로는 별이 정해지지 않고 다른 지표 하나가 받쳐줘야 한다.",
                 "경계 근거: 공수 0.5 MM = 엔지니어 1명 2주 sprint 안에 흡수(1 MM = 한 달 전담이면 로드맵 항목). 비용 $0.5 = 약 25k 토큰 세션 1회 + 재작업 없음, $2는 그 4배(다중 세션·외부 코드 탐색). module 수는 공통 룰 §7.",
                 "이전: DP1·DP2는 frontier $3/$10과 sub-star 중앙값, DP4는 mid $0.5/$2와 man-day 3/7, 비용 우선, 가중 시나리오. 통합 후에도 DP1·DP2·DP4의 별은 같다."],
             ["DP1 C2: 공수 0.506 MM이 경계 0.5를 1.2% 넘고 비용 $0.495가 경계 1% 아래다. 공수가 0.5 아래면 별이 ★★ → ★★★로 바뀐다.",
              "DP4의 man-day 경계(3/7)를 쓰고 단순 평균을 쓰면 후보 2가 ★★★ → ★★로 바뀐다. DP4 시나리오 가중(S2~S4 ×2)은 별에 영향이 없다.",
              "가중 50/30/20과 '50% 초과' 규칙은 소유자 결정이다. 시나리오 집합은 새 메모리 tier 추가를 공통으로 두고 나머지는 DP별로 둔다."],
             y=4.75, h=2.2, lt="별 결정 규칙과 경계 근거", size=9.5)
    new.append(s)

    crit("QA5 기능 정확성 — 공통 별 기준 (DP3)", "F1 상대 하락률 = (압축 전 + full recompute의 F1 − 후보 F1) ÷ 압축 전 + full recompute의 F1 · 같은 task·dataset·모델·precision",
         [["★★★", "하락 ≤ 1%", "소유자가 정한 허용 정확도 저하 한도(압축 전 + full recompute 대비 F1 1% 이내, 상대 비율, 2026-10-09 확정)", "소유자 결정"],
          ["★★", "1% < 하락 ≤ 3%", "제안값: 허용 한도의 3배. '눈에 띄는 열화'가 시작되는 구간으로 본다. task의 run-to-run F1 변동을 재서 하한을 확인해야 한다", "정책 선택, 소유자 결정 필요"],
          ["★", "하락 > 3%", "허용 한도를 크게 넘은 정확도 손실", "정책 선택"]],
         ["DP3 전용 QA다. 다른 DP는 출력이 변하지 않는 것이 전제이므로 bit-exact 확인으로 다룬다(요구사항 QA-06 출력 정확성 시나리오).",
          "Baseline은 '압축 전 + full recompute'다. 후보의 정확도가 Baseline보다 높아지는 경우(F1 상승)는 하락 0%로 본다.",
          "DP3는 평가 데이터가 없어 별을 매기지 않았다(다음 장)."],
         ["두 번째 경계 3%는 데이터로 도출하지 못했다. 문헌은 KV 압축을 '무시할 만한 저하'로 보고하지만(KVzip: 3~4배 KV 축소, 디코딩 지연 약 2배 감소, 세부 F1은 확인 못 함) F1 하락의 일반적 임계는 없다.",
          "query-agnostic(1안)이 query-aware(2안)보다 정확도가 낮다는 설계 가정은 이 기준으로 측정해 확인해야 한다. 지금은 가정이다.",
          "F1을 정의할 task와 dataset이 미정이다."],
         "※ 한도 1%는 PPT DP3 설계 장의 Functional Correctness(F1 score) 지표와 같은 값이다.",
         "출처: 소유자 확정(2026-10-09) · KVzip, NeurIPS 2025 (neurips.cc/virtual/2025/oral/118742) [초록 확인]")

    crit("QA6 확장성 — 공통 별 기준 (DP4)", "SE(N) = Goodput(N) ÷ ((N / N0) × Goodput(N0)) · N0 = 2 노드 · 별은 SE(N_max=16) · N > 2는 모델 외삽(시뮬레이션)",
         [["★★★", "SE(N_max) ≥ 90%", "서버를 N배로 늘릴 때 처리량 손실 10% 이내(16노드에서 14.4노드분 효과). 병렬 효율 논문에서 '효율적'이라 보고할 때 흔히 쓰는 90%", "Cornell CVW · arXiv 1704.03329 [검색 확인]. 90%는 보고 관행이지 공식 표준이 아님"],
          ["★★", "SE(N_max) 70 ~ 90%", "확장은 되지만 조율 계층이 병목 조짐. 70% 하한 = 10대 증설에 7대분 효과(3대분 낭비). 매우 큰 규모에서 70~80%도 '우수'로 보고한 사례", "arXiv 2409.16053 AthenaK (65,536 GPU에서 80%) [검색 확인]"],
          ["★", "SE(N_max) < 70%", "손실 30% 초과 = 증설 비용의 1/3 이상 낭비. 조율 계층이 병목이다", "공통 QA 문서 §13 임시 정의(DP4 초안 값) [내부]"]],
         ["이전: DP4가 이 기준. DP2는 결과 문서에서 η(N=32) 0.90/0.70을 제안값으로 두었다(경계 같음, N_max만 다름. DP2는 PPT 선정 QA가 아님). N_max를 16으로 통일한다.",
          "결과: DP4 후보1 SE 93%(★★★), 후보2 64%(★). Baseline(무상태 LB)은 97%.",
          "레플리카가 상태를 공유하지 않는 후보2는 서버가 N배면 레플리카당 갱신 부하가 약 N배가 되어 결정 지연·상태 낡음이 커진다."],
         ["후보2의 ★은 보완(샤딩 구조) 전 설계 기준이다. 후보1의 Router도 전체 이벤트를 처리하는 구조라 비슷한 한계가 있을 수 있다.",
          "N > 2는 실측이 아니라 모델 외삽이다(C-04). 90%는 HPC 관행이며 공식 기준은 찾지 못했다.",
          "공통 룰 문서에는 §13이 없다(PPT가 인용한 DP4 초안의 절 번호). 공통 문서에 QA6를 정식 추가해야 한다."],
         "※ DP4 PPT 31장의 SE 곡선과 같은 값이다.",
         "출처: cvw.cac.cornell.edu/parallel/efficiency/scaling · arxiv.org/abs/1704.03329 · arxiv.org/abs/2409.16053")

    # ---- consistency check
    s = B.frame("적용 결과 — 정합성 점검 (공통 기준 v1)", "기존 별과 비교: DP1·DP2·DP4의 모든 별과 후보별 별 합계가 같다 · DP3는 평가 자료가 없어 미평가 · 별은 tools/unified_stars.py가 만든다",
                "※ DP2는 PPT 선정 QA(QA1, QA2, QA4)만 합산했다. 결과 문서의 QA3(GPU 사용률)·QA5(확장성)는 PPT 선정 QA가 아니어서 제외했다. DP4는 현 DP4(요청 조율, 옛 DP0)이며 근거는 PPT 26~31장이다.")
    names = {"DP1": ("DP1 C1 자원 상태 기반", "DP1 C2 AI Data 특성 기반"), "DP2": ("DP2 C1 스케줄링 시점", "DP2 C2 사전 계획"), "DP4": ("DP4 후보1 OSS 확장", "DP4 후보2 자체 구현")}
    prevt = {}
    rows = [["DP · 후보", "QA1", "QA2", "QA3", "QA4", "QA5", "QA6", "별 합계 (v1)", "기존 합계", "선택"]]
    sel = {"DP1": ("C1 (10 대 9)", ""), "DP2": ("구분 불가 (동점)", ""), "DP4": ("후보2 (동점, QA2에서 우세)", "")}
    for dp in ("DP1", "DP2", "DP4"):
        for k, c in enumerate(("C1", "C2")):
            st = res[dp][c]
            tot = sum(st.values())
            prev = sum(inp[dp]["previous_stars"][c].values())
            rows.append([names[dp][k]] + [S(st[q]) if q in st else "—" for q in ("QA1", "QA2", "QA3", "QA4", "QA5", "QA6")] + [str(tot), str(prev), sel[dp][0] if k == 0 else ""])
    for k, nm in enumerate(("DP3 1안 Offline", "DP3 2안 Online")):
        rows.append([nm, "—", "미평가", "미평가", "—", "미평가", "—", "—", "—", "보류" if k == 0 else ""])
    B.table(s, 0.4, 1.35, [2.5, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 2.03], rows, size=9, row_h=0.33, align=["l", "c", "c", "c", "c", "c", "c", "c", "c", "l"])
    B.blocks(s, ["DP1: 별이 후보 × QA 모두 같다(C1 10, C2 9, 선택 C1). DP4: 같다(9 대 9, 선택 후보2). DP2: 기존 결과(별 합계 9 대 9)를 같은 기준으로 다시 적용해 같다.",
                 "바뀐 것은 표시 단위다: DP4 QA4 $1.28/$0.38 → $1.02/$0.48, 3.9/2.6 man-day → 0.16/0.15 MM, module 3.1/1.9 → 2.7/2.3(가중 → 단순 평균). DP1·DP2 QA4 비용은 frontier $ → mid tier $(DP1 $1.16/$1.49 → $0.39/$0.49).",
                 "DP2의 두 후보는 QA1·QA2·QA4가 모두 같아 이 평가는 선택을 가르지 못한다(goodput 비 C2/C1 = ×1.002)."],
             ["다른 기준 계열을 쓰면 바뀌는 것: ① 공통 룰 QA1(0.90/1.10) → DP1 C1 합계 10 → 11, ② 개선 배수 QA2를 DP4에 적용 → 후보2 합계 9 → 8(선택 뒤집힘), ③ DP4의 공수 경계 3/7 man-day + 단순 평균 → 후보2 QA4 ★★★ → ★★. 이 세 가지가 v1을 정한 제약이다.",
              "경계 근접: DP1 C1 QA1 ×1.298(경계 1.30), DP1 C2 QA4 공수 0.506 MM(경계 0.5), DP4 QA2 TTFT 2,060 ms(경계 2,000 ms).",
              "DP4 QA1·QA2·QA6의 Baseline은 실측이지만 후보는 시뮬레이터·추정이다 [A+C]. DP1·DP2는 시뮬레이션 [B+C]다."],
             y=4.55, h=2.4, lt="읽는 법", rt="정직한 주의", size=9.5)
    new.append(s)

    # ---- DP2 result (unified format)
    s = B.frame("DP2 평가 결과 — 공통 기준 v1 (H100 + B200 통합)", "C1 스케줄링 시점 결정 vs C2 사전 계획 결정 · 괄호 = 후보 ÷ Baseline · 시뮬레이션 [B+C] · 근거: DP2 평가 결과 2026-10-06",
                "※ 별 경계는 공통 별 기준 v1이다(QA1 0.97/1.30, QA2 절대 SLO, QA4 가중 하한 중앙값). DP2 결과 문서의 QA3(GPU 사용률)·QA5(확장성)는 PPT 선정 QA가 아니라 제외했다.")
    v1, v2 = inp["DP2"]["values"]["C1"], inp["DP2"]["values"]["C2"]
    b = inp["DP2"]["baseline"]
    q1, q2 = res["DP2"]["C1"], res["DP2"]["C2"]
    qa4c = {c: inp["DP2"]["values"][c]["qa4"] for c in ("C1", "C2")}
    rows = [["QA", "평가 metric (↑/↓ 좋음)", "Baseline-PD-fixed", "C1 스케줄링 시점 결정", "C2 사전 계획 결정"],
            ["QA1 처리량", "Max SLO goodput (tok/s) ↑", f"{b['tps']:,}", f"{S(q1['QA1'])}  {v1['qa1_tps']:,} (×{v1['qa1_ratio']:.2f}) [B+C]", f"{S(q2['QA1'])}  {v2['qa1_tps']:,} (×{v2['qa1_ratio']:.2f}) [B+C]"],
            ["QA2 지연 (TTFT)", "TTFT P99 (ms) ↓", f"{b['ttft_p99_ms']:,}", f"{v1['ttft_p99_ms']:,} (×0.40)", f"{v2['ttft_p99_ms']:,} (×0.41)"],
            ["QA2 지연 (TPOT)", "TPOT P99 (ms) ↓", f"{b['tpot_p99_ms']}", f"{v1['tpot_p99_ms']} (×1.34)", f"{v2['tpot_p99_ms']} (×1.35)"],
            ["QA2 별점", "TTFT·TPOT P99 절대 SLO, 낮은 등급", "TTFT ★★ (>2 s)", f"{S(q1['QA2'])} [B+C]", f"{S(q2['QA2'])} [B+C]"],
            ["QA4 ① 비용", "에이전트 비용 $ (mid tier) ↓", "—", f"${qa4c['C1']['usd_t2']:.2f}", f"${qa4c['C2']['usd_t2']:.2f}"],
            ["QA4 ② 공수", "개발 공수 (man-month) ↓", "—", f"{qa4c['C1']['mm']:.2f}", f"{qa4c['C2']['mm']:.2f}"],
            ["QA4 ③ module", "변경 module 수 (4 시나리오 평균) ↓", "—", f"{qa4c['C1']['modules']:.2f}", f"{qa4c['C2']['modules']:.2f}"],
            ["QA4 별점", "가중 하한 중앙값 (50/30/20) [B+C]", "—", S(q1["QA4"]), S(q2["QA4"])],
            ["별 합계", "QA1 + QA2 + QA4", "—", f"{sum(q1.values())}", f"{sum(q2.values())}"]]
    B.table(s, 0.4, 1.35, [2.2, 3.4, 1.8, 2.6, 2.53], rows, size=9.5, row_h=0.36, align=["l", "l", "c", "c", "c"], bold_first_col=True)
    B.blocks(s, ["시스템: SYS-H100(H100×8, HBM3, PCIe 5.0, DDR5-4800) + SYS-B200(B200×8, HBM3e, PCIe 5.0, DDR5-6400) 통합. 노드 = 8-GPU, 6종 메모리, Llama-3.1-70B BF16, 노드 간 링크 RDMA 50 GB/s(ASSUMED). 시나리오 19개 × 2시스템 = 38쌍, 비교 가능 38쌍.",
                 "값은 쌍별 값의 기하평균, 괄호는 후보 ÷ Baseline. QA2는 Baseline의 최적 부하(iso-load)에서 비교했다(결과를 본 뒤 정한 정의).",
                 "선택: 두 후보의 별 합계가 9 대 9로 같고 QA1·QA2·QA4도 모두 같아 이 평가는 선택을 가르지 못한다(goodput 비 C2/C1 = ×1.002)."],
             ["QA1 ×1.66은 Baseline이 SLO를 거의 못 지키도록 설계된 시나리오(링크 경합 ×6~8, 긴 History ×3~4)가 끌어올린다. 중앙 시나리오는 ×1.0~1.7이다.",
              "TPOT P99가 Baseline보다 나쁘다(×1.34, 단 12.0 ms로 SLO 50 ms 이내). 공통 시나리오 3개에서는 TTFT P99도 ×1.50으로 나쁘다.",
              "후보 간 차이는 결정 시점이 아니라 결정 비용 가정(1 ms @ 64 후보)에 의존한다. 노드 수 5~6에서는 시점·지연이 성능을 바꾸지 못했다."],
             y=5.05, h=1.95, lt="평가 환경과 선택", rt="정직한 주의", size=9.5)
    new.append(s)

    # ---- DP3 status
    s = B.frame("DP3 평가 상태 — 별을 채울 자료가 없다", "공통 별 기준 v1의 QA2·QA3·QA5를 DP3에 적용하되, 측정 자료가 없어 별을 매기지 않았다",
                "※ 문헌 수치(KVzip 3~4배 KV 축소, CacheBlend TTFT 2.2~3.3배 감소 등)는 다른 시스템·조건의 값이라 DP3 후보의 별 근거로 쓰지 않았다(평가 규칙: 수치를 지어내지 않는다).")
    rows = [["QA", "지표 (공통 기준 v1)", "필요한 측정", "상태"],
            ["QA2 지연", "TTFT P99, TPOT P99 (절대 SLO)", "TTFT 분해: Comp.KV 로드(I/O) + 선택 재계산 + (2안) Query×KV attention. 재계산·attention 실행 자원은 DP2 결정", "미평가"],
            ["QA3 자원", "HBM 사용량 감소 배수", "Context KV가 차지하는 HBM 시간 평균(GiB). Comp.KV의 저장 tier는 DP1 결정", "미평가"],
            ["QA5 기능 정확성", "F1 상대 하락률 (압축 전 + full recompute 대비)", "task·dataset 확정 후 같은 모델·precision에서 측정. 1안(대표 Query)과 2안(실제 Query) 비교", "미평가 (task·dataset 미정)"]]
    B.table(s, 0.4, 1.35, [1.8, 3.2, 5.6, 1.93], rows, size=9.5, row_h=0.62, align=["l", "l", "l", "l"], bold_first_col=True)
    B.blocks(s, ["Evaluation/DP3는 TBD(README와 benchmark 틀뿐)이고 DP3 설계 문서와 PPT에도 정량 평가가 없다.",
                 "Baseline: As-Is(압축 없이 전체 Context KV 재사용). 정확도 기준값은 압축 전 + full recompute.",
                 "시나리오: 공통 벤치마크(CB-1~3) + DP3 전용(RAG·Agent의 재사용 Context 길이 32K~128K+, 재사용 횟수, 대표 Query와 실제 Query의 거리).",
                 "다음 단계(결정 필요): DP1 simulator를 확장한 시뮬레이션 [B+C]로 할지, 소규모 실측으로 할지."],
             ["DP3의 1안(Offline)은 TTFT 절감이 높고 정확도가 낮고, 2안(Online)은 반대라는 설계 가정은 측정 전까지 가정이다.",
              "QA2가 절대 SLO라서 두 후보가 모두 ★★★일 수 있다. 그러면 DP3의 trade-off는 QA3·QA5와 TTFT 진단 값에서 읽어야 한다.",
              "QA3는 설계 목적상 두 후보 모두 ★★★일 가능성이 높아 후보를 가르지 못할 수 있다(압축 비율이 설계 예산이라 결과가 아니라 입력이다)."],
             y=3.95, h=3.0, lt="현재 상태와 평가 계획", rt="미리 알아 둘 점")
    new.append(s)

    # ---- decisions
    s = B.frame("공통 별 기준 v1 — 소유자 결정이 필요한 사항", "결정 전까지는 proposal이며 기존 별은 모두 그대로다",
                "※ 상세는 doc-mk/Evaluation/qa-star-criteria-unified.md §8.")
    rows = [["#", "결정", "제안", "영향"],
            ["D1", "QA1 ★★★ 상한", "1.30 유지 (보존 구간 1.298 < 상한 ≤ 1.423)", "구간 밖이면 DP1 별 합계가 바뀐다"],
            ["D2", "QA2를 절대 SLO로 통일", "통일, 개선 배수는 진단", "DP1·DP2 별은 같고 설명이 단순해진다"],
            ["D3", "QA4 가중 50/30/20과 '50% 초과' 규칙", "제안대로", "DP1 C2 QA4가 경계 양쪽에 걸려 있다"],
            ["D4", "QA5 두 번째 경계 3%", "제안 3%, task의 F1 변동을 재서 확인", "DP3 별"],
            ["D5", "DP2 결과 세트", "2026-10-06 C1/C2(결정 시점)를 사용. 2026-10-10 Dispatcher 대 Blackboard(노드 내 attention 위치)는 후보 축이 달라 제외", "PPT의 최종 DP2가 어느 쪽인지 확인"],
            ["D6", "DP3 평가 방법", "시뮬레이션 [B+C] (DP1 simulator 확장) 또는 소규모 실측", "DP3 별"],
            ["D7", "QA 번호", "요구사항 덱의 QA-05 확장성 / QA-06 정확성을 PPT(QA5 정확성, QA6 확장성)에 맞춰 교환", "문서 일관성"],
            ["D8", "Evaluation/DP4 폴더", "옛 DP4(현 DP6, CXL 공유 메모리) 자료다. 현 DP4(요청 조율)의 평가 문서 위치를 정한다", "평가 문서 정리"]]
    B.table(s, 0.4, 1.35, [0.6, 3.2, 5.6, 3.13], rows, size=9.5, row_h=0.56, align=["c", "l", "l", "l"], bold_first_col=True)
    new.append(s)
    return new


def move_after(prs, slides, after_index):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    n_orig = len(ids) - len(slides)
    news = ids[n_orig:]
    for el in news:
        lst.remove(el)
    for k, el in enumerate(news):
        lst.insert(after_index + 1 + k, el)


# ------------------------------------------------------------------ main
def main(src, dst):
    inp = U.load()
    res = U.evaluate(inp)
    prs = Presentation(src)
    tpl_slide = prs.slides[2]   # DP1 tradeoff cell with the star/value run structure
    tpl = tradeoff_table(tpl_slide).cell(4, 2)._tc.txBody.findall(qn("a:p"))[0]

    # --- DP2 slide (6): fill tradeoff
    t = tradeoff_table(prs.slides[5])
    for col, c in ((2, "C1"), (4, "C2")):
        v = inp["DP2"]["values"][c]
        st = res["DP2"][c]
        q = v["qa4"]
        fill_value_cell(t.cell(4, col), tpl, [
            (S(st["QA1"]), f"{v['qa1_tps']:,}(x{v['qa1_ratio']:.2f})", None),
            (S(st["QA2"]), f"{v['ttft_p99_ms']:,}(x{v['ttft_p99_ms']/inp['DP2']['baseline']['ttft_p99_ms']:.2f}), {v['tpot_p99_ms']}(x{v['tpot_p99_ms']/inp['DP2']['baseline']['tpot_p99_ms']:.2f})", "FF0000"),
            (S(st["QA4"]), f"${q['usd_t2']:.2f}, {q['mm']:.2f}, {q['modules']:.2f}", None)])
    for idx in (5, 7):
        cols = tradeoff_table(prs.slides[idx]).columns
        for c, w in zip(cols, (0.96, 3.51, 2.44, 3.71, 2.24)):
            c.width = Inches(w)
    # --- DP3 slide (8): not evaluated
    t = tradeoff_table(prs.slides[7])
    for col in (2, 4):
        fill_value_cell(t.cell(4, col), tpl, [("미평가", "(평가 자료 없음)", "7F7F7F")] * 3)

    # --- DP1 slides (3, 4): QA4 line + label order
    for idx in (2, 3):
        t = tradeoff_table(prs.slides[idx])
        for col, c in ((2, "C1"), (4, "C2")):
            q = inp["DP1"]["values"][c]["qa4"]
            p = t.cell(4, col).text_frame.paragraphs[3]
            runs = p.runs
            runs[-1].text = f"${q['usd_t2']:.2f}, {q['mm']:.2f}, {q['modules']:.2f}"
        for col in (1, 3):
            p = t.cell(4, col).text_frame.paragraphs[3]
            r = p.runs
            if len(r) == 4:   # 'Modifiability ', '(module, M/M, ', '토큰 비용', ')'
                r[1].text, r[2].text = "(토큰 비용, M/M, ", "모듈 수"
            else:             # 'Modifiability (module, M/M, ', '토큰 비용', ')'
                r[0].text, r[1].text = "Modifiability (토큰 비용, M/M, ", "모듈 수"
    # --- DP4 slide (10): QA4 line unified basis
    t = tradeoff_table(prs.slides[9])
    for col, c in ((2, "C1"), (4, "C2")):
        q = inp["DP4"]["values"][c]["qa4_equal"]
        p = t.cell(4, col).text_frame.paragraphs[2]
        txt = f"${q['usd_t2']:.2f}, {q['mm']:.2f}, {q['modules']:.1f}"
        rs = p.runs
        rs[-1].text = txt
        for r in rs[2:-1]:
            r.text = ""

    # --- banners on superseded per-DP criteria slides
    layout = [l for l in prs.slide_layouts if l.name == "빈 화면"][0]
    B = Builder(prs, layout)
    def banner(slide, text, y=1.13):
        B.tb(slide, 0.4, y, 12.5, 0.24, [(text, dict(bold=True, color=RED, size=8.5))], name="banner")
    banner(prs.slides[18], "※ 공통 별 기준 v1 QA1과 같은 기준이다(0.97/1.30). 기존 별 그대로.")
    banner(prs.slides[19], "※ 공통 별 기준 v1에서 대체됨: QA2는 절대 SLO(TTFT ≤ 2 s, TPOT ≤ 50 ms)로 통일. 이 장의 개선 배수는 진단으로 병기. DP1 별(C1 ★★★, C2 ★★★)은 그대로.")
    banner(prs.slides[20], "※ 공통 별 기준 v1 QA3과 같은 기준이다(감소 배수 1.25/0.95). 기존 별 그대로.")
    banner(prs.slides[21], "※ 공통 별 기준 v1에서 대체됨: 비용 mid tier $0.5/$2, 공수 0.5/1.0 MM, 별 = 가중 하한 중앙값(토큰 50 · 공수 30 · module 20). DP1 별(C1 ★★★, C2 ★★)은 그대로.")
    banner(prs.slides[26], "※ 공통 별 기준 v1에서 대체됨: QA1 경계는 0.97/1.30(이 장의 0.90/1.10 대신). 이 PPT의 Q1 = 공통 QA1. DP4 별(둘 다 ★★)은 그대로.", y=1.13)
    banner(prs.slides[27], "※ 공통 별 기준 v1 QA2와 같은 기준이다. 이 PPT의 Q2 = 공통 QA2. 기존 별 그대로.")
    banner(prs.slides[28], "※ 공통 별 기준 v1에서 대체됨: 비용·공수·module 우선순위 규칙 대신 가중 하한 중앙값, 공수는 MM(man-day ÷ 21), 시나리오 단순 평균. 이 PPT의 Q3 = 공통 QA4. DP4 별(후보1 ★★, 후보2 ★★★)은 그대로.")
    banner(prs.slides[30], "※ 공통 별 기준 v1 QA6과 같은 기준이다(SE 90/70, N_max=16). 이 PPT의 Q4 = 공통 QA6. 기존 별 그대로.")

    # --- unified section after the Appendix divider (slide 11)
    new = build_slides(prs, res, inp)
    move_after(prs, new, 10)
    prs.save(dst)
    print("saved", dst, len(prs.slides._sldIdLst), "slides; new:", len(new))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
