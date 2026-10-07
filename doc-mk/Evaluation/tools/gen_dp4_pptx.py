#!/usr/bin/env python3
"""Build the DP4 evaluation appendix deck (H15 a) and the complement-design tactics slide (H15 b) from the result data.

    uv run --no-project --with python-pptx --with pillow python doc-mk/Evaluation/tools/gen_dp4_pptx.py [--out-dir doc-mk/DP4] [--render-dir DIR]

Every number comes from tools/gen_dp4_result.py (the generator of the result document), so slides and document cannot diverge.
The palette (STY) and font path come from tools/gen_dp_pptx.py (not modified; its Slide pre-wraps text for a PowerPoint font, so DP4 uses its own Slide that lets the renderer wrap); the frame is slide 2 of doc-mk/DP4/DP4-slides-draft.pptx.
Outputs: DP4-appendix-qa-result.pptx (table + system + 2 'why' slides + scenarios) and DP4-complement-design-tactics.pptx (1 slide).
"""
from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import gen_dp4_result as g  # noqa: E402
import gen_dp_pptx as gp  # noqa: E402  (STY palette, FONT_PATH; the DP1 deck builder stays untouched)

B, C1, C2 = g.B, g.C1, g.C2
TITLE = "DP4. 비일관 CXL 공유 메모리의 KV 일관성 구조 - Appendix"
LATIN_FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"   # wide fallback for Latin glyphs when 맑은 고딕 is not installed (conservative)
FIT_WARNINGS: list[str] = []


def _fonts(size):
    from PIL import ImageFont
    px = max(8, int(size * 110 / 72))
    return ImageFont.truetype(LATIN_FONT, px), ImageFont.truetype(gp.FONT_PATH, px)


def n_lines(text, width_in, size):
    """Conservative number of rendered lines of one paragraph: Latin glyphs measured in DejaVu Sans, CJK in WenQuanYi, 6% safety."""
    from PIL import Image, ImageDraw
    d = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    fl, fc = _fonts(size)
    maxw = (width_in - 0.1) * 110 * 0.94
    lines, cur = 1, 0.0
    for ch in text:
        w = d.textlength(ch, font=fl if ord(ch) < 0x2E80 and ch not in "★·→←↑↓÷≈≤≥" else fc)
        if cur + w > maxw and cur > 0:
            lines += 1
            cur = 0.0
        cur += w
    return lines


class Slide:
    """Like gen_dp_pptx.Slide but WITHOUT pre-wrapping: the renderer wraps, and a conservative line estimate warns when a box is too small."""

    def __init__(self, title):
        self.title, self.E = title, []

    def box(self, x, y, w, h, text, sty, size=8.5, bold=False, align="l", top=True):
        paras = text if isinstance(text, list) else [text]
        need = sum(n_lines(p, w, size) for p in paras) * size * 1.22 / 72 + 0.08
        if need > h + 1e-6:
            FIT_WARNINGS.append(f"[{self.title[:28]}] box at y={y:.2f} needs {need:.2f} in, has {h:.2f}: {paras[0][:40]}")
        self.E.append(dict(x=x, y=y, w=w, h=h, lines=paras, sty=sty, size=size, bold=bold, align=align, top=top))


def qa_slides():
    c = g.final_cells()
    b = c[B]
    o = g.selection()
    m = g.Q4["mean_over_scenarios"]
    sub = g.Q4["sub_stars"]
    ft = g.fit_counts()
    ln = g.link_numbers()
    pf = g.prefill_ms()
    ql = g.qa4_boundary_lines()
    dt, kept, tied = g.decision_table()
    cr = g.cp_ranges()
    hr = g.headroom()
    be = g.SENS["break_even"]
    q, dpd = g.agg_dev_c1_c2()
    cc = g.RM["dp4_benchmark"]["cp_capacity"]
    se = g.RM["dp4_benchmark"]["scaling_efficiency"]
    f0, f1, f2, xr, pct = g.f0, g.f1, g.f2, g.xr, g.pct
    nodeloss = {k.split("@")[1]: g.goodput_ratio("dp4_benchmark", k, C1) for k in g.fail_rows() if "node_loss" in k}
    lh = {k.split("@")[1]: g.goodput_ratio("dp4_benchmark", k, C2) for k in g.fail_rows() if "as_published" in k}
    ag = "d4_agent_multiturn@H100"
    agent = {k.split("@")[1]: g.goodput_ratio("dp4_benchmark", k, C1) for k in g.DPPS if k.startswith("d4_agent_multiturn")}
    pcf = g.pc_facts()
    wp = g.win_pairs()

    # ---------------- slide 1: table + brief system + selection ----------------
    s = Slide("DP4 평가 결과 - QA별 정량 metric (H100 + B200 통합, 별점 = Common 6쌍)")
    cols = [("QA / 평가 metric", 0.4, 2.9), ("Baseline-RDMA", 3.3, 2.1), ("C1 중앙 직렬화", 5.4, 3.75), ("C2 분산 락", 9.15, 3.75)]
    y = 1.15
    for name, x, w in cols:
        s.box(x, y, w, 0.32, name, "head", 10, True, "ctr", False)
    h_prev = [0.34]

    def row(label, base, a, bb, h=0.40, size=9, sty="cell"):
        nonlocal y
        y += h_prev[0]
        s.box(0.4, y, 2.9, h, label, "dp", 8.5, True, "l", False)
        s.box(3.3, y, 2.1, h, base, "cell", size, False, "ctr", False)
        s.box(5.4, y, 3.75, h, a, sty, size, False, "ctr", False)
        s.box(9.15, y, 3.75, h, bb, sty, size, False, "ctr", False)
        h_prev[0] = h + 0.02

    r2 = lambda cd, k: cd[k] / b[k]
    row(["QA1 Throughput", "Max SLO goodput (tok/s) ↑"], f"{f0(b['qa1'])} (Baseline 별 {b['s1']})",
        f"{c[C1]['s1']}  {f0(c[C1]['qa1'])} ({xr(c[C1]['qa1_r'], 3)})", f"{c[C2]['s1']}  {f0(c[C2]['qa1'])} ({xr(c[C2]['qa1_r'], 3)})")
    row(["QA2 TTFT (각 arm의 peak load)", "P99 · P50 (ms) ↓"], f"P99 {f0(b['t99'])} · P50 {f0(b['t50'])}",
        f"P99 {f0(c[C1]['t99'])} ({xr(r2(c[C1], 't99'))}) · P50 {f0(c[C1]['t50'])} ({xr(r2(c[C1], 't50'))})", f"P99 {f0(c[C2]['t99'])} ({xr(r2(c[C2], 't99'))}) · P50 {f0(c[C2]['t50'])} ({xr(r2(c[C2], 't50'))})")
    row(["QA2 TTFT (공통 load = Baseline peak)", "P99 · P50 (ms) ↓"], f"P99 {f0(b['ct99'])} · P50 {f0(b['ct50'])}",
        f"P99 {f0(c[C1]['ct99'])} ({xr(r2(c[C1], 'ct99'))}) · P50 {f0(c[C1]['ct50'])} ({xr(r2(c[C1], 'ct50'))})", f"P99 {f0(c[C2]['ct99'])} ({xr(r2(c[C2], 'ct99'))}) · P50 {f0(c[C2]['ct50'])} ({xr(r2(c[C2], 'ct50'))})")
    row(["QA2 TPOT", "P99 · P50 (ms) ↓"], f"P99 {f2(b['o99'])} · P50 {f2(b['o50'])}",
        f"P99 {f2(c[C1]['o99'])} ({xr(r2(c[C1], 'o99'))}) · P50 {f2(c[C1]['o50'])} ({xr(r2(c[C1], 'o50'))})", f"P99 {f2(c[C2]['o99'])} ({xr(r2(c[C2], 'o99'))}) · P50 {f2(c[C2]['o50'])} ({xr(r2(c[C2], 'o50'))})")
    row(["QA2 별점", "6쌍 중 최악 P99"], f"{b['s2']} (TTFT {f0(b['tw'])} ms)", f"{c[C1]['s2']}  (TTFT {f0(c[C1]['tw'])} · TPOT {f1(c[C1]['ow'])} ms)", f"{c[C2]['s2']}  (TTFT {f0(c[C2]['tw'])} · TPOT {f1(c[C2]['ow'])} ms)")
    row(["QA3 Resource usage", "KV 상주 GiB (Baseline peak load) ↓"], f"{f1(b['res'])} (Baseline 별 {b['s3']})",
        f"{c[C1]['s3']}  {f1(c[C1]['res'])} ({xr(c[C1]['res'] / b['res'])}; 절감 {xr(c[C1]['save'])})", f"{c[C2]['s3']}  {f1(c[C2]['res'])} ({xr(c[C2]['res'] / b['res'])}; 절감 {xr(c[C2]['save'])})")
    row(["QA4 Modifiability", "module · 공수(MM) · 비용($) ↓"], "—",
        f"{g.Q4_STARS[C1]}  {m['C1']['modules']:.2f} · {m['C1']['man_months']:.2f} · ${m['C1']['usd_T1']:.2f}  ({sub['C1']['M1']}/{sub['C1']['M2']}/{sub['C1']['M3']})",
        f"{g.Q4_STARS[C2]}  {m['C2']['modules']:.2f} · {m['C2']['man_months']:.2f} · ${m['C2']['usd_T1']:.2f}  ({sub['C2']['M1']}/{sub['C2']['M2']}/{sub['C2']['M3']})")
    row(["별 합계", ""], "—", f"{o['totals'][C1]}", f"{o['totals'][C2]}", h=0.40, size=11)
    y += h_prev[0] + 0.05
    sf = {sid: g.sys_facts(sid) for sid in g.SYSIDS}
    sysline = " + ".join(f"{sid[4:]}x{sf[sid]['gpus']} ({sf[sid]['gen']['gpu_hbm'].split(' ')[0]}, {sf[sid]['gen']['host_link']}, {sf[sid]['gen']['cxl'].split(' (')[0]})" for sid in g.SYSIDS)
    s.box(0.4, y, 12.5, 0.9, [
        f"시스템: {sysline}. P/D 분리 클러스터(노드 = GPU 8장, CB 기본 P1+D1), Llama-3.1-70B BF16. CXL 풀 8 TiB, 노드 CXL {g.cfgv('cxl_adapters_per_node')} x {g.cfgv('cxl_adapter_bw_Bps') / 1e9:g} GB/s x η {g.cfgv('eta_cxl'):g}(ASSUMED), Baseline RDMA {g.cfgv('nic_bw_Bps') / 1e9:g} GB/s x η {g.cfgv('eta_rdma'):g}(ASSUMED). 메모리 구성은 다음 장.",
        f"집계: 별 = Common 6쌍(CB-1~3 x 2시스템) 기하평균, 전체 {len(g.CBLAB) + len(g.DPLAB)}쌍(비교 가능 {ft['common_benchmark']['comparison_valid'] + ft['dp4_benchmark']['comparison_valid']}, 포화 {ft['common_benchmark']['saturated'] + ft['dp4_benchmark']['saturated']}, 불가 {ft['common_benchmark']['infeasible'] + ft['dp4_benchmark']['infeasible']})는 diagnostic·QA4. 괄호 = 후보 ÷ Baseline, ↑ 높을수록 좋음 ↓ 낮을수록 좋음. 모두 [B+C] (CXL 풀 실측 [A] 없음)."], "note", 9)
    y += 0.97
    s.box(0.4, y, 12.5, 1.1, [
        f"선택 질문은 'CXL 풀을 쓸 것인가'가 아니라 C1(중앙 직렬화) 대 C2(분산 락) = control plane 책임이다. 선택: C1 — 별 합계 {o['totals'][C1]} 대 {o['totals'][C2]} (규칙: 합계가 높은 후보, `dp_selection.py`; QA 우선순위 {' > '.join(g.PRIO['priority'])}는 proposal).",
        f"QA1~QA3은 C1 = C2(차이 {pct(max(q.values()), 3)} 이하), 별 차이는 QA4에서만: 공수 {ql['c1']:.3f} 대 {ql['c2']:.3f} MM이 경계 0.5 MM의 양쪽(C1 {pct(ql['c1_margin'])} 아래, C2 {pct(ql['c2_margin'])} 위). QA4 변형 {len(dt)}가지 중 {tied}가지는 동점(미결정) -> 약한 선택.",
        f"Baseline 대비: QA1 {xr(c[C1]['qa1_r'], 3)}, QA3 절감 {xr(c[C1]['save'])}, 공통 load TTFT P99 {xr(c[C1]['ct99'] / b['ct99'])} (별 ★). 격차는 data plane 상수(η_cxl 대 η_rdma, ASSUMED 대 ASSUMED)에서 오며 C1·C2 공통. under the tested conditions the architecture shows no benefit over the baseline."], "sel", 9.5)

    # ---------------- slide 2: memory configuration ----------------
    s2 = Slide("DP4 평가 시스템 - 메모리 구성 (H100x8 / B200x8, 값이 다르면 H100 / B200)")
    cols2 = [("메모리 / 경로", 0.4, 2.3), ("용량", 2.7, 1.7), ("host 연결 (세대)", 4.4, 2.8), ("대역폭", 7.2, 2.9), ("연산 능력 · 지원 연산", 10.1, 2.8)]
    for name, x, w in cols2:
        s2.box(x, 1.15, w, 0.32, name, "head", 10, True, "ctr", False)
    h_, b_ = sf["SYS-H100"], sf["SYS-B200"]
    node_cxl = g.cfgv("cxl_adapters_per_node") * g.cfgv("cxl_adapter_bw_Bps") / 1e9
    rows2 = [
        ("HBM (GPU 8장 합)", f"{h_['hbm_gib']:,.0f} / {b_['hbm_gib']:,.0f} GiB ({h_['hbm_med']} / {b_['hbm_med']})", f"GPU on-package, host CPU와 {h_['gen']['host_link']}", f"{h_['hbm_bw']:.1f} / {b_['hbm_bw']:.1f} TB/s",
         f"GPU FP16 dense {h_['fp16']:,.0f} / {b_['fp16']:,.0f} TFLOPS/GPU · attention, FFN"),
        ("CXL 공유 풀 (C1, C2의 KV 저장소)", f"{g.cfgv('pool_capacity_bytes') / 2**40:g} TiB", f"{h_['gen']['cxl']} 스위치 풀, 노드당 어댑터 {g.cfgv('cxl_adapters_per_node')}개 x PCIe 5.0 x16",
         f"노드당 {node_cxl:g} GB/s(방향별) x η {g.cfgv('eta_cxl'):g} = {node_cxl * g.cfgv('eta_cxl'):g} GB/s, 풀 집계 {g.cfgv('pool_bw_Bps') / 1e12:g} TB/s", "없음 (순수 저장소)"),
        ("RDMA NIC (Baseline의 KV 경로)", "—", "노드당 4 x 200 Gbps", f"{g.cfgv('nic_bw_Bps') / 1e9:g} GB/s x η {g.cfgv('eta_rdma'):g} = {g.cfgv('nic_bw_Bps') * g.cfgv('eta_rdma') / 1e9:g} GB/s", "없음"),
        ("host DRAM", "—", h_["gen"]["host_link"], "—", "사용하지 않음 (모든 arm 0)"),
        ("ScHBM, CXL-PNM, HBF, SSD-PIM", "—", "—", "—", "DP4 평가 범위 밖 (profile에만 존재)"),
    ]
    yy = 1.5
    for r in rows2:
        for (name, x, w), v in zip(cols2, r):
            s2.box(x, yy, w, 0.7, v, "dp" if name == "메모리 / 경로" else "cell", 9, name == "메모리 / 경로", "l", False)
        yy += 0.72
    s2.box(0.4, yy + 0.08, 12.5, 1.9, [
        "시뮬레이션 반영: 노드 링크와 풀의 유효 대역(η 포함)과 background 부하, 접근 지연에서 오는 control plane 연산 비용(C1 서버 큐, C2 락 scan), 풀 점유율에 따른 퇴출 연산, decode·prefill 계산 시간(DP1 물리 재사용: 8K prefill H100 "
        f"{f0(pf['SYS-H100'])} ms, B200 {f0(pf['SYS-B200'])} ms), KV 상주량 집계.",
        "미반영: HBM·풀의 용량 한계(OOM), 스위치 혼잡과 장치 bank 경합(집계 BW 상한만), RDMA의 작은 블록 fragmentation(강한 baseline), 락 매니저·서버 스레드의 CPU 스케줄링 간섭, 풀 장치 장애, host DRAM tier, 전력.",
        "값의 출처: 풀 8 TiB, 1 TB/s, 어댑터 63 GB/s x 2, NIC 4 x 200 Gbps, CXL 지연, 서버 RPC 시간은 Beluga/TraCT 본문(PAPER, 서버 2대 측정). η_cxl, η_rdma, 중첩 비율, background 부하, C2 락 상수는 ASSUMED. GPU/HBM 값은 PUBLIC(확인 필요). 상세: system-specs.md, cluster_dp4.json"], "note", 9)

    # ---------------- why slides ----------------
    def why_slide(title, items, fs=9.5):
        sl = Slide(title)
        for name, x, w in (("QA / 수치", 0.4, 2.7), ("왜 이런 값이 나왔나", 3.1, 6.7), ("근거 (시나리오·측정값)", 9.8, 3.1)):
            sl.box(x, 1.15, w, 0.32, name, "head", 10, True, "ctr", False)
        yy = 1.5
        for h, a, bb, cc_ in items:
            sl.box(0.4, yy, 2.7, h, a, "dp", 10, True, "l", False)
            sl.box(3.1, yy, 6.7, h, bb, "cell", fs, False, "l", False)
            sl.box(9.8, yy, 3.1, h, cc_, "note", fs - 1, False, "l", False)
            yy += h + 0.05
        return sl

    kb = "cb_kv_8k_b32@B200"
    vb_ = g.CBPS[kb]
    w3 = g.CBPS["cb_mixed_8k_b32@H100"]
    s3 = why_slide("DP4 평가 결과 - QA별로 왜 이런 값이 나왔나 (1/2: 처리량, 지연)", [
        (2.15, ["QA1 처리량", f"C1 = C2 {xr(c[C1]['qa1_r'], 3)} ±{c[C1]['qa1_ci']:.3f}", f"H100 {xr(g.R['SYS-H100']['common_benchmark']['qa_feasible'][C1]['qa1_ratio_geomean'], 3)} · B200 {xr(g.R['SYS-B200']['common_benchmark']['qa_feasible'][C1]['qa1_ratio_geomean'], 3)}"],
         [f"- 두 후보 모두 같은 data plane을 쓴다. 8K 요청의 KV {ln['kv_gib']:.2f} GiB를 Baseline은 RDMA로 한 번(링크 시간 어림 {f0(ln['t_rdma_ms'])} ms), 후보는 풀에 쓴 뒤 publish·pin을 거쳐 다시 읽는다(쓰기 {f0(ln['t_cxl_ms'])} + 읽기 {f0(ln['t_cxl_ms'])} ms; config 상수로 계산한 어림값, bg {ln['bg']:g}).",
          f"- 노드 CXL 유효 대역 {ln['cxl']:g} GB/s(η {g.cfgv('eta_cxl'):g})가 RDMA {ln['nic']:g} GB/s(η {g.cfgv('eta_rdma'):g})보다 낮다. B200은 prefill이 {f0(pf['SYS-B200'])} ms(H100 {f0(pf['SYS-H100'])} ms)로 빨라 링크가 병목이라 격차가 크다.",
          f"- 별은 두 후보가 같다: 일관성 구조(C1 대 C2)는 값을 바꾸지 않는다(차이 {pct(q['qa1_abs_goodput_geomean_tps'], 3)})."],
         [f"{kb}: Baseline {f0(vb_[B]['max_goodput_tps'])} -> C1 {f0(vb_[C1]['max_goodput_tps'])} tok/s ({xr(vb_[C1]['max_goodput_tps'] / vb_[B]['max_goodput_tps'], 2)})", f"cb_mixed@H100: {xr(w3[C1]['max_goodput_tps'] / w3[B]['max_goodput_tps'], 2)}", f"break-even η_cxl* {be[g.MERGED][C1]['bg0.85']['qa1_ratio_parity']['x']:.3f} (대역폭 동등점 {g.SENS['meta']['eta_bw_parity']:.3f})"]),
        (2.1, ["QA2 TTFT", f"공통 load P99 {xr(c[C1]['ct99'] / b['ct99'])}", f"각 arm peak load P99 {xr(c[C1]['t99'] / b['t99'])}"],
         [f"- TTFT는 KV 이동이 노출되는 경로다. 같은 offered load(Baseline peak)에서 후보 P99 {f0(c[C1]['ct99'])} ms 대 Baseline {f0(b['ct99'])} ms: 후보가 그 load에서 이미 포화(공통 6쌍 모두 악화, 최악 {xr(g.CBQ[C1]['tail_check_common_load']['ttft']['worst_ratio'], 1)}).",
          f"- 각 arm의 peak load끼리 비교하면 후보가 낮아 보이지만({f0(c[C1]['t99'])} ms) 후보의 peak가 더 낮은 load에 있어서이며 이득이 아니다.",
          f"- 별은 둘 다 ★: 6쌍 중 최악 P99가 Baseline {f0(b['tw'])} ms, 후보 {f0(c[C1]['tw'])} ms로 4 s 초과."],
         [f"{kb} 공통 load: Baseline {f0(vb_[B]['at_base_peak']['ttft_p99_ms'])} -> C1 {f0(vb_[C1]['at_base_peak']['ttft_p99_ms'])} ms, 그 load goodput {f0(vb_[C1]['at_base_peak']['goodput_tps'])} tok/s", f"악화 쌍 {g.CBQ[C1]['tail_check_common_load']['ttft']['n_worse']}/6 (공통 load)"]),
        (1.15, ["QA2 TPOT", f"P99 {xr(c[C1]['o99'] / b['o99'])}"],
         [f"- 거의 같다({f2(b['o99'])} vs {f2(c[C1]['o99'])} ms). TPOT는 decode의 HBM 대역폭이 정하고 DP4는 decode 방식을 바꾸지 않으며 KV는 decode 시작 전에 HBM에 도착한다. SLO(50 ms)보다 훨씬 낮아 개선 여지도 없다."],
         [f"TPOT P99 악화 쌍 {g.CBQ[C1]['tail_check_common_load']['tpot']['n_worse']}/6"]),
    ], fs=10.5)
    pb = c[B]["comp"]["p_buffer"]
    s4 = why_slide("DP4 평가 결과 - QA별로 왜 이런 값이 나왔나 (2/2: 자원, 변경 용이성, C1 대 C2)", [
        (1.45, ["QA3 KV 상주", f"{xr(c[C1]['res'] / b['res'])} (절감 {xr(c[C1]['save'])})", f"own-peak {f1(c[C1]['res_own'])} vs {f1(b['res_own'])} GiB"],
         [f"- 사전 등록 정의는 Baseline peak load에서 같은 요청 구간을 비교한다. 후보는 그 load에서 이미 포화해 P 노드 전송 버퍼에 요청이 쌓인다(P buffer {f1(c[C1]['comp']['p_buffer'])} 대 {f1(pb)} GiB). 풀 사본은 평균 {f1(c[C1]['comp']['pool'])} GiB로 작다.",
          f"- 각 arm의 own-peak load에서는 {f1(c[C1]['res_own'])} 대 {f1(b['res_own'])} GiB(절감 {xr(b['res_own'] / c[C1]['res_own'])}, ★★ 구간). 정의 민감도일 뿐 C1 = C2라 선택은 불변."],
         [f"{kb}: KV 상주 {f1(vb_[B]['at_base_peak']['kv_resident_gib'])} -> {f1(vb_[C1]['at_base_peak']['kv_resident_gib'])} GiB", f"QA3 하한 0.95는 Baseline 잡음 안(H100 {pct(g.SB['SYS-H100']['qa3_noise_max_dev'])})"]),
        (1.5, ["QA4 변경 용이성", f"C1 {g.Q4_STARS[C1]} · C2 {g.Q4_STARS[C2]}", f"공수 {ql['c1']:.3f} vs {ql['c2']:.3f} MM"],
         [f"- 변경 시나리오 4종(신규 HW capability, 객체 class, 정책 교체, topology)을 시뮬레이터 proxy에 구현해 diff를 쟀다. module 평균은 같다({m['C1']['modules']:.2f}). 차이는 S4(풀 2개)에서 C2의 LOC {g.Q4['scenarios']['S4']['C2']['loc_added']} 대 {g.Q4['scenarios']['S4']['C1']['loc_added']}, S1에서 C2 module 크기가 더 큰 데서 온다.",
          f"- 이 차이가 별을 가르는 것은 경계 0.5 MM가 두 값 사이라서다. 상수·집계·경계 변형 {len(dt)}가지 중 {tied}가지에서 C1 = C2(동점)."],
         [f"S4 공수: C1 {g.Q4['scenarios']['S4']['C1']['man_months']:.3f} · C2 {g.Q4['scenarios']['S4']['C2']['man_months']:.3f} MM", "측정 기준 소스가 Iteration 2 이전(재측정 필요)", "[B+C] proxy, 실제 에이전트 세션 아님"]),
        (2.45, ["C1 대 C2 trade-off", "성능은 같고 장애·확장 특성이 다름", f"model check: 2노드 전수 통과, 3노드 불완전"],
         [f"- control plane 연산은 C1 {f1(min(x['c1'][0] for x in cr.values()))}~{f1(max(x['c1'][1] for x in cr.values()))} µs, C2 {f1(min(x['c2'][0] for x in cr.values()))}~{f0(max(x['c2'][1] for x in cr.values()))} µs로 C2가 최대 {f0(max(x['ratio'][1] for x in cr.values()))}배 느리지만 TTFT는 수백 ms 이상이라 QA에 안 드러난다.",
          f"- 확장: C2 포화 rate는 S x N에 따라 감소(stripe 1 -> 512: {f0(cc['d4_lock_stripes_1@H100'][C2]['saturation_rps'])} -> {f0(cc['d4_lock_stripes_512@H100'][C2]['saturation_rps'])} req/s), C1은 스레드로 증가({f0(cc['d4_lock_stripes_1@H100'][C1]['saturation_rps'])} -> {f0(cc['d4_server_threads_4@H100'][C1]['saturation_rps'])}). 제공 rate보다 최소 {f0(hr[0][0])}배 위라 N <= 16에서 포화 없음.",
          f"- 장애: C1 서버 노드 손실(30 s)에서 goodput {', '.join(f'{k} {xr(v, 2)}' for k, v in nodeloss.items())}(C2 영향 없음); C2 as-published 락 고착(요청 {g.RM['dp4_benchmark']['failure']['d4_fail_lockholder[as_published]@H100'][C2]['n_incomplete_stuck']:.0f}건 정지, goodput {', '.join(f'{k} {xr(v, 3)}' for k, v in lh.items())}), lease 보완안은 회복.",
          f"- 정확성 [C]: 정상 C1·C2는 2노드 전수, 결함 변종 {len(pcf['defects'])}개 모두 검출. 3노드는 상한 {pcf['cap']:,} 상태에서 모두 불완전(검출 {len(pcf['detected3'])}/{len(pcf['defects'])})."],
         [f"포화 rate와 장애: `d4_lock_stripes_*`, `d4_server_threads_*`, `d4_fail_*`", f"C2 락 상수는 ASSUMED, lease는 평가자의 가정", "stale slot-body 가정이 C1_no_clflush 결과를 좌우"]),
    ])

    # ---------------- scenarios ----------------
    ncb, ndp = len(g.CBLAB), len(g.DPLAB)
    sat = sorted(k for k, v in g.RM["dp4_benchmark"]["fit"].items() if v == "saturated")
    s5 = Slide("DP4 평가에서 고려한 시나리오")
    s5.box(0.4, 1.15, 12.5, 1.0, [
        f"prefill 서버가 만든 KV cache(대화 문맥의 중간 결과)를 decode 서버로 넘기는 경로를 비교했다. Baseline은 서버 사이를 RDMA로 복사하고, 후보는 서버들이 함께 쓰는 CXL 공유 메모리 풀에 한 번 두고 서로 읽는다. 공유 풀은 서버 간 캐시 일관성을 하드웨어가 주지 않아 '목록(메타데이터)을 누가 관리하나'가 문제이며 답이 C1(한 서버가 직렬 처리), C2(모든 노드가 락으로 직접 갱신)다.",
        f"공통 3개 시나리오 x 2시스템 = {ncb}쌍, DP4 전용 {g.n_dp4_scen()}개(변종 포함 {ndp // 2}행) x 2시스템 = {ndp}쌍. 비교 가능 {ft['common_benchmark']['comparison_valid'] + ft['dp4_benchmark']['comparison_valid']}쌍, 포화 {ft['common_benchmark']['saturated'] + ft['dp4_benchmark']['saturated']}쌍({', '.join(sat)}), Baseline도 SLO 불가인 쌍 {ft['common_benchmark']['infeasible'] + ft['dp4_benchmark']['infeasible']}쌍. 승/무/패(Baseline 대비) C1 {'/'.join(map(str, g.tally_counts(C1)))}, C2 {'/'.join(map(str, g.tally_counts(C2)))}; goodput이 95% CI 밖으로 좋아진 쌍 {g.n_goodput_wins()}."], "note", 9.5)
    blocks = [
        ("기본 서비스 상황 (공통 3개, 별점 산출)", f"8K 입력·256 생성의 대화 서비스에서 서버 사이 링크가 다른 트래픽으로 거의 차 있는 경우(background 0.85), 점점 막히는 경우(0.2에서 0.9), KV에 LoRA·MoE·Agent·Tool 데이터가 섞이는 경우. 결과: 6쌍 모두 후보가 Baseline보다 낮다(QA1 {xr(c[C1]['qa1_r'], 2)}). 대표 예: 링크가 {ln['bg'] * 100:.0f}% 차 있을 때 KV {ln['kv_gib']:.1f} GiB를 Baseline은 한 번({f0(ln['t_rdma_ms'])} ms), 후보는 쓰고 다시 읽어 두 번({f0(2 * ln['t_cxl_ms'])} ms)."),
        ("접근이 쏠린 경우", "요청의 90%가 같은 긴 prefix를 공유(`d4_hot_prefix_fanout`): C1 서버 대기열, C2 락 경합을 건드린다. 두 구조 모두 영향이 없었다. 고르게 흩어진 접근만 따로 본 시나리오는 아직 없다."),
        ("데이터 종류별", f"KV cache 중심. 8턴 에이전트의 History KV 재사용(`d4_agent_multiturn`, 별도 행): H100 {xr(agent['H100'], 2)}, B200 {xr(agent['B200'], 2)}(H100은 Baseline seed 일부 붕괴로 확정 불가). 혼합(CB-3)은 풀 트래픽·객체 수만 반영하고 종류별 접근 패턴은 모델링하지 않았다."),
        ("규모·자원 조건의 변화", "서버 4·8·16대, 블록 크기 16·256 토큰, 작은 prompt를 높은 rate로(512 in/64 out), 풀 95% 점유에서 매번 퇴출, C2 락 stripe 1·8·512, C1 서버 스레드 2·4. 모두 C1 = C2이고 control plane 포화 rate는 제공 rate보다 크게 위(최소 " + f"{f0(hr[0][0])}배)."),
        ("장애", "C1 메타데이터 서버 재시작(0.5 s)과 노드 손실(인덱스 재구성 30 s), C2 락 보유 노드 장애(as-published는 락이 풀리지 않음, lease 보완안은 풀림). C1 SPOF와 C2 락 고착이 별점 밖에서 둘을 가른다."),
        ("정확성 (별점 밖) / 아직 없는 것", f"비일관 캐시 추상 모델에서 두 프로토콜과 결함 변종 {len(pcf['defects'])}개를 검사. 아직 없음: 실제 CXL 풀 측정, 서버 2대 초과 실측, 다중 테넌트, 풀 장치 장애, 스위치 혼잡, KV 외 객체의 종류별 패턴, 균일 접근 전용 시나리오, 실제 trace."),
    ]
    yy = 2.25
    for t, body in blocks:
        s5.box(0.4, yy, 2.9, 0.75, t, "dp", 10, True, "l", False)
        s5.box(3.3, yy, 9.6, 0.75, body, "cell", 10, False, "l", False)
        yy += 0.79
    return [s, s2, s3, s4, s5]


def tactics_slide():
    """Complement design for the selected structure (C1): weakness (evaluation evidence) -> tactic, with verification status ([B] measured / [C] argued, unimplemented)."""
    o = g.selection()
    sel = "C1" if o["winner"] == C1 else ("C2" if o["winner"] == C2 else "C1")
    c = g.final_cells()
    cc = g.RM["dp4_benchmark"]["cp_capacity"]
    nodeloss = {k.split("@")[1]: g.goodput_ratio("dp4_benchmark", k, C1) for k in g.fail_rows() if "node_loss" in k}
    ql = g.qa4_boundary_lines()
    xr, f0, pct = g.xr, g.f0, g.pct
    agent = {k.split("@")[1]: g.goodput_ratio("dp4_benchmark", k, C1) for k in g.DPPS if k.startswith("d4_agent_multiturn")}
    ncl = g.pc_facts()["r2"]["C1_no_clflush_on_server"]["invariants"]["I1"]["cex_length"]
    be = g.SENS["break_even"][g.MERGED][C1]["bg0.85"]["qa1_ratio_parity"]["x"]
    hr = [h for h in g.headroom() if h[2] == C1][0]
    s = Slide(f"DP4 보완 설계 택틱 - 선택 구조 {sel}(중앙 직렬화) 기준")
    cols = [("#", 0.4, 0.45), ("약점 (평가 근거)", 0.85, 4.1), ("보완 택틱", 4.95, 4.3), ("개선 QA", 9.25, 1.15), ("검증 상태", 10.4, 2.5)]
    for name, x, w in cols:
        s.box(x, 1.2, w, 0.34, name, "head", 9, True, "ctr", False)
    rows = [
        ("W1", f"메타데이터 서버 SPOF: 노드 손실(재구성 30 s)에서 goodput {', '.join(f'{k} {xr(v, 2)}' for k, v in nodeloss.items())} (C2 영향 없음). 프로세스 재시작(0.5 s)은 영향 없음",
         "스탠바이 서버를 다른 노드에 두고 CXL의 인덱스를 그대로 인계(재구성 풀 스캔 회피)", "QA1·QA2 (장애 시)", "[B] 재시작 경로(CXL 인덱스 보존)는 구현·측정. [C] 스탠바이 인계는 미구현"),
        ("W2", f"서버 포화 rate 유한(1 스레드 {f0(cc['d4_lock_stripes_1@H100'][C1]['saturation_rps'])} req/s, block 16에서 {f0(cc['d4_block16@H100'][C1]['saturation_rps'])}). 제공 rate보다 최소 {f0(hr[0])}배 위, 코어 1개 상시 busy-poll",
         "서버 스레드 증설, 큰 block(256)으로 연산 수 감소", "QA1 (확장 시)", f"[B] 구현·측정: 스레드 2/4에서 {f0(cc['d4_server_threads_2@H100'][C1]['saturation_rps'])}/{f0(cc['d4_server_threads_4@H100'][C1]['saturation_rps'])} req/s"),
        ("W3", f"정확성이 서버 CLFLUSH-before-read와 슬롯 프로토콜에 의존: `C1_no_clflush_on_server`는 I1 위반(반례 {ncl} step). stale slot-body 저자 가정에 의존, 3노드는 탐색 상한으로 불완전",
         "슬롯 body에 sequence/검증값을 넣어 stale body를 서버가 감지(누락 flush를 safety가 아닌 liveness 문제로), conformance 시험 추가", "정확성 (별점 밖)", "[C] 논증·미구현, 효과 수치 없음"),
        ("W4", f"QA1 {xr(c[C1]['qa1_r'], 3)}, QA3 절감 {xr(c[C1]['save'])}, 공통 load TTFT P99 {xr(c[C1]['ct99'] / c[B]['ct99'])}: data plane(풀 경유 쓰기 + 읽기 순차)이 원인, C2도 동일. break-even η_cxl* {be:.3f} (ASSUMED 대 ASSUMED)",
         "(a) 필요한 블록만 읽는 partial read와 풀 사본 제거, (b) 청크 단위 publish로 쓰기·읽기 겹침(publish 경계 유지), (c) 단일 사용 KV는 RDMA 직접, 풀은 공유·재사용 객체에만(하이브리드)", "QA1·QA2·QA3", "[C] 논증·미구현. 새 iteration으로 사전 등록 후 측정해야 하며 효과 수치를 주장하지 않음"),
        ("W5", f"reuse 이득이 시스템 의존·미확정: `d4_agent_multiturn` {', '.join(f'{k} {xr(v, 2)}' for k, v in agent.items())}, H100은 Baseline seed 붕괴(CV {pct(g.DPPS['d4_agent_multiturn@H100'][B]['goodput_cv'], 0)}), 판정 tie",
         "reuse 객체(History KV, 공유 prefix)에만 풀 사용(W4-c와 같은 계열), seed 추가로 Baseline 붕괴 빈도 확인", "QA1·QA2", "[C] 미구현, seed 추가 미실시"),
        ("W6", f"QA4 별이 경계 근처: C1 {ql['c1']:.3f} MM이 0.5 MM보다 {pct(ql['c1_margin'])} 아래, C2 {ql['c2']:.3f} MM. 측정 기준 소스가 Iteration 2 이전",
         "Iteration 2 이후 소스에서 QA4 시나리오 4종 재측정, 가능하면 실제 vLLM 통합에서 측정", "QA4", "[C] 미실시. 재측정 전 QA4 별 확정 안 함"),
    ]
    y = 1.6
    for r_ in rows:
        h = 0.78
        for (name, x, w), v in zip(cols, r_):
            s.box(x, y, w, h, v, "dp" if name == "#" else ("c2" if name == "보완 택틱" else "cell"), 9.5, name == "#", "ctr" if name in ("#", "개선 QA") else "l", False)
        y += h + 0.04
    s.box(0.4, y + 0.02, 12.5, 0.56, [
        "[B] = 시뮬레이터에서 구현·측정됨([B+C]), [C] = 논증·미구현. 미구현 택틱의 효과는 수치로 주장하지 않는다(H14). 선택 C1은 QA4 한 칸(★★★ 대 ★★, 경계 근처)에 기댄 약한 선택이며, 성능 격차 대 Baseline은 택틱 W4가 다루는 data plane에서 온다."], "note", 8.5)
    return s


def build(slides, out, base_pptx):
    """Same construction as gen_dp_pptx.build, with the DP4 frame (slide 2 of DP4-slides-draft.pptx) and DP4 title."""
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    FONT = "맑은 고딕"
    rgb = RGBColor.from_string
    prs = Presentation(str(base_pptx))
    base = prs.slides[1]
    keep = {3, 5, 8, 39, 57}
    # the draft's layout carries a static '/ 23' text box (total of the draft deck); drop it from this deck's layout copy
    for sh in list(base.slide_layout.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("/ "):
            sh._element.getparent().remove(sh._element)
    originals = list(prs.slides._sldIdLst)
    for sl in slides:
        s = prs.slides.add_slide(base.slide_layout)
        for sh in list(s.shapes):
            sh._element.getparent().remove(sh._element)
        for sh in base.shapes:
            if sh.shape_id in keep:
                s.shapes._spTree.append(copy.deepcopy(sh._element))
        for sh in s.shapes:
            if sh.shape_id == 3:
                r = sh.text_frame.paragraphs[0].runs
                r[0].text = TITLE
                for x in r[1:]:
                    x.text = ""
            if sh.shape_id == 8:
                r = sh.text_frame.paragraphs[0].runs
                r[0].text = sl.title
                for x in r[1:]:
                    x.text = ""
                sh.width = Inches(12.5)
        for e in sl.E:
            f, l, t = gp.STY[e["sty"]]
            b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
            b.shadow.inherit = False
            b.fill.solid()
            b.fill.fore_color.rgb = rgb(f)
            b.line.color.rgb = rgb(l)
            b.line.width = Pt(0.75)
            tf = b.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.05)
            tf.margin_top = tf.margin_bottom = Inches(0.03)
            tf.vertical_anchor = MSO_ANCHOR.TOP if e["top"] else MSO_ANCHOR.MIDDLE
            for i, ln in enumerate(e["lines"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER if e["align"] == "ctr" else PP_ALIGN.LEFT
                r = p.add_run()
                r.text = ln
                r.font.size = Pt(e["size"])
                r.font.name = FONT
                r.font.bold = bool(e["bold"]) or (i == 0 and e["sty"] in ("note", "sel") and not ln.startswith("-") and not ln.startswith("평가"))
                r.font.color.rgb = rgb(t)
                ea = etree.SubElement(r._r.get_or_add_rPr(), "{%s}ea" % A)
                ea.set("typeface", FONT)
    lst = prs.slides._sldIdLst
    for el in originals:
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    prs.save(str(out))


def render(pptx_paths, out_dir):
    """LibreOffice -> PDF -> PNG (pdftoppm) for visual overlap / cut-off checks."""
    import subprocess
    out_dir.mkdir(parents=True, exist_ok=True)
    for pp in pptx_paths:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(pp)], check=True, capture_output=True, timeout=300)
        pdf = out_dir / (pp.stem + ".pdf")
        subprocess.run(["pdftoppm", "-r", "80", "-png", str(pdf), str(out_dir / pp.stem)], check=True)
        print("rendered", pdf)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir", type=Path, default=g.ROOT.parent / "DP4")
    ap.add_argument("--render-dir", type=Path, default=None, help="also render PDF/PNG with LibreOffice into this directory")
    a = ap.parse_args()
    base = g.ROOT.parent / "DP4" / "DP4-slides-draft.pptx"
    a.out_dir.mkdir(parents=True, exist_ok=True)
    jobs = [("DP4-appendix-qa-result.pptx", qa_slides()), ("DP4-complement-design-tactics.pptx", [tactics_slide()])]
    paths = []
    for name, sls in jobs:
        build(sls, a.out_dir / name, base)
        paths.append(a.out_dir / name)
        print("wrote", a.out_dir / name, len(sls), "slide(s)")
    for w in FIT_WARNINGS:
        print("FIT-WARNING", w)
    if a.render_dir:
        render(paths, a.render_dir)


if __name__ == "__main__":
    main()
