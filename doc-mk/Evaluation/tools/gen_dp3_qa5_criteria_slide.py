#!/usr/bin/env python3
"""DP3 QA 기능 정확성(F1 상대 하락률) 평가 기준 슬라이드 1장 생성.

    python gen_dp3_qa5_criteria_slide.py --template <사용자 덱.pptx> [--slide 55] [--scheme 3-1|1-0.1] --out <out.pptx>

사용자 덱의 '처리량 - 평가 기준' 슬라이드(기본 55번)를 템플릿으로 삼아, 그 슬라이드의 표·구간 막대·마커 도형을 복제해
같은 스타일(맑은 고딕, 2F5597, 구간 색)로 만든다. 기준값과 근거 수치는 doc-mk/Requirements/DP3-QS-accuracy-throughput-v2.md 에서 원문 확인한 것만 쓴다.
기준(--scheme):
  3-1   : 하락 <=3% 최소(★1), <=1% 목표(★3), 사이(★2)는 기하 중점 sqrt(1*3)=1.7% (선형 축 0~7%).
  1-0.1 : 하락 <=1% 최소(★1, MLPerf 99% tier), <=0.1% 목표(★3, MLPerf 99.9% tier), 사이(★2)는 기하 중점 sqrt(1*0.1)=0.3% (로그 축 0.05~20%).
          3%(97%)는 잡음 큰 과제용 완화 기준으로 표시만 하고 채택하지 않는다.
두 체계 모두 ★2 경계는 슬라이드 55와 같은 기하 중점 규칙이며 문헌 근거가 없다.
"""
import argparse
import copy
import math

from lxml import etree
from pptx import Presentation

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}

X0, X1 = 411480, 11750039            # 원본 구간 막대의 좌우 끝
AXIS = {"kind": "linear", "lo": 0.0, "hi": 7.0}   # main()에서 scheme별로 설정


def xp(el, path):
    return etree._Element.xpath(el, path, namespaces=NS)


FILL = {"3": "E2EFDA", "2": "FFF2CC", "1": "FBE0D6", "x": "EDEDED"}
BLUE, RED, GRAY = "2F5597", "C00000", "7F7F7F"


def xpos(pct):
    lo, hi = AXIS["lo"], AXIS["hi"]
    if AXIS["kind"] == "log":
        f = (math.log10(max(pct, lo)) - math.log10(lo)) / (math.log10(hi) - math.log10(lo))
    else:
        f = (min(pct, hi) - lo) / (hi - lo)
    return int(X0 + min(max(f, 0.0), 1.0) * (X1 - X0))


def set_runs(par, texts):
    """paragraph의 첫 run들 서식을 유지하며 texts(list)로 교체. 남는 run은 삭제."""
    runs = par.runs
    for i, t in enumerate(texts):
        runs[min(i, len(runs) - 1)].text = t
    for r in runs[len(texts):]:
        r._r.getparent().remove(r._r)


def set_par(par, text):
    set_runs(par, [text])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--slide", type=int, default=55)
    ap.add_argument("--scheme", choices=["3-1", "1-0.1"], default="3-1")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    B = a.scheme == "1-0.1"
    if B:
        AXIS.update(kind="log", lo=0.05, hi=20.0)

    prs = Presentation(a.template)
    lst = prs.slides._sldIdLst
    for i, sid in enumerate(list(lst)):
        if i != a.slide - 1:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
    # 삭제한 슬라이드를 가리키는 section 목록(p14)도 제거하지 않으면 PowerPoint가 복구를 요구한다
    pres = prs.part._element
    for ext in xp(pres, "./p:extLst"):
        pres.remove(ext)
    s = prs.slides[0]
    sh = list(s.shapes)
    title, sub, _rule, metric, btn, src, tbl_shape, ctitle = sh[0], sh[1], sh[2], sh[3], sh[4], sh[5], sh[6], sh[7]

    # 템플릿 도형(복제 원본)
    tpl = dict(zone=sh[8], zlabel=sh[9], tick=sh[20], vline=sh[26], diamond=sh[31], rlabel=sh[32],
               bar=sh[33], cap=sh[34], oval=sh[41], lg_oval=sh[61], lg_text=sh[62])
    tpl = {k: copy.deepcopy(v._element) for k, v in tpl.items()}
    # 구간 차트 영역(8번~)은 모두 지우고 새로 그린다. 이전 페이지 이동 단추(4번)도 지운다(단독 슬라이드라 링크가 끊김).
    for shape in sh[4:5] + sh[8:]:
        shape._element.getparent().remove(shape._element)
    tree = s.shapes._spTree

    # ---- 제목/부제/Metric -------------------------------------------------------------
    set_runs(title.text_frame.paragraphs[0], ["기능 정확성(Accuracy) - 평가 기준"])
    set_runs(sub.text_frame.paragraphs[0],
             ["압축·재사용 전(full recompute) 대비 F1 상대 하락률 — 정확도 기준을 만족하는 조건에서만 처리량 이득을 인정"])
    ps = metric.text_frame.paragraphs
    set_runs(ps[0], ["[Metric] F1 하락률 = (F1_full − F1_후보) ÷ F1_full  ",
                     "(F1_full: 같은 모델·task·precision에서 압축 전 + full recompute 의 F1)"])
    set_par(ps[1], "task별로 판정(평균 통과 + task 위반 = 미충족), 95% CI 상한이 기준 이내일 때만 통과")
    set_par(ps[2], "처리량 연계: 정확도 기준을 통과한 구성에서만 Max SLO goodput 인정 (최소 R ≥ 2.0×, 목표 R ≥ 3.6×)")
    set_runs(src.text_frame.paragraphs[0],
             ["출처: MLPerf Inference Rules (MLCommons) · KVServe 2605.13734 · EVICPRESS 2512.14946 · FastKV 2502.01068 · "
              "CacheGen 2310.07240 · CacheBlend 2405.16444 · KVzip 2505.23416 · Give Me BF16 2411.02355 (ACL 2025)"])

    # ---- 표 ---------------------------------------------------------------------------
    tbl = tbl_shape.table
    rows = list(xp(tbl._tbl, "./a:tr"))
    # 원본 행: 0 헤더, 1 ★3, 2 ★2.5, 3 ★2, 4 ★1.5, 5 ★1, 6 미충족  -> 유지: 0, 1, 3, 5, 6
    for i in (2, 4):
        rows[i].getparent().remove(rows[i])
    MID = round(math.sqrt(1.0 * 0.1), 1) if B else round(math.sqrt(1.0 * 3.0), 1)
    spec_a = [
        ("★ 3", "하락 ≤ 1%  (목표)",
         ["목표(최대): 정확도를 reference의 99% 이상으로 묶고 그 조건에서만 성능을 인정하는 업계 표준과 같은 수준",
          "MLPerf Inference LLM 행 99%(Llama3.1-405B 장문맥 포함) · 양자화 8-bit 평균 recovery 99.75% · FastKV 평균 1% 이내에서 처리량 최대 4.82×(up to)"],
         ["MLPerf Inference Rules, Give Me BF16 (2411.02355), FastKV (2502.01068)"]),
        ("★ 2", f"1% < 하락 ≤ {MID}%",
         [f"★1과 ★3의 기하 중점(√(1×3)={MID}%). 문헌 근거 없음(구간 분할용)."],
         ["—"]),
        ("★ 1", f"{MID}% < 하락 ≤ 3%  (최소)",
         ["최소(허용): 지표 잡음이 큰 과제의 업계 기준 97%와 KV 서빙 논문들의 설계 제약 값",
          "MLPerf RAG(LLM judge ±3%)·agent function-calling 97% · KVServe 97% 상대 정확도 제약 · EVICPRESS 3% 품질 하락 허용"],
         ["MLPerf Inference Rules, KVServe (2605.13734), EVICPRESS (2512.14946)"]),
        ("미충족", "하락 > 3%",
         ["허용 한도로 정의한 문헌·표준을 찾지 못함. 관측값: KIVI task 최악 5.3%, DuoAttention Multi-News 13.9% (task 하나의 큰 손실이 평균에 가려짐)"],
         ["KVServe Table 1"]),
    ]
    spec_b = [
        ("★ 3", "하락 ≤ 0.1%  (목표)",
         ["목표(최대): MLPerf Inference의 99.9% 정확도 tier와 같은 수준. 압축·재사용 연구 중에서는 CacheGen의 F1 하락 < 0.1%가 이 수준",
          "※ 0.1%는 현재 실측(n=93, CI 반폭 1.5~13%)으로 판정할 수 없음. task당 약 40만 개 예시가 필요(실측 SD 기준 계산)"],
         ["MLPerf Inference Rules, CacheGen (2310.07240)"]),
        ("★ 2", f"0.1% < 하락 ≤ {MID}%",
         [f"★1과 ★3의 기하 중점(√(1×0.1)={math.sqrt(0.1):.2f}%). 문헌 근거 없음(구간 분할용)."],
         ["—"]),
        ("★ 1", f"{MID}% < 하락 ≤ 1%  (최소)",
         ["최소(허용): MLPerf Inference의 99% 정확도 tier. LLM 대부분 행이 사용(Llama3.1-405B 장문맥 포함)",
          "양자화 8-bit 평균 recovery 99.75%(하락 0.25%) · FastKV 평균 1% 이내에서 처리량 최대 4.82×(up to)"],
         ["MLPerf Inference Rules, Give Me BF16 (2411.02355), FastKV (2502.01068)"]),
        ("미충족", "하락 > 1%",
         ["99% 기준을 넘는 손실. 잡음이 큰 과제(RAG LLM judge ±3%, agent)와 KVServe·EVICPRESS는 97%(하락 3%)를 쓰나 본 기준은 채택하지 않음",
          "관측값: KIVI task 최악 5.3%, DuoAttention Multi-News 13.9% (task 하나의 큰 손실이 평균에 가려짐)"],
         ["MLPerf Inference Rules, KVServe (2605.13734), EVICPRESS (2512.14946)"]),
    ]
    spec = spec_b if B else spec_a
    heights = [219456, 520000, 300000, 460000, 520000] if B else [219456, 460000, 300000, 460000, 360000]
    trs = list(xp(tbl._tbl, "./a:tr"))
    for tr, h in zip(trs, heights):
        tr.set("h", str(h))
    for ri, (star, crit, why, srcs) in enumerate(spec, start=1):
        cells = tbl.rows[ri].cells
        set_par(cells[0].text_frame.paragraphs[0], star)
        set_par(cells[1].text_frame.paragraphs[0], crit)
        for cell, texts in ((cells[2], why), (cells[3], srcs)):
            tps = cell.text_frame.paragraphs
            for k, t in enumerate(texts):
                if k < len(tps):
                    set_runs(tps[k], [t])
                else:  # 문단 추가: 마지막 문단 복제
                    new = copy.deepcopy(tps[-1]._p)
                    tps[-1]._p.addnext(new)
                    cell.text_frame.paragraphs[-1].runs[0].text = t
            for extra in cell.text_frame.paragraphs[len(texts):]:
                extra._p.getparent().remove(extra._p)
    tbl_shape.height = sum(heights)

    # ---- 구간 대비 근거 위치 (선형 축, F1 하락 %) -----------------------------------------
    set_runs(ctitle.text_frame.paragraphs[0],
             ["구간 대비 근거 위치 (F1 상대 하락 %" + (", log 축" if B else "") + ", 압축·재사용 전 대비, 작을수록 좋음) — 표준 1 + 문헌 5편"])
    next_id = [900]

    def add(el, x=None, y=None, w=None, h=None, name=None):
        el = copy.deepcopy(el)
        cnv = xp(el, ".//p:cNvPr")[0]
        for ext in xp(cnv, "./a:extLst"):
            cnv.remove(ext)
        cnv.set("id", str(next_id[0]))
        cnv.set("name", name or f"qa5-{next_id[0]}")
        next_id[0] += 1
        xf = xp(el, ".//a:xfrm")[0]
        off, ext = xf.find(f"{{{A}}}off"), xf.find(f"{{{A}}}ext")
        if x is not None:
            off.set("x", str(int(x)))
        if y is not None:
            off.set("y", str(int(y)))
        if w is not None:
            ext.set("cx", str(int(w)))
        if h is not None:
            ext.set("cy", str(int(h)))
        tree.append(el)
        return el

    def fill(el, hexcol):
        for c in xp(el, "./p:spPr/a:solidFill/a:srgbClr")[:1]:
            c.set("val", hexcol)

    def text(el, t, color=None, align=None):
        runs = xp(el, ".//a:r")
        runs[0].find(f"{{{A}}}t").text = t
        for r in runs[1:]:
            r.getparent().remove(r)
        if color:
            for c in xp(runs[0], "./a:rPr/a:solidFill/*"):
                c.getparent().remove(c)
            sf = runs[0].find(f"{{{A}}}rPr").find(f"{{{A}}}solidFill")
            etree.SubElement(sf, f"{{{A}}}srgbClr").set("val", color)

    CHAR = 78000   # 7pt 라벨의 문자당 대략 폭(EMU, 한글·영문 혼합)

    def put_label(t, x, y, color, align="l", width=None):
        w = width or min(len(t) * CHAR, 11950000 - x)
        el = add(tpl["rlabel"], x=x, y=y, w=w)
        bp = xp(el, ".//a:bodyPr")[0]
        bp.set("wrap", "square")
        xp(el, ".//a:pPr")[0].set("algn", align)
        text(el, t, color=color)
        return el

    ZY, ZH = 4393368, 274320
    if B:
        bounds = [(0.05, 0.1, "3", "★3"), (0.1, 0.3, "2", "★2"), (0.3, 1.0, "1", "★1"), (1.0, 20.0, "x", "☆ 미충족")]
        ticks = ((0.05, "≤0.05%"), (0.1, "0.1%"), (0.3, "0.3%"), (1, "1%"), (3, "3%"), (10, "10%"), (20, "20%"))
        vlines = (0.1, 0.3, 1.0)
    else:
        bounds = [(0.0, 1.0, "3", "★3"), (1.0, 1.7, "2", "★2"), (1.7, 3.0, "1", "★1"), (3.0, 7.0, "x", "☆ 미충족")]
        ticks = ((0, "0%"), (1, "1%"), (1.7, "1.7%"), (3, "3%"), (5, "5%"), (7, "7%+"))
        vlines = (1.0, 1.7, 3.0)
    for lo, hi, key, label in bounds:
        z = add(tpl["zone"], x=xpos(lo), y=ZY, w=xpos(hi) - xpos(lo), h=ZH)
        fill(z, FILL[key])
        zl = add(tpl["zlabel"], x=xpos(lo), y=ZY, w=xpos(hi) - xpos(lo), h=ZH)
        text(zl, label)
    for pct, lab in ticks:
        tk = add(tpl["tick"], x=xpos(pct) - 320040, y=4667688, w=640080, h=146304)
        text(tk, lab)
    for pct in vlines:
        add(tpl["vline"], x=xpos(pct), y=4709160, h=1625803)

    ROW0, STEP = 4850000, 148000
    LBL_DX = 70000
    rows_a = [
        # kind, a, b, color, label
        ("diamond", 1.0, None, RED, "MLPerf Inference: LLM 행 대부분 reference의 99% 이상일 때만 성능 인정 (하락 ≤ 1%)"),
        ("diamond", 3.0, None, RED, "MLPerf: RAG(LLM judge, ±3% 잡음)·agent function-calling 은 reference의 97% (하락 ≤ 3%)"),
        ("diamond", 3.0, None, RED, "KVServe: 97% 상대 정확도 제약 하에서 압축 선택 · EVICPRESS: 3% 품질 하락 허용 하에서 처리량 보고"),
        ("bar", 0.0, 1.0, BLUE, "FastKV: 평균 하락 < 1% 에서 처리량 최대 4.82× (up to)"),
        ("bar", 0.0, 2.0, BLUE, "CacheGen: accuracy 하락 ≤ 2% (F1 하락 < 0.1%), 지연 3.1~4.7× 감소"),
        ("bar", 0.25, 3.12, BLUE, "양자화(Give Me BF16): 8-bit 평균 0.25% ~ W4A16 task 최악 3.1%"),
        ("bar", 2.57, 5.3, GRAY, "KIVI(4.4×) 2.6% → 5.3%"),
        ("barL", 4.52, 13.9, GRAY, "DuoAttention(3.1×) 4.5% → 13.9% (축 밖)"),
        ("bar", 1.8, 5.7, GRAY, "KVServe-Unified(7.4×) 1.8% → 5.7%"),
    ]
    rows_b = [
        ("diamond", 0.1, None, RED, "MLPerf Inference 99.9% tier(Llama2-70b 등): reference의 99.9% 이상 (하락 ≤ 0.1%)"),
        ("diamond", 1.0, None, RED, "MLPerf Inference 99% tier: LLM 행 대부분 reference의 99% 이상 (하락 ≤ 1%)"),
        ("diamond", 3.0, None, GRAY, "완화 기준(미채택): MLPerf RAG·agent 97% · KVServe 97% 제약 · EVICPRESS 3% 허용"),
        ("bar", 0.05, 1.0, BLUE, "FastKV: 평균 하락 < 1% 에서 처리량 최대 4.82× (up to)"),
        ("bar", 0.05, 0.1, BLUE, "CacheGen: F1 하락 < 0.1% (accuracy 하락은 ≤ 2%), 지연 3.1~4.7× 감소"),
        ("bar", 0.25, 3.12, BLUE, "양자화(Give Me BF16): 8-bit 평균 0.25% ~ W4A16 task 최악 3.1%"),
        ("bar", 2.57, 5.3, GRAY, "KIVI(4.4×) 2.6% → 5.3%"),
        ("bar", 4.52, 13.9, GRAY, "DuoAttention(3.1×) 4.5% → 13.9%"),
        ("bar", 1.8, 5.7, GRAY, "KVServe-Unified(7.4×) 1.8% → 5.7%"),
    ]
    rows_def = rows_b if B else rows_a
    for i, (kind, lo, hi, col, txt) in enumerate(rows_def):
        yc = ROW0 + i * STEP
        if kind == "diamond":
            d = add(tpl["diamond"], x=xpos(lo) - 59436, y=yc - 59436)
            fill(d, col)
            put_label(txt, xpos(lo) + 130000, yc - 52322, col)
        else:
            x_lo, x_hi = xpos(lo), xpos(hi)
            b = add(tpl["bar"], x=x_lo, y=yc - 36576, w=x_hi - x_lo, h=73152)
            fill(b, col)
            for xe in (x_lo, x_hi if kind == "bar" else None):
                if xe is not None:
                    c = add(tpl["cap"], x=xe, y=yc - 64008, h=128016)
                    for ln in xp(c, ".//a:ln/a:solidFill/a:srgbClr"):
                        ln.set("val", col)
            need = len(txt) * CHAR
            if kind == "bar" and x_hi + LBL_DX + need <= 11950000:
                put_label(txt, x_hi + LBL_DX, yc - 52322, col)
            else:  # 오른쪽에 자리가 없으면(또는 축 밖으로 나가는 막대) 막대 왼쪽에 우측 정렬
                put_label(txt, x_lo - LBL_DX - need, yc - 52322, col, align="r", width=need)
    # 범례
    LY = 6262000
    lg = [("oval", BLUE, "문헌 보고 범위(허용 기준 아님)"), ("diamond", RED, "허용 기준으로 채택된 값(표준·설계 제약)"),
          ("oval", GRAY, "참고: KVServe Table 1의 6개 task 평균 → task 최악 하락(허용 기준 아님)")]
    if B:
        lg = [("oval", BLUE, "문헌 보고 범위(허용 기준 아님)"), ("diamond", RED, "본 기준에 채택된 값(표준)"),
              ("diamond", GRAY, "완화 기준(미채택)"), ("oval", GRAY, "참고: KVServe Table 1 평균 → task 최악")]
    x = 384048
    for kind, col, t in lg:
        if kind == "oval":
            o = add(tpl["lg_oval"], x=x, y=LY)
            fill(o, col)
        else:
            o = add(tpl["diamond"], x=x - 4572, y=LY - 4572, w=91440, h=91440)
            fill(o, col)
        tx = add(tpl["lg_text"], x=x + 118872, y=6220852, w=len(t) * 80000)
        bp = xp(tx, ".//a:bodyPr")[0]
        bp.set("wrap", "square")
        text(tx, t)
        x += 118872 + len(t) * 80000 + 180000

    # ---- 발표자 노트 ------------------------------------------------------------------
    notes = s.notes_slide.notes_text_frame
    if B:
        notes.text = ("[백업(Appendix) 스크립트 | 본 발표 40분에는 미포함, 질의응답 시 사용]\n"
                      "기능 정확성 평가 기준입니다. 압축·재사용 전 full recompute 대비 F1 상대 하락률이 1% 이하면 ★1(최소), 0.1% 이하면 ★3(목표)입니다. "
                      "1%는 MLPerf Inference가 LLM 대부분 행에 쓰는 reference 99% tier, 0.1%는 같은 규칙의 99.9% tier입니다. "
                      f"★2({MID}%)는 두 값의 기하 중점으로 문헌 근거가 없습니다. 잡음이 큰 과제(RAG, agent)의 97%와 KVServe·EVICPRESS의 3% 허용은 완화 기준으로 표시만 하고 채택하지 않았습니다.\n"
                      "처리량은 정확도 기준을 통과한 구성에서만 인정합니다(MLPerf, KVServe, EVICPRESS가 같은 방식). 문헌에서 정확도 제약 하 처리량 이득은 하한 약 2×, 상한 3.6~5×입니다.\n"
                      "한계: 0.1%는 현재 정확도 실측(n=93)으로 판정할 수 없고(task당 약 40만 개 필요), 문헌 수치는 대부분 7~8B 모델·단일 노드입니다. 1%와 0.1%는 업계 관행 값이며 통계적 근거는 문헌에 없습니다.")
    else:
        notes.text = ("[백업(Appendix) 스크립트 | 본 발표 40분에는 미포함, 질의응답 시 사용]\n"
                      "기능 정확성 평가 기준입니다. 압축·재사용 전 full recompute 대비 F1 상대 하락률이 3% 이하면 ★1(최소), 1% 이하면 ★3(목표)입니다. "
                      "1%는 MLPerf Inference의 LLM 대부분 행이 쓰는 reference 99% 기준, 3%는 MLPerf가 잡음 있는 지표에 쓰는 97%와 KVServe·EVICPRESS가 설계 제약으로 쓰는 값입니다. "
                      f"★2({MID}%)는 두 값의 기하 중점으로 문헌 근거가 없습니다. 3%를 넘는 값을 허용으로 정의한 문헌은 찾지 못했습니다.\n"
                      "처리량은 정확도 기준을 통과한 구성에서만 인정합니다(MLPerf, KVServe, EVICPRESS가 같은 방식). 문헌에서 정확도 제약 하 처리량 이득은 하한 약 2×, 상한 3.6~5×입니다.\n"
                      "한계: 문헌 수치는 대부분 7~8B 모델·단일 노드이고, 정확도 한도 자체의 통계적 근거는 문헌에 없습니다(업계 관행 값).")
    prs.save(a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
