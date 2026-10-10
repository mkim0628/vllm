#!/usr/bin/env python3
"""DP1 Migration Executor 이후의 실행 구조를 참조 deck(DP1-final-style-reference.pptx) 스타일로 한 장에 그린다.

질문: "Migration Executor 다음에 migration을 어떻게 실행하나, logic 블록과 물리 위치의 연결은 어떻게 갱신하고,
       inference engine에는 물리 위치를 어떻게 알리나?"
근거: doc-mk/vllm-ai-data-migration-architecture.md (§2 Boundary, §4 Module, §7 State machine, §10 Component, §11 SchedulerOutput,
      §12/§13 Sequence, §15 KV identity, §17 Transfer handler, §19 Failure, §22 Code change points, §25 Phase-1 flow).

사용:  /usr/local/bin/python3 doc-mk/Evaluation/tools/gen_dp1_exec_flow_pptx.py [--ref ...] [--out ...]
스타일 helper(add_box/arrow/표 셀 서식)는 gen_c1_complement_styled_pptx.py 를 재사용한다.
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gen_c1_complement_styled_pptx as G
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE

SLIDE_W, SLIDE_H = G.SLIDE_W, G.SLIDE_H
EMU_PT = 12700.0

# 색 (참조 palette)
BLUE, BLUE_L = G.BLUE_FILL, G.BLUE_LINE        # DP1 기존 블록
NEW, NEW_L = G.NEW_FILL, G.NEW_LINE            # 신규 (공통 migration subsystem)
CHG, CHG_L = G.CHG_FILL, G.CHG_LINE            # 기존 vLLM 수정
GRAY, GRAY_L = G.GRAY_FILL, G.GRAY_LINE        # 기존 vLLM 재사용
TXT, SUB = G.TXT, G.SUBTXT

OUT_DEFAULT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "DP1",
                                            "DP1-migration-execution-structure.pptx"))


def T(t, sz=800, b=False, c=None):
    return (t, sz, b, c or TXT)


def build(ref, out):
    prs = Presentation(ref)
    slide = prs.slides[0]
    sp_tree = slide.shapes._spTree
    KEEP = {3, 5, 8, 39, 48, 31, 32, 38, 70, 57, 137, 59, 60, 61, 62, 63, 64, 67, 97, 98, 101}
    for sh in list(slide.shapes):
        if sh.shape_id not in KEEP:
            sp_tree.remove(sh._element)
    by_id = {sh.shape_id: sh for sh in slide.shapes}

    # ---- 제목 / 부제 (참조 run 서식 유지)
    title = by_id[3]
    r0 = title.text_frame.paragraphs[0].runs
    for r, t in zip(r0, ["DP1. ", "Migration 실행 구조 ", "(Executor 이후)", " ", "- ", "설계"]):
        r.text = t
    sub = by_id[8]
    texts = ["Migration 실행", " · ", "Location commit", " · ", "Inference engine 위치 통지", "", ""]
    for i, r in enumerate(sub.text_frame.paragraphs[0].runs):
        r.text = texts[i] if i < len(texts) else ""
    G.TEXT_RECORDS.append(dict(name="title", w=title.width, h=title.height, ins=(91440, 45720, 91440, 45720), wrap="square",
                               paras=[[("DP1. Migration 실행 구조 (Executor 이후) - 설계", 2000, False, None)]], title=True))

    # ---- 구조 표 (참조 표 스타일): 헤더 1행 + 다이어그램 1행
    X0, Y0, TW, LBL = 249035, 1044054, 11747501, 875393
    ROW_H = [342786, 4730000]
    gf = G.add_table(slide, "표 구조", X0, Y0, [LBL, TW - LBL], ROW_H)
    tb = gf.table
    G.fmt_cell(tb.cell(0, 0), [[("설계 안", 1200, False, G.TBL_LBL)]], fill=G.TBL_FILL, algn="ctr", name="hdr-label", width=LBL)
    G.fmt_cell(tb.cell(0, 1), [[("DP1 결정 → 공통 Migration subsystem ", 1200, True, None),
                                ("(WHAT/WHERE는 DP1, HOW/WHEN/COMMIT은 Migration subsystem)", 1200, True, G.RED)]],
               name="hdr-text", width=TW - LBL)
    G.fmt_cell(tb.cell(1, 0), [[("구조", 1200, False, G.TBL_LBL)]], fill=G.TBL_FILL, algn="ctr", name="row-label", width=LBL)
    G.fmt_cell(tb.cell(1, 1), [[("", 1050, False, None)]], name="diagram-cell")
    diag_top, diag_bot = Y0 + ROW_H[0], Y0 + sum(ROW_H)

    q = slide.shapes.add_textbox(Emu(1111250), Emu(1397380), Emu(10700000), Emu(261610))
    q.name = "headline"
    G.set_text(q, [[("“Policy decides ", 1100, False, None), ("WHAT and WHERE", 1100, False, G.LINK_BLUE),
                    (". Migration subsystem decides ", 1100, False, None), ("HOW, WHEN, and COMMIT", 1100, False, G.LINK_BLUE),
                    ("“", 1100, False, None)]],
               algn="l", anchor="t", name="headline")

    # ---- 레이아웃 좌표 (EMU)
    LA_X, LA_W = 1150000, 1750000                     # Lane A: Decision Plane (DP1)
    LB_X, LB_W = 3100000, 4900000                     # Lane B: Control Plane (EngineCore)
    LC_X, LC_W = 8250000, 3700000                     # Lane C: Data Plane (Worker)
    ST_X, ST_W = 3100000, 850000                      # DataLocationStore (tall)
    CL_X, CL_W = 4100000, 1900000                     # lane B col L
    CR_X, CR_W = 6100000, 1900000                     # lane B col R
    R1, R2, R3, R4, R5 = 2060000, 2660000, 3380000, 4100000, 4800000
    H1, H2, H3 = 400000, 540000, 500000

    # lane header labels
    for nm, x, w, t in (("lane A", LA_X, LA_W, "Decision Plane (DP1, 기존)"),
                        ("lane B", LB_X, LB_W, "Migration Control Plane  -  EngineCore process (신규·공통)"),
                        ("lane C", LC_X, LC_W, "Transfer Data Plane  -  Worker process")):
        G.add_label(slide, nm, x, 1730000, w, 230000, [[T(t, 800, True)]], algn="ctr", border=GRAY_L)

    def bx(name, x, y, w, h, fill, line, lines, lw=12700, algn="ctr"):
        return G.add_box(slide, name, x, y, w, h, fill, line, lines, lw=lw, algn=algn, ins=(36000, 18000, 36000, 18000))

    # ---- Lane A (DP1 기존: blue)
    bx("Migration Data Selector", LA_X, R1, LA_W, 400000, BLUE, BLUE_L, [[T("Migration Data", 850, True)], [T("Selector", 850, True)]])
    bx("Migration Executor", LA_X, 2740000, LA_W, 560000, BLUE, BLUE_L,
       [[T("Migration Executor", 900, True)], [T("(DP1 출구 boundary:", 700, False, SUB)], [T("결정만 넘기고 복사는 안 함)", 700, False, SUB)]])
    G.add_label(slide, "intent", LA_X, 3400000, LA_W, 1250000,
                [[T("MigrationIntent", 800, True, G.RED)],
                 [T("object id (logical)", 700)], [T("source tier → target tier", 700)],
                 [T("action: MOVE · REPLICATE ·", 700)], [T("DROP · REMAP · RECLASSIFY", 700)],
                 [T("reason · priority", 700)], [T("(foreground / background)", 700)]], algn="l", border=G.RED, dash="dash")

    # ---- Lane B (신규: orange; 기존 vLLM 수정: green; 재사용: gray)
    bx("MigrationCoordinator", LB_X, R1, LB_W, H1, NEW, NEW_L,
       [[T("① MigrationCoordinator", 880, True), T("  Intent 수신·검증, job lifecycle 관리", 720)]])
    bx("DataLocationStore", ST_X, R2, ST_W, R5 + 580000 - R2, NEW, NEW_L,
       [[T("③ Data", 800, True)], [T("Location", 800, True)], [T("Store", 800, True)], [T("", 600)],
        [T("logical id →", 680)], [T("physical location", 680)], [T("(resource,", 680)], [T("alloc, offset,", 680)], [T("version)", 680)], [T("", 600)],
        [T("copy 중에는", 680)], [T("source가", 680)], [T("authoritative", 680, True)], [T("", 600)],
        [T("⑨ commit 시에만", 680, True, G.RED)], [T("target으로 전환", 680, True, G.RED)], [T("(version+1)", 680, True, G.RED)]])
    bx("KVDataAdapter", CL_X, R2, CL_W, H2, CHG, CHG_L,
       [[T("② KVDataAdapter", 850, True)], [T("can_migrate_now (sealed?)", 700)], [T("source pin / 이동 중 free 금지", 700)]])
    bx("MigrationPlanner", CR_X, R2, CR_W, H2, NEW, NEW_L,
       [[T("④ MigrationPlanner", 850, True)], [T("target reserve (MemoryResourceRegistry)", 680)], [T("path: direct / staged / multi-hop", 680)]])
    bx("MigrationScheduler", CR_X, R3, CR_W, H3, NEW, NEW_L,
       [[T("⑤ MigrationScheduler", 850, True)], [T("foreground(blocking) 우선, background는 link 몫 내", 680)]])
    bx("SchedulerOutput", CL_X, R4, CR_X + CR_W - CL_X, H3, CHG, CHG_L,
       [[T("⑥ SchedulerOutput (step boundary 공유)", 850, True)],
        [T("migration_commands[ job, transfer_steps, src/dst spec ]  +  migration_dependencies[ request, job_ids, wait_before=MODEL_FORWARD ]", 680)]])
    bx("CompletionReconciler", CL_X, R5, CR_X + CR_W - CL_X, 580000, NEW, NEW_L,
       [[T("⑧⑨ CompletionReconciler", 850, True), T("  version·source 유효 검증 → atomic commit(location=target, version+1) → source 해제·unpin", 700)]])

    # ---- Lane C (Worker)
    G.add_label(slide, "worker note", LC_X, 2060000, LC_W, 1950000,
                [[T("Worker = copy executor", 820, True)],
                 [T("위치 결정·commit 권한 없음", 720)],
                 [T("(logical migration authority는 EngineCore)", 720)],
                 [T("", 500)],
                 [T("Job 상태 (EngineCore가 관리)", 780, True)],
                 [T("PENDING → PREPARING → COPYING", 700)],
                 [T("→ VERIFYING → COMMITTING → COMPLETED", 700)],
                 [T("", 500)],
                 [T("실패·stale·worker crash", 780, True, G.RED)],
                 [T("→ target 폐기, source authoritative 유지", 700)],
                 [T("→ 재계획 (rollback)", 700)]], algn="l", border=GRAY_L, dash="dash")
    bx("MigrationExecutor (Worker)", LC_X, R4, LC_W, H3, NEW, NEW_L,
       [[T("MigrationExecutor (Worker)", 850, True)], [T("Command batch 수신, 비동기 handle 발급", 700)]])
    bx("TransferRouter", LC_X, R5, LC_W, 580000, GRAY, GRAY_L,
       [[T("⑦ TransferRouter → TransferHandler", 830, True)], [T("CUDA P2P · Host DMA · CXL · NVMe 중 선택, source 유지한 채 target에 쓰기", 680)],
        [T("HBM · ScHBM · CXL-PNM · DRAM · HBF · SSD-PIM 사이 (Phase 1: 기존 OffloadingWorker 재사용)", 650, False, SUB)]])

    # ---- 화살표
    def top(n): x, y, w, h = G.NODES[n]; return (x + w // 2, y)
    def bot(n): x, y, w, h = G.NODES[n]; return (x + w // 2, y + h)
    def lf(n): x, y, w, h = G.NODES[n]; return (x, y + h // 2)
    def rt(n): x, y, w, h = G.NODES[n]; return (x + w, y + h // 2)
    BLU, ORG, RED, GREY = ("s", "accent1", {}), G.NEW_LINE, G.RED, G.GRAY_LINE

    G.arrow(slide, "a-sel-exec", bot("Migration Data Selector"), top("Migration Executor"), BLU, from_node="Migration Data Selector", to_node="Migration Executor")
    # ① Executor -> Coordinator (polyline, lane 사이 통로)
    xe, ye = rt("Migration Executor")
    xc, yc = lf("MigrationCoordinator")
    G.polyline(slide, "a-exec-coord", [(xe, ye), (3000000, ye), (3000000, yc), (xc, yc)], RED, 12700, "Migration Executor", "MigrationCoordinator")
    # Coordinator -> store / kv adapter / planner
    x, y = bot("MigrationCoordinator")
    G.arrow(slide, "a-coord-store", (ST_X + ST_W // 2, y), top("DataLocationStore"), ORG, from_node="MigrationCoordinator", to_node="DataLocationStore")
    G.arrow(slide, "a-coord-kv", (CL_X + CL_W // 2, y), top("KVDataAdapter"), ORG, from_node="MigrationCoordinator", to_node="KVDataAdapter")
    G.arrow(slide, "a-coord-plan", (CR_X + CR_W // 2, y), top("MigrationPlanner"), ORG, from_node="MigrationCoordinator", to_node="MigrationPlanner")
    # Planner -> Scheduler -> SchedulerOutput
    G.arrow(slide, "a-plan-sched", bot("MigrationPlanner"), top("MigrationScheduler"), ORG, from_node="MigrationPlanner", to_node="MigrationScheduler")
    sx, sy = bot("MigrationScheduler")
    G.arrow(slide, "a-sched-so", (sx, sy), (sx, G.NODES["SchedulerOutput"][1]), ORG, from_node="MigrationScheduler", to_node="SchedulerOutput")
    # SchedulerOutput -> Worker executor
    G.arrow(slide, "a-so-exec", rt("SchedulerOutput"), lf("MigrationExecutor (Worker)"), ORG, 15000, "SchedulerOutput", "MigrationExecutor (Worker)")
    # Executor(W) -> Router -> Memory
    G.arrow(slide, "a-exec-router", bot("MigrationExecutor (Worker)"), top("TransferRouter"), ORG, from_node="MigrationExecutor (Worker)", to_node="TransferRouter")
    # ⑧ completion: Router -> Reconciler (lane C -> B)
    G.arrow(slide, "a-complete", lf("TransferRouter"), rt("CompletionReconciler"), RED, 15000, "TransferRouter", "CompletionReconciler")
    # ⑨ commit: Reconciler -> Store
    G.arrow(slide, "a-commit", lf("CompletionReconciler"), (ST_X + ST_W, G.NODES["CompletionReconciler"][1] + G.NODES["CompletionReconciler"][3] // 2),
            RED, 15000, "CompletionReconciler", "DataLocationStore")

    # ---- 하단: Inference engine에 물리 위치를 알리는 경로
    SY, SH = 5580000, 500000
    G.add_label(slide, "strip title", 1150000, 5420000, 5600000, 150000,
                [[T("Inference engine에 물리 위치를 알리는 경로 (logical block id는 그대로, 위치만 교체)", 780, True, G.RED)]], algn="l")
    G.add_label(slide, "legend", 6900000, 5420000, 5050000, 150000,
                [[T("색: ", 650, True), T("파랑=DP1 기존 · ", 650, False, ("s", "accent1", {"lumMod": 75000})),
                  T("주황=신규(공통 migration) · ", 650, False, G.NEW_LINE), T("초록=기존 vLLM 수정 · ", 650, False, G.CHG_LINE),
                  T("회색=재사용", 650, False, SUB)]], algn="r")
    boxes = [
        ("N1", "⑨ commit 결과", "logical id → 새 physical location (resource, alloc, offset, version+1). 이전 HBM block unpin·release", NEW, NEW_L),
        ("N2", "KVCacheManager / BlockPool", "Scheduler는 logical block id만 봄. 위치는 KVDataAdapter가 DataLocationStore에서 조회", CHG, CHG_L),
        ("N3", "다음 step: Scheduler.schedule()", "필요 KV가 하위 tier면 foreground promotion job + dependency를 SchedulerOutput에 실음", CHG, CHG_L),
        ("N4", "Worker: dependency 대기", "migration job 완료까지 MODEL_FORWARD 진입 전 wait (Phase 1: forward 앞 barrier)", CHG, CHG_L),
        ("N5", "GPUModelRunner.execute_model()", "required KV가 HBM의 새 physical block에 있음 → attention 수행", CHG, CHG_L),
    ]
    bw, gap = 1960000, 250000
    for i, (n, h, d, f, l) in enumerate(boxes):
        bx(n, 1150000 + i * (bw + gap), SY, bw, SH, f, l, [[T(h, 760, True)], [T(d, 660)]])
    for i in range(len(boxes) - 1):
        a, b = boxes[i][0], boxes[i + 1][0]
        G.arrow(slide, f"a-{a}-{b}", rt(a), lf(b), GREY, 12700, a, b)
    # ---- takeaway bar
    TY = 6195000
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(234949), Emu(TY), Emu(11747501), Emu(300000))
    bar.name = "takeaway"
    G.strip_style(bar)
    G.set_fill(bar, ("s", "bg1", {"lumMod": 50000}))
    G.set_line(bar, None)
    G.set_text(bar, [[("핵심: ", 1000, True, G.WHITE),
                      ("copy 중엔 source가 authoritative, 검증 후에만 atomic commit, 새 위치는 다음 step의 SchedulerOutput으로 전달 (설계 단계, 미구현)", 900, True, G.WHITE)]],
               algn="l", anchor="ctr", name="takeaway")
    G.NODES["__takeaway__"] = (234949, TY, 11747501, 300000)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    return prs, (diag_top, diag_bot)


def verify(prs, diag):
    probs = []
    slide = prs.slides[0]
    for sh in slide.shapes:
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > SLIDE_W or sh.top + sh.height > SLIDE_H:
            probs.append(f"OUT OF SLIDE: {sh.name}")
    for rec in G.TEXT_RECORDS:
        if rec.get("title"):
            continue
        l, t, r, b = rec["ins"]
        wpt = (rec["w"] - l - r) / EMU_PT
        hpt = (rec["h"] - t - b) / EMU_PT
        tot = 0.0
        for runs in rec["paras"]:
            n = G.n_lines(runs, wpt) if rec["wrap"] == "square" else 1
            mx = max((sz for _, sz, _, _ in runs), default=900) / 100.0
            tot += n * mx * 1.2
        if tot > hpt + 0.5:
            probs.append(f"TEXT OVERFLOW: {rec['name']} need {tot:.1f}pt > {hpt:.1f}pt")
    names = [n for n in G.NODES if not n.startswith("__")]
    d0, d1 = diag
    for i, a in enumerate(names):
        ax, ay, aw, ah = G.NODES[a]
        if ay < d0 or ay + ah > d1:
            probs.append(f"NODE outside diagram cell: {a}")
        for b in names[i + 1:]:
            bx_, by, bw, bh = G.NODES[b]
            if ax < bx_ + bw and bx_ < ax + aw and ay < by + bh and by < ay + ah:
                probs.append(f"NODE OVERLAP: {a} <-> {b}")

    def on_edge(pt, n):
        x, y, w, h = G.NODES[n]
        px, py = pt
        inx, iny = x - 2 <= px <= x + w + 2, y - 2 <= py <= y + h + 2
        return (inx and (abs(py - y) <= 2 or abs(py - (y + h)) <= 2)) or (iny and (abs(px - x) <= 2 or abs(px - (x + w)) <= 2))
    for name, p1, p2, fn, tn in G.ARROWS:
        if fn and not on_edge(p1, fn):
            probs.append(f"ARROW start not on edge: {name}")
        if tn and not on_edge(p2, tn):
            probs.append(f"ARROW end not on edge: {name}")
        for nn in names:
            if nn in (fn, tn):
                continue
            x, y, w, h = G.NODES[nn]
            (x1, y1), (x2, y2) = p1, p2
            if x1 == x2 and x < x1 < x + w and min(y1, y2) < y + h and max(y1, y2) > y:
                probs.append(f"ARROW crosses node: {name} x {nn}")
            if y1 == y2 and y < y1 < y + h and min(x1, x2) < x + w and max(x1, x2) > x:
                probs.append(f"ARROW crosses node: {name} x {nn}")
    return probs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ref", default=G.DEFAULT_REF)
    ap.add_argument("--out", default=OUT_DEFAULT)
    a = ap.parse_args()
    prs, diag = build(a.ref, a.out)
    probs = verify(prs, diag)
    print("saved:", a.out)
    if probs:
        print("VERIFY PROBLEMS:")
        for p in probs:
            print(" -", p)
        sys.exit(1)
    print("verify OK: bounds / text fit / node overlap / arrow edges")


if __name__ == "__main__":
    main()
