#!/usr/bin/env python3
"""DP1 설계 구조를 '로직/알고리즘'이 아니라 아키텍처 뷰/스타일로 다시 표현한 두 버전(각 1장)을 만든다.

  A: 제어 루프(MAPE-K) 스타일 + 3 Plane 분리      -> DP1-architecture-view-A-controlloop-3plane.pptx
  B: Plug-in(ports & adapters) + 변이점/확장 뷰    -> DP1-architecture-view-B-plugin-ports.pptx

참조 deck(DP1-final-style-reference.pptx)의 master/theme/header/palette를 그대로 쓰고 helper는
gen_c1_complement_styled_pptx.py 를 재사용한다.
근거: dp1-ai-data-migration-decision-architecture.md, vllm-ai-data-migration-architecture.md(§2 plane, §10 component),
      Evaluation/DP1/qa4-modifiability.md (확장 시나리오 module 수, [B+C]).
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gen_c1_complement_styled_pptx as G
import gen_dp1_exec_flow_pptx as X
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, "..", "..", "DP1"))

BLUE, BLUE_L = G.BLUE_FILL, G.BLUE_LINE
NEW, NEW_L = G.NEW_FILL, G.NEW_LINE
CHG, CHG_L = G.CHG_FILL, G.CHG_LINE
GRAY, GRAY_L = G.GRAY_FILL, G.GRAY_LINE
TXT, SUB = G.TXT, G.SUBTXT
T = X.T


def reset():
    G.NODES.clear(); G.ARROWS.clear(); G.TEXT_RECORDS.clear(); G.CELL_RECORDS.clear()


def begin(ref, title_runs, sub_texts, hdr_runs, headline_runs, row_h2=4730000):
    """참조 header 유지, 본문 제거, 구조 표(헤더 행 + 다이어그램 행)와 headline 을 만든다."""
    prs = Presentation(ref)
    slide = prs.slides[0]
    sp_tree = slide.shapes._spTree
    KEEP = {3, 5, 8, 39, 48, 31, 32, 38, 70, 57, 137, 59, 60, 61, 62, 63, 64, 67, 97, 98, 101}
    for sh in list(slide.shapes):
        if sh.shape_id not in KEEP:
            sp_tree.remove(sh._element)
    by_id = {sh.shape_id: sh for sh in slide.shapes}
    title = by_id[3]
    for r, t in zip(title.text_frame.paragraphs[0].runs, title_runs):
        r.text = t
    sub = by_id[8]
    for i, r in enumerate(sub.text_frame.paragraphs[0].runs):
        r.text = sub_texts[i] if i < len(sub_texts) else ""
    G.TEXT_RECORDS.append(dict(name="title", w=title.width, h=title.height, ins=(91440, 45720, 91440, 45720), wrap="square",
                               paras=[[("".join(title_runs), 2000, False, None)]], title=True))
    X0, Y0, TW, LBL = 249035, 1044054, 11747501, 875393
    ROW_H = [342786, row_h2]
    gf = G.add_table(slide, "표 구조", X0, Y0, [LBL, TW - LBL], ROW_H)
    tb = gf.table
    G.fmt_cell(tb.cell(0, 0), [[("설계 안", 1200, False, G.TBL_LBL)]], fill=G.TBL_FILL, algn="ctr", name="hdr-label", width=LBL)
    G.fmt_cell(tb.cell(0, 1), [hdr_runs], name="hdr-text", width=TW - LBL)
    G.fmt_cell(tb.cell(1, 0), [[("구조", 1200, False, G.TBL_LBL)]], fill=G.TBL_FILL, algn="ctr", name="row-label", width=LBL)
    G.fmt_cell(tb.cell(1, 1), [[("", 1050, False, None)]], name="diagram-cell")
    q = slide.shapes.add_textbox(Emu(1111250), Emu(1397380), Emu(10700000), Emu(261610))
    q.name = "headline"
    G.set_text(q, [headline_runs], algn="l", anchor="t", name="headline")
    return prs, slide, (Y0 + ROW_H[0], Y0 + sum(ROW_H))


def bx(slide, name, x, y, w, h, fill, line, lines, lw=12700):
    return G.add_box(slide, name, x, y, w, h, fill, line, lines, lw=lw, ins=(36000, 18000, 36000, 18000))


def ends(slide):
    TY = 6195000
    return TY


def takeaway(slide, text):
    from pptx.enum.shapes import MSO_SHAPE
    TY = 6195000
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(234949), Emu(TY), Emu(11747501), Emu(300000))
    bar.name = "takeaway"
    G.strip_style(bar)
    G.set_fill(bar, ("s", "bg1", {"lumMod": 50000}))
    G.set_line(bar, None)
    G.set_text(bar, [[("한 줄 요약: ", 1000, True, G.WHITE), (text, 900, True, G.WHITE)]], algn="l", anchor="ctr", name="takeaway")


def top(n): x, y, w, h = G.NODES[n]; return (x + w // 2, y)
def bot(n): x, y, w, h = G.NODES[n]; return (x + w // 2, y + h)
def lf(n): x, y, w, h = G.NODES[n]; return (x, y + h // 2)
def rt(n): x, y, w, h = G.NODES[n]; return (x + w, y + h // 2)


# =====================================================================================================
# 버전 A: 제어 루프(MAPE-K) 스타일 + 3 Plane
# =====================================================================================================
def build_A(ref, out):
    reset()
    prs, slide, diag = begin(
        ref, ["DP1. ", "아키텍처 ", "(제어 루프 + 3 Plane)", " ", "- ", "설계"], ["제어 루프 스타일", " · ", "Plane 분리", " · ", "Plug-in 정책", "", ""],
        [("Decision Plane의 DP1 구조 ", 1200, True, None), ("(자기 적응 제어 루프: Observe → Assess → Decide → Request)", 1200, True, G.RED)],
        [("“Policy decides ", 1100, False, None), ("WHAT and WHERE", 1100, False, G.LINK_BLUE), (". Migration subsystem decides ", 1100, False, None),
         ("HOW, WHEN, and COMMIT", 1100, False, G.LINK_BLUE), ("“", 1100, False, None)])
    # lanes
    LD_X, LD_W = 1150000, 6000000
    LC_X, LC_W = 7550000, 2000000
    LW_X, LW_W = 9950000, 2000000
    for nm, x, w, t in (("lane D", LD_X, LD_W, "Decision Plane (DP1)  -  제어 루프 (MAPE-K 스타일)"),
                        ("lane C", LC_X, LC_W, "Control Plane (EngineCore, 공통)"),
                        ("lane W", LW_X, LW_W, "Data Plane (Worker, 공통)")):
        G.add_label(slide, nm, x, 1730000, w, 230000, [[T(t, 800, True)]], algn="ctr", border=GRAY_L)
    R1, H1 = 2060000, 950000
    bw, gp = 1080000, 150000
    x0 = LD_X
    bx(slide, "Event", x0, R1, 900000, H1, GRAY, GRAY_L,
       [[T("Event", 800, True)], [T("Resource", 700)], [T("Manager", 700)], [T("TELEMETRY", 600, False, SUB)], [T("ALLOCATED·FREED", 600, False, SUB)], [T("ACCESSED", 600, False, SUB)]])
    xs = [x0 + 900000 + gp + i * (bw + gp) for i in range(4)]
    bx(slide, "Observation", xs[0], R1, bw, H1, BLUE, BLUE_L,
       [[T("① Observation", 800, True)], [T("무슨 일이 있었나", 680, False, SUB)], [T("C1: 용량·BW·부하", 660)], [T("C2: data별 접근 행동", 660)]])
    bx(slide, "Assessment", xs[1], R1, bw, H1, NEW, NEW_L,
       [[T("② Situation", 800, True)], [T("Assessment", 800, True)], [T("개입이 필요한가", 680, False, SUB)], [T("C1: 압력·추세", 660)], [T("C2: 재사용·수명 예측", 660)]])
    bx(slide, "Decision", xs[2], R1, bw, H1, NEW, NEW_L,
       [[T("③ Placement", 800, True)], [T("Decision", 800, True)], [T("무엇을 어디로", 680, False, SUB)], [T("capability·SLO 필터", 660)], [T("C1: Data-Memory Affinity", 640)]])
    bx(slide, "Request", xs[3], R1, bw, H1, BLUE, BLUE_L,
       [[T("④ Actuation", 800, True)], [T("Request", 800, True)], [T("결정을 계약으로 발행", 660, False, SUB)], [T("MigrationIntent", 700, True, G.RED)], [T("(복사는 하지 않음)", 640, False, SUB)]])
    G.arrow(slide, "e-obs", rt("Event"), lf("Observation"), G.GRAY_LINE, from_node="Event", to_node="Observation")
    G.arrow(slide, "obs-ass", rt("Observation"), lf("Assessment"), ("s", "accent1", {}), from_node="Observation", to_node="Assessment")
    G.arrow(slide, "ass-dec", rt("Assessment"), lf("Decision"), ("s", "accent1", {}), from_node="Assessment", to_node="Decision")
    G.arrow(slide, "dec-req", rt("Decision"), lf("Request"), ("s", "accent1", {}), from_node="Decision", to_node="Request")
    # Shared Knowledge
    KY, KH = 3300000, 600000
    bx(slide, "Knowledge", LD_X, KY, LD_W, KH, BLUE, BLUE_L,
       [[T("Shared Knowledge", 850, True)], [T("Data Object Registry: 위치·크기·tier·이동 가능 상태 (C1 type-agnostic / C2 type-aware)", 680)],
        [T("Memory Registry: 메모리별 capability descriptor (용량·BW·지연·primitive)", 680)]])
    for n in ("Observation", "Assessment", "Decision"):
        x, y = bot(n)
        G.arrow(slide, f"k-{n}", (x, y), (x, KY), ("s", "accent1", {}), 9525, n, "Knowledge", both=True)
    # Backend I/F + plug-in label
    BY = 4150000
    bx(slide, "Backend IF", LD_X, BY, 2700000, 600000, GRAY, GRAY_L,
       [[T("Memory Backend I/F  (plug-in)", 780, True)], [T("HBM · ScHBM · CXL-PNM · DRAM · HBF · SSD-PIM", 640)], [T("descriptor·telemetry를 같은 계약으로 제공", 640, False, SUB)]])
    x, y = top("Backend IF")
    G.arrow(slide, "be-k", (x, y), (x, KY + KH), G.GRAY_LINE, from_node="Backend IF", to_node="Knowledge")
    G.add_label(slide, "plug-in", LD_X + 2900000, BY, 3100000, 600000,
                [[T("정책 plug-in 영역: ②③", 780, True, G.NEW_LINE)], [T("C1(Resource State-driven)과 C2(Behavior-driven)가 교체", 680)],
                 [T("①④, Knowledge, 계약은 두 후보 공통", 680)]], algn="l", border=G.NEW_LINE)
    # Control / Data planes
    bx(slide, "Control", LC_X, R1, LC_W, H1, NEW, NEW_L,
       [[T("Migration Control", 820, True)], [T("Coordinator: Intent 검증·job 관리", 660)], [T("Planner: target 예약·경로", 660)], [T("Scheduler: 우선순위·link 예산", 660)]])
    bx(slide, "Data", LW_X, R1, LW_W, H1, NEW, NEW_L,
       [[T("Migration Executor", 820, True)], [T("Transfer Handlers", 780, True)], [T("CUDA P2P · Host DMA", 660)], [T("CXL · NVMe", 660)], [T("copy만 수행", 660, False, SUB)]])
    bx(slide, "Location", LC_X, KY, LC_W, KH, NEW, NEW_L,
       [[T("Location Store", 820, True)], [T("copy 검증 후 atomic commit", 660)], [T("(version 증가, source 해제)", 650, False, SUB)]])
    bx(slide, "Tiers", LW_X, KY, LW_W, KH, GRAY, GRAY_L,
       [[T("Memory tiers", 820, True)], [T("source 유지한 채 target에 쓰기", 650, False, SUB)]])
    G.arrow(slide, "req-ctl", rt("Request"), lf("Control"), G.RED, 15000, "Request", "Control")
    G.arrow(slide, "ctl-data", rt("Control"), lf("Data"), G.NEW_LINE, 15000, "Control", "Data")
    G.arrow(slide, "data-tier", bot("Data"), top("Tiers"), G.GRAY_LINE, from_node="Data", to_node="Tiers")
    G.arrow(slide, "tier-loc", lf("Tiers"), rt("Location"), G.RED, 15000, "Tiers", "Location")
    G.arrow(slide, "loc-k", lf("Location"), rt("Knowledge"), G.RED, 15000, "Location", "Knowledge")
    # 계약 라벨 (화살표 위)
    for nm, x, y, w, t, c in (("lb1", 7150000, R1 + H1 // 2 - 190000, 400000, "Intent", G.RED),
                              ("lb2", 9550000, R1 + H1 // 2 - 190000, 400000, "명령", G.NEW_LINE),
                              ("lb3", 9550000, KY + KH // 2 - 190000, 400000, "완료", G.RED),
                              ("lb4", 7150000, KY + KH // 2 - 190000, 400000, "갱신", G.RED)):
        G.add_label(slide, nm, x, y, w, 150000, [[T(t, 620, True, c)]], algn="ctr")
    G.add_label(slide, "contracts", LC_X, BY, 4400000, 600000,
                [[T("Plane 사이 계약 (이 계약만 넘나듦)", 780, True)],
                 [T("MigrationIntent: object id · src→dst · action(MOVE/REPLICATE/DROP/REMAP/RECLASSIFY) · priority", 640)],
                 [T("MigrationCommand → Completion → Location 갱신. 결정 주체와 commit 주체는 분리", 640)]], algn="l", border=GRAY_L)
    # 스타일 strip
    G.add_label(slide, "styles", LD_X, 4900000, 5000000, 150000, [[T("적용한 아키텍처 스타일", 780, True, G.RED)]], algn="l")
    SW, SG, SY, SH = 2550000, 200000, 5080000, 900000
    styles = [
        ("Control loop (MAPE-K)", "관찰 → 판단 → 결정 → 실행 요청을 닫힌 루프로 구성. 루프 구조는 고정하고 판단 방식만 바꾼다"),
        ("3-Plane 분리", "결정(WHAT/WHERE) · 제어(HOW/WHEN/COMMIT) · 데이터(복사). 책임을 나누고 계약으로만 통신한다"),
        ("Plug-in (변이점)", "정책(C1/C2), 메모리 backend, AI data adapter는 인터페이스 뒤에서 교체된다"),
        ("Shared Knowledge", "위치와 capability는 한 곳에 둔다. location은 commit 결과로만 바뀌어 일관성을 지킨다"),
    ]
    for i, (h, d) in enumerate(styles):
        bx(slide, f"S{i}", LD_X + i * (SW + SG), SY, SW, SH, GRAY, GRAY_L, [[T(h, 800, True)], [T(d, 680)]])
    takeaway(slide, "알고리즘은 정책 plug-in으로 내리고, 구조는 제어 루프 + 3 Plane + 계약으로 보인다. 알고리즘 상세는 부록. 설계 단계, 미구현")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    return prs, diag


# =====================================================================================================
# 버전 B: Plug-in (ports & adapters) + 변이점 / 확장 뷰
# =====================================================================================================
def build_B(ref, out):
    reset()
    prs, slide, diag = begin(
        ref, ["DP1. ", "아키텍처 ", "(Plug-in / Port & Adapter)", " ", "- ", "설계"], ["고정 Core", " · ", "교체 가능한 변이점", " · ", "확장 비용", "", ""],
        [("고정 Core + Policy 변이점 ", 1200, True, None), ("(C1/C2는 같은 Policy SPI의 두 구현, 새 메모리·데이터는 adapter 추가)", 1200, True, G.RED)],
        [("“", 1100, False, None), ("Core는 고정하고", 1100, False, G.LINK_BLUE), (", 정책·메모리·AI data는 ", 1100, False, None),
         ("port 뒤 plug-in으로 교체", 1100, False, G.LINK_BLUE), ("“", 1100, False, None)], row_h2=4730000)
    LX, LW_ = 1150000, 2700000                       # 입력 port/adapters
    CX = 4150000                                     # core
    RX, RW = 8750000, 3200000                        # 출력
    for nm, x, w, t in (("h-in", LX, LW_, "입력 Port  ←  Adapter (plug-in)"),
                        ("h-core", CX, 4300000, "DP1 Core (고정) + Policy 변이점"),
                        ("h-out", RX, RW, "출력 Port  →  공통 Migration subsystem")):
        G.add_label(slide, nm, x, 1730000, w, 230000, [[T(t, 800, True)]], algn="ctr", border=GRAY_L)

    def pair(i, pname, ptext, aname, alines):
        y = 2050000 + i * 820000
        bx(slide, pname, LX, y, LW_, 260000, BLUE, BLUE_L, [[T(ptext, 760, True)]])
        bx(slide, aname, LX, y + 380000, LW_, 400000, NEW, NEW_L, alines)
        G.arrow(slide, f"a-{aname}", top(aname), (G.NODES[aname][0] + LW_ // 2, y + 300000), G.NEW_LINE, 9525, aname, pname) if False else None
        return y

    y1 = pair(0, "P1", "Telemetry / Event port", "A1", [[T("Resource Manager telemetry", 700, True)], [T("vLLM scheduler·block pool 통계, 이벤트", 640)]])
    y2 = pair(1, "P2", "Memory Descriptor port", "A2", [[T("Memory Backend adapter x6", 700, True)], [T("HBM·ScHBM·CXL-PNM·DRAM·HBF·SSD-PIM (새 메모리 = adapter 1개)", 620)]])
    y3 = pair(2, "P3", "AI Data port", "A3", [[T("AI Data adapter", 700, True)], [T("KV · LoRA · MoE · RAG · Agent memory", 640)]])
    for p, a in (("P1", "A1"), ("P2", "A2"), ("P3", "A3")):
        x, y = bot(p)
        G.arrow(slide, f"v-{p}", (G.NODES[a][0] + LW_ // 2, G.NODES[a][1]), (x, y), G.NEW_LINE, 9525, a, p)
    # core
    bx(slide, "K1", CX, 2050000, 4300000, 280000, BLUE, BLUE_L, [[T("Event · Migration Scheduler  (고정)", 780, True)]])
    bx(slide, "K2", CX, 2500000, 4300000, 300000, CHG, CHG_L,
       [[T("Policy SPI  (변이점)", 800, True), T("   assess() · decide()  →  MigrationIntent", 680)]])
    bx(slide, "I1", CX, 2950000, 2100000, 750000, NEW, NEW_L,
       [[T("C1  Resource State-driven", 780, True)], [T("압력·추세 감지", 660)], [T("Eviction · Destination Selector", 660)], [T("Data-Memory Affinity (static hint)", 650, False, G.RED)]])
    bx(slide, "I2", CX + 2200000, 2950000, 2100000, 750000, NEW, NEW_L,
       [[T("C2  Behavior-driven", 780, True)], [T("data별 행동 관찰 · 예측", 660)], [T("benefit-vs-cost gating", 660)], [T("type-aware registry", 650, False, G.RED)]])
    bx(slide, "K3", CX, 3850000, 4300000, 400000, BLUE, BLUE_L,
       [[T("Registry port · Selector → Executor boundary  (고정)", 760, True)], [T("Data Object Registry schema는 고정, 정보 모델은 C1 type-agnostic / C2 type-aware", 620, False, SUB)]])
    G.arrow(slide, "p1-k1", (LX + LW_, 2180000), (CX, 2180000), ("s", "accent1", {}), from_node="P1", to_node="K1")
    G.arrow(slide, "k1-k2", bot("K1"), top("K2"), ("s", "accent1", {}), from_node="K1", to_node="K2")
    x2, y2b = bot("K2")
    G.arrow(slide, "k2-i1", (CX + 1050000, y2b), (CX + 1050000, G.NODES["I1"][1]), G.NEW_LINE, 12700, "K2", "I1")
    G.arrow(slide, "k2-i2", (CX + 3250000, y2b), (CX + 3250000, G.NODES["I2"][1]), G.NEW_LINE, 12700, "K2", "I2")
    G.arrow(slide, "i1-k3", (CX + 1050000, G.NODES["I1"][1] + 750000), (CX + 1050000, G.NODES["K3"][1]), G.NEW_LINE, 12700, "I1", "K3")
    G.arrow(slide, "i2-k3", (CX + 3250000, G.NODES["I2"][1] + 750000), (CX + 3250000, G.NODES["K3"][1]), G.NEW_LINE, 12700, "I2", "K3")
    G.arrow(slide, "p2-i1", (LX + LW_, G.NODES["P2"][1] + 130000), (CX, G.NODES["P2"][1] + 130000), ("s", "accent1", {}), 9525, "P2", "I1")
    G.arrow(slide, "p3-k3", (LX + LW_, 3890000), (CX, 3890000), ("s", "accent1", {}), 9525, "P3", "K3")
    # 출력
    bx(slide, "O1", RX, 3850000, RW, 400000, BLUE, BLUE_L, [[T("Migration port", 780, True)], [T("MigrationIntent  (object · src→dst · action · priority)", 640)]])
    bx(slide, "O2", RX, 2050000, RW, 1660000, NEW, NEW_L,
       [[T("공통 Migration subsystem", 840, True)], [T("Control Plane (EngineCore)", 760, True)], [T("Coordinator · Planner · Scheduler", 660)], [T("Location Store (atomic commit)", 660)],
        [T("", 400)], [T("Data Plane (Worker)", 760, True)], [T("Executor · Transfer Handlers", 660)], [T("", 400)],
        [T("commit 결과로만 Registry 갱신", 660, True, G.RED)]])
    G.arrow(slide, "k3-o1", rt("K3"), lf("O1"), G.RED, 15000, "K3", "O1")
    G.arrow(slide, "o1-o2", top("O1"), bot("O2"), G.RED, 15000, "O1", "O2")
    # 확장 표
    G.add_label(slide, "ext-title", LX, 4560000, 10800000, 150000,
                [[T("확장 시나리오별로 고치는 곳과 module 수 (시뮬레이터 복사본에 구현해 측정, [B+C], `qa4-modifiability.md`)", 760, True, G.RED)]], algn="l")
    cw = [3000000, 5900000, 1900000]
    ext = G.add_table(slide, "표 확장", LX, 4730000, cw, [230000] * 5)
    t = ext.table
    mar = (60000, 10000, 60000, 10000)
    hdr = ["확장 시나리오", "고치는 곳 (어디가 변하나)", "module 수 (C1 / C2)"]
    for c, h in enumerate(hdr):
        G.fmt_cell(t.cell(0, c), [[(h, 780, True, None)]], fill=G.TBL_FILL, algn="ctr" if c != 1 else "l", mar=mar, name=f"eh{c}", width=cw[c])
    rows = [
        ("새 메모리 (CXL.mem expander)", "Memory Backend adapter (공통). C2는 type별 선호 목록에 이름이 있어 Selector도 수정", "1 / 2"),
        ("새 AI data 종류 (SPARSE_EMBED)", "AI Data adapter와 op-class hint 표 (공통). C2는 type-aware registry class와 Selector도 수정", "1 / 3"),
        ("정책 교체·조정", "Policy SPI 구현 내부 (C1 Affinity Mapper / C2 Predictor)", "2 / 2"),
        ("새 event 종류 (SLO_ALERT)", "Event schema와 Event Source (두 후보가 공유하는 Core)", "3 / 3"),
    ]
    for r, (a, b, c_) in enumerate(rows, start=1):
        G.fmt_cell(t.cell(r, 0), [[(a, 760, True, None)]], mar=mar, name=f"e{r}0", width=cw[0])
        G.fmt_cell(t.cell(r, 1), [[(b, 740, False, None)]], mar=mar, name=f"e{r}1", width=cw[1])
        G.fmt_cell(t.cell(r, 2), [[(c_, 800, True, G.RED)]], mar=mar, algn="ctr", name=f"e{r}2", width=cw[2])
    # legend
    G.add_label(slide, "legend", RX, 4560000, RW, 150000,
                [[T("파랑=고정 Core · 주황=plug-in/신규 · 초록=변이점(SPI)", 620, False, SUB)]], algn="r")
    takeaway(slide, "C1/C2는 같은 Policy SPI의 두 구현이고, 새 메모리·데이터는 adapter 추가로 확장된다 (module 1~3개). 설계 단계, 미구현")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    return prs, diag


def check_cells(maxrows_h=230000):
    probs = []
    for rec in G.CELL_RECORDS:
        l, t, r, b = rec["mar"]
        wpt = (rec["w"] - l - r) / G.EMU_PT if hasattr(G, "EMU_PT") else (rec["w"] - l - r) / 12700.0
        need = sum(G.n_lines(runs, wpt) * max(sz for _, sz, _, _ in runs) / 100.0 * 1.2 for runs in rec["paras"]) + (t + b) / 12700.0
        if rec["name"].startswith("e") and need * 12700 > maxrows_h + 1000:
            probs.append(f"CELL too tall: {rec['name']} need {need:.1f}pt")
    return probs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ref", default=G.DEFAULT_REF)
    ap.add_argument("--outdir", default=OUT_DIR)
    a = ap.parse_args()
    ok = True
    for fn, name in ((build_A, "DP1-architecture-view-A-controlloop-3plane.pptx"), (build_B, "DP1-architecture-view-B-plugin-ports.pptx")):
        out = os.path.join(a.outdir, name)
        prs, diag = fn(a.ref, out)
        probs = X.verify(prs, diag) + check_cells()
        print("saved:", out)
        for p in probs:
            ok = False
            print(" -", p)
    if not ok:
        sys.exit(1)
    print("verify OK: bounds / text fit / node overlap / arrow edges / table cells")


if __name__ == "__main__":
    main()
