#!/usr/bin/env python3
"""Build the C1 QA1 complement-structure slides (block diagram + component table).

    python3 doc-mk/Evaluation/tools/gen_c1_complement_pptx.py [--out PATH] [--preview-dir DIR]

Standalone (python-pptx + Pillow only; does not import gen_dp1_result / gen_dp_pptx, which other agents edit).
Numbers are copied from doc-mk/DP1/dp1-c1-qa1-complement-design.md (diagnosis from results/data/*/qa_result.json);
all effect statements on the proposed components are Evidence [C] (reasoning, unimplemented).
After saving, the script verifies: every shape inside the 16:9 slide, text fits its box (PIL measurement with the
CJK font), and no two boxes overlap. Exit code 1 if a check fails.
"""
from __future__ import annotations

import argparse
import math
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
DEFAULT_OUT = HERE.parent.parent / "DP1" / "DP1-c1-qa1-complement-structure.pptx"
FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
FONT_NAME = "맑은 고딕"
SW, SH = 13.333, 7.5

# (fill, line, text, line width pt, dash)
STY = {
    "exist": ("DCE9F5", "4D7EA8", "1A1A1A", 1.0, False),
    "chg": ("E3F1E3", "3F8A3F", "1A1A1A", 2.25, False),
    "new": ("FDEBD3", "C0701A", "1A1A1A", 2.25, False),
    "head": ("1F4E79", "1F4E79", "FFFFFF", 0.75, False),
    "cell": ("FFFFFF", "BFBFBF", "1A1A1A", 0.75, False),
    "note": ("F2F2F2", "BFBFBF", "404040", 0.75, False),
    "warn": ("FFF8EC", "C98A3C", "1A1A1A", 0.75, False),
    "label": (None, None, "404040", 0, False),
}

LINE_H = 1.22  # line height factor used for both layout and fit check


def _font(sz):
    from PIL import ImageFont
    return ImageFont.truetype(FONT_PATH, max(8, int(sz * 110 / 72)))


_MEAS = None


def text_w(s, size):
    global _MEAS
    from PIL import Image, ImageDraw
    if _MEAS is None:
        _MEAS = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    return _MEAS.textlength(s, font=_font(size)) / 110.0  # inches


def wrap(text, width_in, size, margin=0.12):
    """Greedy char wrap to the box width (Korean has no spaces worth breaking on; 6% safety)."""
    maxw = (width_in - margin) * 0.94
    out, cur = [], ""
    for ch in text:
        if cur and text_w(cur + ch, size) > maxw:
            out.append(cur)
            cur = ch.lstrip()
        else:
            cur += ch
    if cur:
        out.append(cur)
    return out or [""]


class Box:
    def __init__(self, key, x, y, w, h, paras, sty, align="l", anchor="top"):
        """paras: list of (text, size, bold). Pre-wrapped into lines at build time."""
        self.key, self.x, self.y, self.w, self.h = key, x, y, w, h
        self.sty, self.align, self.anchor = sty, align, anchor
        self.lines = []  # (text, size, bold)
        for t, size, bold in paras:
            for ln in wrap(t, w, size):
                self.lines.append((ln, size, bold))

    def need_h(self):
        return sum(s * LINE_H / 72.0 for _, s, _ in self.lines) + 0.08

    def cx(self):
        return self.x + self.w / 2

    def cy(self):
        return self.y + self.h / 2


def edge_point(b, tx, ty):
    """Point where the ray from b's center toward (tx, ty) leaves b's rectangle."""
    cx, cy = b.cx(), b.cy()
    dx, dy = tx - cx, ty - cy
    if dx == 0 and dy == 0:
        return cx, cy
    sx = (b.w / 2) / abs(dx) if dx else math.inf
    sy = (b.h / 2) / abs(dy) if dy else math.inf
    s = min(sx, sy)
    return cx + dx * s, cy + dy * s


class Slide:
    def __init__(self, title):
        self.title = title
        self.boxes: list[Box] = []
        self.arrows = []  # (from_key, to_key, dashed)
        self.by = {}

    def add(self, b):
        self.boxes.append(b)
        self.by[b.key] = b
        return b

    def arrow(self, a, b, dashed=False):
        self.arrows.append((a, b, dashed))


# --------------------------------------------------------------------------------------------------------
def slide_diagram():
    s = Slide("C1 QA1 보완 구조 - 기존 C1 pipeline에 Activity Gate와 Paced Migration을 추가 ([C] 설계 제안, 미구현)")
    P = lambda t, sz=8.5, b=False: (t, sz, b)
    s.add(Box("banner", 0.4, 0.98, 12.5, 0.4, [P("진단 (qa_result.json, 통합 21쌍): QA1 격차 C1 x1.298 대 C2 x1.422의 ln 격차 96.5%가 B200 같은 class KV 3쌍에 있다. C1은 이 3쌍 15 seed-run 중 10개가 Baseline과 동일(이동 없음).", 9, False)], "note", anchor="mid"))
    cols = [0.4, 3.55, 6.7, 9.85]
    W, H = 2.75, 0.66
    rows = [1.55, 2.5, 3.45, 4.4, 5.35]

    def at(c, r):
        return cols[c], rows[r]

    def mk(key, c, r, title, detail, sty):
        x, y = at(c, r)
        s.add(Box(key, x, y, W, H, [P(title, 9.5, True), P(detail, 8, False)], sty, align="ctr", anchor="mid"))

    mk("EV", 0, 0, "Event Source", "TELEMETRY / ALLOCATED / FREED / ACCESSED", "exist")
    mk("RSM", 1, 0, "Resource State Monitor", "capacity / BW / load -> pressure", "exist")
    mk("RTA", 2, 0, "Resource Trend Analyzer", "pressure source tier 선정", "exist")
    mk("B2", 3, 0, "B2 No-benefit Rebalance Gate", "[신규] 용량 부족도 SLO 위반도 없으면 rebalance 보류", "new")
    mk("A1", 0, 1, "A1 Activity Tag Store", "[신규] last_access_ts + 8 s window 횟수 (Registry 밖)", "new")
    mk("A3", 3, 1, "A3 Data Eviction Manager", "[변경] victim을 idle 우선 정렬", "chg")
    mk("DOR", 0, 3, "Data Object Registry", "위치 / 크기 / tier, type-agnostic (변경 없음)", "exist")
    mk("A2", 2, 2, "A2 Promotion / Swap Pass", "[변경] gain x act(p) >= margin x loss x act(v)", "chg")
    mk("DMA", 3, 2, "Data-Memory Affinity Mapper", "static hint: op class, shape, sensitivity", "exist")
    mk("ACE", 2, 3, "Access Cost Estimator", "descriptor 기반 비용 (B2, A2도 사용)", "exist")
    mk("DTS", 3, 3, "Destination Tier Selector", "SLO filter, do-no-harm, affinity score", "exist")
    mk("MB", 3, 4, "Migration Budget", "[변경] link-time token bucket", "chg")
    mk("B1", 2, 4, "B1 Staged Migration Pacer", "[신규] tick당 share만 진행, copy-then-switch", "new")
    mk("ME", 1, 4, "Migration Executor", "common boundary (C1/C2 공유)", "exist")

    for a, b in (("EV", "RSM"), ("RSM", "RTA"), ("RTA", "B2"), ("B2", "A3"), ("A3", "DMA"), ("DMA", "DTS"),
                 ("DTS", "MB"), ("MB", "B1"), ("B1", "ME"), ("EV", "A1"), ("A1", "A3"), ("A1", "A2"),
                 ("ACE", "DTS"), ("DOR", "A2"), ("A2", "DTS")):
        s.arrow(a, b)

    def label(key, x, y, w, text):
        s.add(Box(key, x, y, w, 0.24, [P(text, 7.5, False)], "label", anchor="mid"))

    label("L1", 1.85, 2.23, 1.7, "ACCESSED (현재 C1은 버림)")
    label("L2", 4.2, 2.56, 1.6, "idle-first (A3)")
    label("L3", 4.1, 3.12, 1.5, "activity (A2)")
    # legend + fence
    ly = 6.12
    s.add(Box("lg1", 0.4, ly, 1.45, 0.3, [P("기존 블록", 8.5, True)], "exist", align="ctr", anchor="mid"))
    s.add(Box("lg2", 1.95, ly, 1.45, 0.3, [P("변경 블록", 8.5, True)], "chg", align="ctr", anchor="mid"))
    s.add(Box("lg3", 3.5, ly, 1.45, 0.3, [P("신규 블록", 8.5, True)], "new", align="ctr", anchor="mid"))
    s.add(Box("fence", 0.4, 6.5, 6.55, 0.78, [
        P("정체성 fence (C2와의 경계)", 8.5, True),
        P("객체당 상태는 last_access_ts + window 횟수뿐(EWMA, 이력, 예측, class prior, type 없음). activity는 새 trigger가 아니라 기존 trigger의 후보 축소/가중/정렬에만 쓴다. 새 정책 상수 없음(window=COOLDOWN_S, margin=AFFINITY_MARGIN, share=LINK_SHARE).", 8, False)], "warn"))
    s.add(Box("keep", 7.1, 6.5, 5.8, 0.78, [
        P("보존해야 할 C1 이점 (측정, 통합 21쌍)", 8.5, True),
        P("이동 138 대 1,339 GiB (C2), 링크 점유 1.6% 대 10.5%, 결정 연산 3 대 111 ms/run, HBM x0.97 대 x1.21, QA4 module 1.75 대 2.50. 보완 후 같은 항목을 재측정해 guard로 쓴다.", 8, False)], "note"))
    s.add(Box("srcnote", 5.1, ly, 7.8, 0.3, [P("A1/A2/A3 = QA1 격차(D1~D3), B1/B2 = 꼬리와 큰 객체 admission(D4, D5). A1 신호: C1이 관측(A1-observe) 또는 runtime hint push(A1-hint)", 7.5, False)], "label", anchor="mid"))
    return s


def slide_table():
    s = Slide("C1 QA1 보완 컴포넌트 표 - 겨냥 진단, 기대 방향, 검증 상태 (전부 [C] 논증, 수치 이득 주장 없음)")
    P = lambda t, sz=8.5, b=False: (t, sz, b)
    cols = [("ID / 컴포넌트", 0.4, 2.35), ("겨냥 진단 (측정 [B+C], 출처 qa_result.json)", 2.75, 3.15),
            ("기대 방향 (↑ 개선 ↓ 악화 ≈ 중립 ? 불확실)", 5.9, 2.2), ("리스크 / trade-off", 8.1, 2.65), ("검증 상태 / 실험", 10.75, 2.15)]
    for name, x, w in cols:
        s.add(Box("h_" + name, x, 1.0, w, 0.42, [P(name, 9, True)], "head", align="ctr", anchor="mid"))
    rows = [
        ("A1 Activity Tag Store", "신규 side-car: ACCESSED -> last_access_ts + 8 s window 횟수",
         "D1: 같은 class KV 3쌍(B200)이 QA1 ln 격차의 96.5% (recency_shift 44.4, idle_kv 33.6, rotating 18.5%). C1 x1.11/1.48/1.19 대 C2 x2.62/2.82/1.70, C1 이동 0.8~1.6회/run 대 C2 21~30회",
         "QA1 ↑ · QA2 ↑? (TTFT P50) · QA3 ≈ · QA4 ≈ (C1 로컬 1 module 예상)",
         "C2화 drift. fence 필수. 낮은 접근률 noise (hot 2.8회/window면 0회 확률 약 6%)",
         "[C] 미구현. Exp-1: DP1_C1_ACTIVITY 변형, off == 본 결과여야 함"),
        ("A2 Promotion / Swap (변경)", "gain x (n_p+1) >= margin x loss x (n_v+1), 활성 객체만 승격 후보",
         "D2: 같은 class는 static gain == loss라 margin 2.0 통과 불가. D3: H100 1 TiB RAG C1 x1.06 대 C2 x1.41 (이동 1회 대 18회, CI 겹침)",
         "QA1 ↑ · QA2 ↑? · QA3 ≈/↓ (직접 승격은 HBM 증가, C2는 HBM x1.21) · QA4 ≈",
         "새 trigger가 되면 C2화. rotating hotset thrash (기존 cooldown 8 s, margin 2.0이 방어선)",
         "[C] Exp-1 + fast-rotation 대조 시나리오, QA3 HBM 비율 <= x1.05 guard"),
        ("A3 Eviction (변경)", "victim 정렬: idle 우선 (현재 key = size x residency age)",
         "D2: 현재 victim 순서가 접근 시각과 무관 (설계 8.8의 basic age/LRU와 구현 차이)",
         "QA1 ≈ · QA3 ↑ (cold가 HBM을 먼저 떠남) · QA2 ≈",
         "작은 객체를 여러 개 고르면 이동 횟수 증가 가능",
         "[C] Exp-1 variant (이동 횟수/GiB 모니터링)"),
        ("B2 No-benefit Rebalance Gate", "신규: 용량 부족도 SLO 위반도 없는 압박이면 rebalance 보류",
         "D4: Common 4쌍에서 C1 이동 231~284 GiB, QA1 x0.99~1.00, TTFT P99 x3.76~4.32 (C1 악화 6쌍, C2 5쌍)",
         "QA1 ≈ · QA2 ↑ (꼬리) · QA3 ≈ · QA4 ≈",
         "host path 경합의 C1 승리(H100 x1.63, B200 x1.37)를 막을 위험. 대역폭 충격을 비용 추정에 반영해야 함",
         "[C] Exp-2. 반증: dyn_host_path win 소실"),
        ("B1 Staged Migration Pacer", "신규: tick당 min(LINK_SHARE, 1 - bw_util)만 진행, copy-then-switch",
         "D5: 이동이 burst cap보다 크면 admit 불가. cap 1.0 s에서 C1 QA1 x1.206, 0.5 s에서 x1.000 (2.0 s x1.298)",
         "QA1 ↑? (큰 객체 통과) · QA2 ↑ (꼬리) · QA3 ≈ · QA4 ↓? (공유 경계 변경)",
         "simulator executor에 in-flight 모델 필요 (M-class, C2에도 영향). 평균장 간섭 과대 가능성과 혼재",
         "[C] Exp-3, 별도 iteration. 수정 전 결과는 pre_ 폴더에 보존"),
    ]
    y = 1.46
    for rid, sub, diag, qa, risk, val in rows:
        h = 0.9
        vals = [([P(rid, 9, True), P(sub, 8, False)], "new" if rid.startswith(("A1", "B2", "B1")) else "chg"),
                ([P(diag, 8, False)], "cell"), ([P(qa, 8.5, False)], "cell"), ([P(risk, 8, False)], "cell"), ([P(val, 8, False)], "cell")]
        for (name, x, w), (paras, sty) in zip(cols, vals):
            s.add(Box(f"r{rid[:2]}_{x}", x, y, w, h, paras, sty))
        y += h + 0.04
    s.add(Box("limits", 0.4, y + 0.04, 12.5, 1.1, [
        P("한계 (솔직하게)", 9, True),
        P("격차의 96.5%를 차지하는 3쌍은 B200 한 시스템에만 있고(H100은 Baseline이 SLO 불가) hot/cold 대비 10~30배와 객체 8~9개를 시나리오가 심었다. 이득은 상한에 가깝다. A1은 per-object 관측이라 C1의 '객체별 behavior 비추적' 원칙을 부분 양보하는 것이며, 양보 없이 닫히는 부분은 B(QA2)뿐이다. "
          "회계식(예측 아님): 세 쌍의 격차를 비율 f만큼 닫으면 통합 QA1 = 1.298 x exp(0.088 f) (f=0.5이면 x1.357). C2의 이득 일부는 HBM 증가(x1.21)와 이동 증가로 얻은 것이라 같은 일을 하면 QA3와 이동량 이점이 줄어든다. QA1 별 경계 1.30은 결과를 본 뒤 정한 값이라 별 변화는 성과가 아니다.", 8.5, False)], "warn"))
    return s


# --------------------------------------------------------------------------------------------------------
def build(slides, out: Path):
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    rgb = RGBColor.from_string
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]

    def style_runs(p, text, size, bold, color):
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bool(bold)
        r.font.name = FONT_NAME
        r.font.color.rgb = rgb(color)
        rpr = r._r.get_or_add_rPr()
        ea = etree.SubElement(rpr, qn("a:ea"))
        ea.set("typeface", FONT_NAME)

    for sl in slides:
        s = prs.slides.add_slide(blank)
        # title band
        band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.82))
        band.fill.solid(); band.fill.fore_color.rgb = rgb("1F4E79"); band.line.fill.background()
        band.shadow.inherit = False
        tf = band.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.4); tf.margin_right = Inches(0.3)
        for i, ln in enumerate(wrap(sl.title, SW - 0.7, 15, 0.1)):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            style_runs(p, ln, 15, True, "FFFFFF")
        for b in sl.boxes:
            fill, line, tc, lw, _ = STY[b.sty]
            shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(b.x), Inches(b.y), Inches(b.w), Inches(b.h))
            shp.shadow.inherit = False
            if fill is None:
                shp.fill.background(); shp.line.fill.background()
            else:
                shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
                shp.line.color.rgb = rgb(line); shp.line.width = Pt(lw)
            tf = shp.text_frame; tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE if b.anchor == "mid" else MSO_ANCHOR.TOP
            for i, (ln, size, bold) in enumerate(b.lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER if b.align == "ctr" else PP_ALIGN.LEFT
                p.line_spacing = 1.0
                style_runs(p, ln, size, bold, tc)
        for a, bkey, dashed in sl.arrows:
            A, B = sl.by[a], sl.by[bkey]
            x1, y1 = edge_point(A, B.cx(), B.cy())
            x2, y2 = edge_point(B, A.cx(), A.cy())
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
            c.line.color.rgb = rgb("404040"); c.line.width = Pt(1.25)
            ln = c.line._get_or_add_ln()
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    prs.save(str(out))


def verify(path: Path, slides) -> list[str]:
    """Re-open the saved pptx and check bounds, text fit, and box overlaps from our layout model + file geometry."""
    from pptx import Presentation
    from pptx.util import Emu
    errs = []
    prs = Presentation(str(path))
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    if abs(W / H - 16 / 9) > 0.01:
        errs.append(f"slide aspect {W:.3f}x{H:.3f} is not 16:9")
    if len(prs.slides) != len(slides):
        errs.append("slide count mismatch")
    for si, (ps, sl) in enumerate(zip(prs.slides, slides), 1):
        for shp in ps.shapes:
            x, y, w, h = (Emu(v).inches for v in (shp.left, shp.top, shp.width, shp.height))
            if x < -1e-6 or y < -1e-6 or x + w > W + 1e-6 or y + h > H + 1e-6:
                errs.append(f"slide {si}: shape out of bounds ({shp.shape_type}, {x:.2f},{y:.2f},{w:.2f},{h:.2f})")
        for b in sl.boxes:
            if b.need_h() > b.h + 1e-6:
                errs.append(f"slide {si}: text overflow in '{b.key}' need {b.need_h():.2f} in > box {b.h:.2f} in")
            for ln, size, _ in b.lines:
                if text_w(ln, size) > b.w - 0.1 + 1e-6:
                    errs.append(f"slide {si}: line wider than box in '{b.key}': {ln[:20]}")
        bs = [b for b in sl.boxes if b.sty != "label"]
        for i, a in enumerate(bs):
            for c in bs[i + 1:]:
                if a.x < c.x + c.w - 1e-6 and c.x < a.x + a.w - 1e-6 and a.y < c.y + c.h - 1e-6 and c.y < a.y + a.h - 1e-6:
                    errs.append(f"slide {si}: boxes overlap '{a.key}' / '{c.key}'")
        for lb in [b for b in sl.boxes if b.sty == "label"]:
            for c in bs:
                if lb.x < c.x + c.w and c.x < lb.x + lb.w and lb.y < c.y + c.h and c.y < lb.y + lb.h:
                    errs.append(f"slide {si}: label '{lb.key}' overlaps box '{c.key}'")
        # title band fit (15 pt, 0.82 in band)
        n = len(wrap(sl.title, SW - 0.7, 15, 0.1))
        if n * 15 * LINE_H / 72 > 0.82:
            errs.append(f"slide {si}: title needs {n} lines, exceeds band")
        # arrows must not run through unrelated boxes
        for a, bk, _ in sl.arrows:
            A, B = sl.by[a], sl.by[bk]
            x1, y1 = edge_point(A, B.cx(), B.cy())
            x2, y2 = edge_point(B, A.cx(), A.cy())
            for c in bs:
                if c.key in (a, bk):
                    continue
                for t in [i / 40 for i in range(1, 40)]:
                    px, py = x1 + (x2 - x1) * t, y1 + (y2 - y1) * t
                    if c.x + 0.02 < px < c.x + c.w - 0.02 and c.y + 0.02 < py < c.y + c.h - 0.02:
                        errs.append(f"slide {si}: arrow {a}->{bk} crosses box '{c.key}'")
                        break
    return errs


def preview(sl, out: Path):
    from PIL import Image, ImageDraw
    S = 110
    hx = lambda h: tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
    im = Image.new("RGB", (int(SW * S), int(SH * S)), "white")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, int(SW * S), int(0.82 * S)], fill=hx("1F4E79"))
    for i, ln in enumerate(wrap(sl.title, SW - 0.7, 15, 0.1)):
        d.text((int(0.4 * S), 8 + i * 22), ln, font=_font(15), fill=(255, 255, 255))
    for b in sl.boxes:
        fill, line, tc, lw, _ = STY[b.sty]
        x, y, w, h = [int(v * S) for v in (b.x, b.y, b.w, b.h)]
        if fill:
            d.rectangle([x, y, x + w, y + h], fill=hx(fill), outline=hx(line), width=max(1, int(lw)))
        lh = [s * LINE_H / 72 * S for _, s, _ in b.lines]
        ty = y + (h - sum(lh)) / 2 if b.anchor == "mid" else y + 4
        for (ln, size, bold), l_h in zip(b.lines, lh):
            f = _font(size)
            tw = d.textlength(ln, font=f)
            d.text((x + (w - tw) / 2 if b.align == "ctr" else x + 5, ty), ln, font=f, fill=hx(tc))
            ty += l_h
    for a, bk, _ in sl.arrows:
        A, B = sl.by[a], sl.by[bk]
        x1, y1 = edge_point(A, B.cx(), B.cy())
        x2, y2 = edge_point(B, A.cx(), A.cy())
        d.line([x1 * S, y1 * S, x2 * S, y2 * S], fill=(64, 64, 64), width=2)
        ang = math.atan2(y2 - y1, x2 - x1)
        for da in (2.7, -2.7):
            d.line([x2 * S, y2 * S, (x2 + 0.11 * math.cos(ang + da)) * S, (y2 + 0.11 * math.sin(ang + da)) * S], fill=(64, 64, 64), width=2)
    im.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--preview-dir", type=Path, default=None)
    a = ap.parse_args()
    slides = [slide_diagram(), slide_table()]
    a.out.parent.mkdir(parents=True, exist_ok=True)
    build(slides, a.out)
    print("wrote", a.out, len(slides), "slide(s)")
    errs = verify(a.out, slides)
    for e in errs:
        print("CHECK FAIL:", e)
    if a.preview_dir:
        a.preview_dir.mkdir(parents=True, exist_ok=True)
        for i, sl in enumerate(slides, 1):
            preview(sl, a.preview_dir / f"c1-complement-{i}.png")
    print("verify:", "OK" if not errs else f"{len(errs)} problem(s)")
    sys.exit(1 if errs else 0)


if __name__ == "__main__":
    main()
