"""DP PPT 제작 예시 스크립트 (doc-mk/slides/README.md 참고).

reference/DP-PPT-style-reference.pptx 를 템플릿으로 열어 기존 슬라이드를 지우고, 레이아웃 "1_제목 및 내용"에
내비게이션/띠/칩/점선 패널/표 헬퍼로 슬라이드를 만든다. 내용은 DP2 후보 구조(A / B′ / C) 예시이며,
다른 DP를 만들 때는 아래 슬라이드 정의부(# ==== 1 배경 이후)만 바꿔 쓴다.

실행: python build_dp_deck_example.py <reference.pptx> <out.pptx>   (python-pptx 필요, .venv 사용)
"""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

REF, OUT = sys.argv[1], sys.argv[2]
FONT = "맑은 고딕"
C = lambda h: RGBColor.from_string(h)
BLACK, WHITE = C("000000"), C("FFFFFF")
G40, G59, G7F, GBF, GD9, GF2, GF0 = C("404040"), C("595959"), C("7F7F7F"), C("BFBFBF"), C("D9D9D9"), C("F2F2F2"), C("F0F1F3")
BLUE, LBLUE = C("2E75B6"), C("DEEBF7")
ORG, LORG = C("C55A11"), C("FBE5D6")
GRN, LGRN = C("548235"), C("E2F0D9")
KEY = C("0066FF")
RED = C("FF0000")
CF, CQ, CC = C("203864"), C("833C0B"), C("375623")   # FR / QA / 제약 chips
TOTAL = 11

prs = Presentation(REF)
# ---- drop the reference slides (keep master/layout/theme)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
LAY = [l for l in prs.slide_layouts if l.name == "1_제목 및 내용"][0]
for sh in list(LAY.shapes):
    if sh.shape_type == 6:          # nav group -> drawn per slide
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "/ 23":
        r = sh.text_frame.paragraphs[0].runs[0]
        r.text = f"/ {TOTAL}"
        for extra in sh.text_frame.paragraphs[0].runs[1:]:
            extra.text = ""
NAV = ["과제 소개", "요구사항", "설계", "구현/검증", "결론"]
NX = [8.227, 9.203, 10.186, 11.162, 12.145]
NW = [1.128, 1.128, 1.128, 1.128, 0.976]


def set_font(run, size, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=0.8, dash=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
        if dash:
            ln = s.line._get_or_add_ln()
            d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    s.shadow.inherit = False
    return s


def write(s, paras, size=10, color=BLACK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.05, space=2):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for p in paras:
        t, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = o.get("align", align)
        para.space_after = Pt(o.get("space", space))
        for st, so in (o.get("segs") or [(t, {})]):
            r = para.add_run(); r.text = st
            set_font(r, o.get("size", size), so.get("bold", o.get("bold", bold)), so.get("color", o.get("color", color)))
    return s


def tbox(slide, x, y, w, h, paras, **kw):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return write(s, paras, **kw)


def rect(slide, x, y, w, h, paras, fill=WHITE, line=G59, lw=0.8, size=9, bold=False, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, color=BLACK, dash=False, kind=MSO_SHAPE.RECTANGLE, **kw):
    s = shape(slide, kind, x, y, w, h, fill, line, lw, dash)
    return write(s, paras, size=size, bold=bold, align=align, anchor=anchor, color=color, **kw)


def arrow(slide, x1, y1, x2, y2, color=G7F, w=1.0, dash=False, both=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    if both:
        h = etree.SubElement(ln, qn("a:headEnd")); h.set("type", "triangle")
    t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle")
    return c


def new_slide(title, nav):
    s = prs.slides.add_slide(LAY)
    s.shapes.title.text = title
    for i, (n, x, w) in enumerate(zip(NAV, NX, NW)):
        cur = i == nav
        sh = shape(s, MSO_SHAPE.CHEVRON, x, 0.15, w, 0.391, G59 if cur else GBF)
        write(sh, [n], size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    shape(s, MSO_SHAPE.RECTANGLE, 8.234, 0.15, 0.232, 0.391, G59 if nav == 0 else GBF)
    shape(s, MSO_SHAPE.RECTANGLE, 12.893, 0.15, 0.232, 0.391, G59 if nav == 4 else GBF)
    return s


def pageno(s, n):
    tbox(s, 5.85, 7.2, 0.68, 0.26, [(str(n), {"align": PP_ALIGN.RIGHT})], size=10, color=G7F)


def band(s, text, chips=()):
    shape(s, MSO_SHAPE.RECTANGLE, 0.26, 0.65, 12.85, 0.43, G7F)
    shape(s, MSO_SHAPE.RECTANGLE, 0.44, 0.74, 0.07, 0.24, WHITE)
    tbox(s, 0.51, 0.68, 8.3, 0.36, [text], size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    x = 13.01
    for name, col in reversed(chips):
        w = 0.78
        x -= w + 0.04
        b = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.69, w, 0.35, col)
        write(b, [name], size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)


def dashed(s, x, y, w, h, label, lw=2.6):
    shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, None, GBF, 1.0, dash=True)
    r = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + (w - lw) / 2, y - 0.17, lw, 0.34, GF2, G7F, 0.8)
    write(r, [label], size=12, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)


def container(s, x, y, w, h, title, fill, line, tsize=9):
    shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, 1.0)
    r = shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.24, line, line)
    write(r, [title], size=tsize, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)


def keymsg(s, y, text, issues):
    ar = shape(s, MSO_SHAPE.RIGHT_ARROW, 0.31, y, 0.45, 0.3, BLUE, None)
    tbox(s, 0.8, y - 0.02, 12.2, 0.34, [text], size=12, bold=True, color=KEY, anchor=MSO_ANCHOR.MIDDLE)
    tbox(s, 0.42, y + 0.34, 12.5, 0.6, issues, size=11, bold=True, space=1)


def table(s, x, y, colw, rows, heights, hdr=True, size=10.5, fills=None, first_color=G7F, align_first=PP_ALIGN.CENTER):
    nr, nc = len(rows), len(colw)
    gs = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)), Inches(sum(heights)))
    tb = gs.table
    tbl = gs._element.graphic.graphicData.tbl
    sid = tbl.tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # no style, no grid
    tbl.tblPr.set("bandRow", "0"); tbl.tblPr.set("firstRow", "0")
    for i, w in enumerate(colw):
        tb.columns[i].width = Inches(w)
    for r in range(nr):
        tb.rows[r].height = Inches(heights[r])
        for c in range(nc):
            cell = tb.cell(r, c)
            val = rows[r][c]
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            items = val if isinstance(val, list) else [val]
            for k, it in enumerate(items):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                text, o = (it, {}) if isinstance(it, str) else it
                para.space_after = Pt(1)
                head = hdr and r == 0
                if c == 0 or head:
                    para.alignment = o.get("align", PP_ALIGN.CENTER if (head or align_first == PP_ALIGN.CENTER) else PP_ALIGN.LEFT)
                for st, so in (o.get("segs") or [(text, {})]):
                    run = para.add_run(); run.text = st
                    col = first_color if (c == 0 or head) else so.get("color", o.get("color", BLACK))
                    set_font(run, o.get("size", size + (1 if head else 0)), so.get("bold", o.get("bold", head and c > 0)), col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (fills or {}).get((r, c), WHITE)
            tcPr = cell._tc.get_or_add_tcPr()
            fill_el = tcPr.find(qn("a:solidFill"))
            if fill_el is not None:
                tcPr.remove(fill_el)
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(tag)); ln.set("w", "6350")
                sf = etree.SubElement(ln, qn("a:solidFill")); cl = etree.SubElement(sf, qn("a:srgbClr")); cl.set("val", "BFBFBF")
            if fill_el is not None:
                tcPr.append(fill_el)
    return gs


def bullets(items, size=10.5):
    return [("• " + t if isinstance(t, str) else ("• " + t[0], t[1])) for t in items]


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ================================================================ 1 배경
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 배경", 0)
tbox(s, 0.39, 0.73, 12.7, 1.0, [
    "- 이기종 메모리 환경에서는 KV가 여러 Tier(HBM, CXL-PNM, ScHBM, HBF 등)에 분산 배치되고, 일부 Tier는 attention을 직접 실행할 수 있음",
    "- KV를 HBM으로 끌어와 GPU에서 계산하면 단일 요청은 빠를 수 있으나, 동시성이 높으면 HBM 용량·대역폭·링크를 소모해 처리량이 떨어질 수 있음",
    "- 따라서 Turn 시작 시 KV 구간별로 attention을 어디서 실행할지(GPU_STAGE / IN_SITU)를 정하는 결정 구조가 필요함"], size=12, color=G40, space=3)
dashed(s, 0.38, 2.0, 6.05, 3.95, "DP2의 위치 (노드 내 결정)")
for i, (t, d, f, ln_) in enumerate([
        ("DP1  메모리 배치", "데이터를 어느 Tier에 둘지 (남는 배치: promote / evict / prefetch)", GF0, G7F),
        ("DP2  실행 위치 (본 설계)", "그 배치를 입력으로 attention을 어디서 실행할지 (노드 내, Turn 시작 시)", LBLUE, BLUE),
        ("DP4  멀티 노드 확장", "노드 간으로 확장해 P 노드 / D 노드를 분별", GF0, G7F)]):
    y = 2.35 + i * 1.12
    rect(s, 0.75, y, 5.3, 0.82, [(t, {"bold": True, "size": 12}), (d, {"size": 10})], fill=f, line=ln_, lw=1.5 if ln_ == BLUE else 0.8)
    if i < 2:
        arrow(s, 3.4, y + 0.82, 3.4, y + 1.12, G59, 1.5)
tbox(s, 0.6, 5.5, 5.6, 0.4, ["DP1은 DP1대로, DP2는 DP2대로 평가. 경계: DP1 = 이동 후 남는 데이터, DP2 = 실행을 위한 일시적 staging"], size=9, color=G7F)
dashed(s, 6.75, 2.0, 6.05, 3.95, "결정 값: GPU_STAGE vs IN_SITU", 3.6)
rect(s, 6.95, 2.4, 5.65, 0.26, ["GPU_STAGE : KV를 HBM staging 영역으로 옮겨 GPU에서 attention"], fill=LORG, line=ORG, size=10, bold=True)
for k, (t, f, ln_) in enumerate([("Tier (KV)", LGRN, GRN), ("HBM staging", LBLUE, BLUE), ("GPU attention", LBLUE, BLUE)]):
    rect(s, 7.1 + k * 1.95, 2.8, 1.45, 0.45, [t], fill=f, line=ln_, size=10)
    if k < 2:
        arrow(s, 8.55 + k * 1.95, 3.02, 9.05 + k * 1.95, 3.02, ORG, 1.5)
tbox(s, 8.35, 3.27, 1.2, 0.2, ["PCIe 이동"], size=8, color=RED)
rect(s, 6.95, 3.65, 5.65, 0.26, ["IN_SITU(tier) : KV가 있는 Tier에서 q를 받아 fused attention"], fill=LGRN, line=GRN, size=10, bold=True)
rect(s, 7.1, 4.1, 1.7, 0.45, ["GPU (QKV 투영)"], fill=LBLUE, line=BLUE, size=10)
rect(s, 9.45, 4.1, 1.7, 0.45, ["Tier fused attention"], fill=LGRN, line=GRN, size=10)
rect(s, 11.35, 4.1, 1.15, 0.45, ["LSE 병합"], fill=LBLUE, line=BLUE, size=10)
arrow(s, 8.8, 4.2, 9.45, 4.2, GRN, 1.5); arrow(s, 11.15, 4.2, 11.35, 4.2, GRN, 1.5)
tbox(s, 8.85, 3.93, 0.6, 0.2, ["q"], size=8, color=RED); tbox(s, 11.1, 3.93, 0.5, 0.2, ["(m,l,o)"], size=8, color=RED)
tbox(s, 6.95, 4.75, 5.65, 1.1, [
    ("끌어오기(GPU_STAGE)의 비용", {"bold": True, "size": 10}),
    "HBM 용량 점유(running batch 감소) / HBM write BW / 공유 링크의 직렬 상한",
    "128K 컨텍스트 KV ≈ 43GB → PCIe 64 GB/s에서 ≈ 0.67 s [추정]"], size=10, color=G40, space=2)
keymsg(s, 6.12, "Turn 시작 시 KV 구간별 attention 실행 위치를 HBM 기회비용까지 고려해 결정하는 구조 설계 필요",
       ["설계 쟁점 1. 처리량(QA1)·꼬리 지연(QA2)·HBM 점유(QA3)를 함께 만족하는 attention 실행 위치 결정",
        "설계 쟁점 2. 결정 주체·정보 범위에 따른 구조 선택(A / B′)과 약점을 보완하는 하이브리드 C"])
pageno(s, 1)
notes(s, "핵심: DP2는 P/D를 구분하지 않고 Turn 시작 시 KV 구간별 attention을 어디서 실행할지(GPU_STAGE / IN_SITU)만 정한다.\n"
         "근거: 끌어오기는 단일 요청에는 빠르나 동시성이 높으면 HBM·링크를 소모한다. 수치(43GB, 0.67 s)는 Llama-3.1-70B, PCIe 64 GB/s 가정의 추정이다.\n"
         "연결: DP1(배치)을 입력으로 받고, DP4가 멀티 노드로 확장한다.")

# ================================================================ 2 요구사항
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 요구사항", 1)
band(s, "DP2가 판별해야 하는 QA와 제약", [("Q1", CQ), ("Q2", CQ), ("Q3", CQ), ("C1", CC), ("C2", CC)])
table(s, 0.27, 1.25, [0.8, 1.9, 4.6, 5.55], [
    ["ID", "QA", "정의 (지표)", "이 DP에서 보는 이유"],
    ["Q1", "Throughput", "Max SLO Goodput (지연은 SLO 조건으로 포함)", "끌어오기(GPU_STAGE)는 고동시성에서 HBM·링크를 소모해 처리량을 깎을 수 있음"],
    ["Q2", "Latency", "TTFT · TPOT의 P99 / P50. SLO와 달리 꼬리 분포", "배치 오류는 평균이 아니라 꼬리로 나타남 (예: 느린 Tier 배정 → TPOT 꼬리)"],
    ["Q3", "Resource Utilization", "HBM 점유율, 낮을수록 좋음 (iso-load, SLO 통과 조건, weight 제외)", "같은 성능이면 HBM을 덜 쓰는 편이 좋음. DP1의 Q3와 정의 통일"]],
    [0.38, 0.62, 0.62, 0.62], size=10.5,
    fills={(1, 0): GF2, (2, 0): GF2, (3, 0): GF2, (0, 0): GF2, (0, 1): GF2, (0, 2): GF2, (0, 3): GF2})
dashed(s, 0.38, 4.1, 6.1, 2.85, "제약 (C)", 1.6)
tbox(s, 0.5, 4.35, 5.9, 2.6, bullets([
    "C-A1  연산은 항상 GPU에서 시작 (attention만 이동 가능)",
    "C-A2  iteration 단위 결정, layer 스트리밍은 배제 (현실성 없음)",
    "C-A3  연산 가능 메모리의 지원 연산은 attention 계열뿐 (GC-5)",
    "C-A4  위치와 무관하게 모델 출력 동일 (bit-exact)",
    "C-A5  시뮬레이션 [B+C], 실측 아님 (GC-1)",
    "공통 가정  Tier별 부분 결과 (m, l, o)를 log-sum-exp로 병합. 부분 결과 kernel은 주어진 것 (GC-5, 확인 필요)"]), size=10.5, space=4)
dashed(s, 6.75, 4.1, 6.1, 2.85, "제외한 QA와 이유", 2.2)
tbox(s, 6.87, 4.35, 5.9, 2.6, bullets([
    "변경 용이성: 신규 Tier는 Tier descriptor 추가로 처리 → 후보 간 변별력 낮음",
    "확장성: 노드 수 확장은 DP4 소관",
    "단, 결정 비용(0.1 / 1 / 10 ms)은 Q1 sweep의 민감도로 유지",
    "수치 임계값·별점은 통합 단계(memo-threshold-unification)에서 정함. 이 문서는 점수를 쓰지 않음"]), size=10.5, space=4)
pageno(s, 2)
notes(s, "핵심: Q1·Q2·Q3을 본다. Q3은 HBM 점유율(낮을수록 좋음)로 재정의했다.\n"
         "근거: 끌어오기는 HBM 용량·대역폭·링크를 소모하므로 지연만 보면 손해가 처리량에서는 이득일 수 있다.\n"
         "연결: 변경 용이성·확장성은 후보 간 변별력이 없거나 DP4 소관이라 제외했다. 결정 비용은 민감도로 남긴다.")

# ================================================================ 3 후보 축
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 설계", 2)
band(s, "DP-02. 후보를 가르는 축 선정  [1/3]", [("Q1", CQ), ("Q2", CQ), ("Q3", CQ), ("C1", CC), ("C2", CC)])
table(s, 0.27, 1.25, [2.6, 7.65, 2.6], [
    ["후보 축", "판단", "결과"],
    ["결정 시점 (C1 inline / C2 사전 계획)", "노드 내 + Turn 단위에서는 후보가 Tier 수 이하라 결정 비용이 무시할 만함 → C2의 존재 이유(결정 지연 은닉)가 사라져 C1이 앞설 가능성 [가설]. 기존 평가에서도 Q1 동점 (954 / 956 tok/s)", "축으로 약함"],
    ["실행 방식 (이동 후 GPU / 제자리 / 분할-병합)", "결정의 출력 값. 결정 구조와 직교하므로 모든 후보가 같은 값을 냄", "공통 요소"],
    ["iteration 경계 스케줄링, layer 스트리밍, 배치 분할", "연산은 GPU에서 시작해야 하므로 현실성 없음 (사용자 결정)", "기각"],
    [[("결정 주체와 정보 범위", {"bold": True})], [("전역 정보로 Cost를 예측해 중앙에서 정할 것인가, 로컬 측정 신호로 Tier가 분산해 정할 것인가. 후보 간 QA가 실제로 갈릴 수 있음", {"bold": True})], [("채택", {"bold": True, "color": RED})]]],
    [0.35, 0.95, 0.55, 0.6, 0.75], size=10.5, fills={(4, 1): LBLUE, (4, 2): LBLUE})
dashed(s, 0.38, 5.05, 12.47, 1.85, "두 후보 (같은 ExecTarget 값, 같은 병합 가정)", 4.8)
rect(s, 0.7, 5.4, 5.8, 1.3, [("[A안] 전역 Cost 기반 중앙 오케스트레이터", {"bold": True, "size": 12, "color": BLUE}),
                              ("전역 상태로 Cost를 예측해 Turn 시작 시 모든 구간의 실행 위치를 한 번에 확정", {"size": 10})], fill=LBLUE, line=BLUE, lw=1.2)
rect(s, 6.8, 5.4, 5.8, 1.3, [("[B′안] 로컬 신호·규칙 기반 분산 admission", {"bold": True, "size": 12, "color": GRN}),
                              ("각 Tier가 자기 KV에 대해 로컬 측정 신호로 접수/거절, 거절은 GPU 폴백", {"size": 10})], fill=LGRN, line=GRN, lw=1.2)
tbox(s, 6.0, 5.85, 1.0, 0.4, [("vs", {"align": PP_ALIGN.CENTER})], size=14, bold=True, color=G59)
keymsg_y = 6.95
tbox(s, 0.42, 6.9, 12.5, 0.26, ["바뀌는 변수는 '누가, 어떤 정보로 정하는가' 하나다."], size=10, bold=True, color=KEY)
pageno(s, 3)
notes(s, "핵심: 결정 시점 축은 노드 내에서 변별력을 잃으므로 결정 주체·정보 범위를 후보 축으로 삼는다.\n"
         "근거: 기존 C1/C2는 Q1에서 동점이었고 확장성에서만 갈렸는데 확장성은 DP4로 이동했다.\n"
         "결정: 같은 출력값을 쓰고 결정 구조만 바꿔 단일 변수 비교가 되도록 한다.")

# ================================================================ 4 A vs B' 비교표
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 설계", 2)
band(s, "Turn 시작 시 KV 구간별 attention을 어디서, 누가 정할 것인가", [("Q1", CQ), ("Q2", CQ), ("Q3", CQ), ("C1", CC), ("C2", CC)])
HDR = [("설계 안", {})]
rows = [
    ["설계 안", [("(1안) 전역 Cost 기반 중앙 오케스트레이터  ", {"segs": [("(1안) 전역 Cost 기반 중앙 오케스트레이터 ", {"bold": True}), ("+ HBM 기회비용 항", {"bold": True, "color": RED})]})],
     [("(2안) 로컬 신호·규칙 기반 분산 admission  ", {"segs": [("(2안) 로컬 신호·규칙 기반 분산 admission ", {"bold": True}), ("+ GPU HBM 예산 admission", {"bold": True, "color": RED})]})]],
    ["구조", "", ""],
    ["장점", bullets(["전역 자원(링크, HBM)을 함께 최적화 [가설]", "HBM 기회비용을 목적함수에 명시 → Q1·Q3 유리 [가설]"]),
     bullets(["측정값 기반이라 Cost 오차·jitter에 강건 [가설]", "중앙 결정 비용 거의 없음, Tier별 책임 분리"])],
    ["단점", bullets([("Cost Model 예측 오차에 의존, Turn 중 jitter 반영 불가", {"color": RED}), "λ_HBM 보정 필요"]),
     bullets([("전역 HBM·링크를 못 봄, 동시 거절 시 폴백 폭주 가능", {"color": RED}), "Tier별 임계치 보정 필요"])],
    ["Tradeoff\n(평가 전 가설)", [("Q1 고동시성: 우세 예상", {"size": 10}), ("Q2 꼬리: Cost 오차·jitter 크면 열세", {"size": 10}), ("Q3 HBM: 우세 예상", {"size": 10})],
     [("Q1 고동시성: 폭주 시 열세 가능", {"size": 10}), ("Q2 꼬리: 측정 기반이라 강건할 수 있음", {"size": 10}), ("Q3 HBM: 예산 admission 필요", {"size": 10})]]]
table(s, 0.27, 1.14, [0.96, 5.95, 5.94], rows, [0.35, 3.0, 0.72, 0.78, 0.85], size=10.5)
# --- diagrams in 구조 row (y 1.49 .. 4.49)
# A
x0, y0 = 1.3, 1.6
rect(s, x0, y0 + 0.15, 0.95, 0.5, ["vLLM Scheduler", ("Turn 시작", {"size": 7.5})], fill=GF0, line=G7F, size=8, bold=True)
container(s, x0 + 1.25, y0, 4.5, 1.5, "DP2 Planner (중앙) — 동기, Turn 시작 시 1회", LBLUE, BLUE, 8.5)
rect(s, x0 + 1.37, y0 + 0.32, 1.9, 0.5, ["Global State Table", ("Tier 점유 · 링크 BW · HBM 예산", {"size": 7.5})], size=8, bold=True)
rect(s, x0 + 3.45, y0 + 0.32, 2.2, 0.5, ["Cost Evaluator", ("T_attn + T_rt + T_link + λ·ΔHBM", {"size": 7.5})], size=8, bold=True)
rect(s, x0 + 1.37, y0 + 0.95, 4.28, 0.42, ["ExecutionPlan [(req, seg) → GPU_STAGE | IN_SITU(tier)]"], fill=C("FFF2CC"), line=C("BF9000"), size=8, bold=True)
arrow(s, x0 + 0.95, y0 + 0.4, x0 + 1.25, y0 + 0.4); arrow(s, x0 + 3.27, y0 + 0.57, x0 + 3.45, y0 + 0.57); arrow(s, x0 + 4.5, y0 + 0.82, x0 + 4.5, y0 + 0.95)
rect(s, x0 + 1.25, y0 + 1.85, 2.1, 0.55, ["GPU Worker", ("QKV · staging 후 attention", {"size": 7.5})], fill=GF0, line=G7F, size=8, bold=True)
rect(s, x0 + 3.65, y0 + 1.85, 2.1, 0.55, ["Tier Worker", ("q 수신 → fused attention", {"size": 7.5})], fill=LGRN, line=GRN, size=8, bold=True)
arrow(s, x0 + 2.3, y0 + 1.5, x0 + 2.3, y0 + 1.85, BLUE, 1.2); arrow(s, x0 + 4.7, y0 + 1.5, x0 + 4.7, y0 + 1.85, BLUE, 1.2)
rect(s, x0 + 1.25, y0 + 2.6, 4.5, 0.32, ["LSE 병합 (m, l, o) → context vector → FFN"], fill=GF2, line=G7F, size=8, bold=True)
arrow(s, x0 + 2.3, y0 + 2.4, x0 + 2.3, y0 + 2.6); arrow(s, x0 + 4.7, y0 + 2.4, x0 + 4.7, y0 + 2.6)
tbox(s, x0 - 0.1, y0 + 1.0, 1.35, 0.6, [("예측(Cost) 기반", {"color": RED}), ("Turn 중 재결정 없음", {"color": RED})], size=7.5)
# B'
x1 = 7.4
rect(s, x1, y0 + 0.15, 0.95, 0.5, ["Thin Scheduler", ("Task 발행", {"size": 7.5})], fill=GF0, line=G7F, size=8, bold=True)
container(s, x1 + 1.25, y0, 4.5, 1.5, "Tier Admission Agent (CXL-PNM / ScHBM) — Turn 시작 시 1회", LGRN, GRN, 8.5)
rect(s, x1 + 1.37, y0 + 0.32, 4.28, 0.42, ["측정 신호(큐 깊이 · in-situ 점유 · 링크) vs 임계치"], size=8, bold=True)
rect(s, x1 + 1.37, y0 + 0.9, 2.0, 0.45, ["접수 → IN_SITU"], fill=C("FFF2CC"), line=C("BF9000"), size=8, bold=True)
rect(s, x1 + 3.65, y0 + 0.9, 2.0, 0.45, ["거절 → GPU_STAGE"], fill=LORG, line=ORG, size=8, bold=True, color=RED)
arrow(s, x1 + 0.95, y0 + 0.4, x1 + 1.25, y0 + 0.4); arrow(s, x1 + 2.37, y0 + 0.74, x1 + 2.37, y0 + 0.9); arrow(s, x1 + 4.65, y0 + 0.74, x1 + 4.65, y0 + 0.9)
rect(s, x1 + 1.25, y0 + 1.85, 2.1, 0.55, ["Tier Worker", ("q 수신 → fused attention", {"size": 7.5})], fill=LGRN, line=GRN, size=8, bold=True)
rect(s, x1 + 3.65, y0 + 1.7, 2.1, 0.85, ["GPU HBM Budget Admission", ("예산 내 staging → GPU attention", {"size": 7.5})], fill=C("FFF2CC"), line=C("BF9000"), size=8, bold=True)
arrow(s, x1 + 2.3, y0 + 1.35, x1 + 2.3, y0 + 1.85, GRN, 1.2); arrow(s, x1 + 4.65, y0 + 1.35, x1 + 4.65, y0 + 1.7, ORG, 1.2)
rect(s, x1 + 1.25, y0 + 2.6, 4.5, 0.32, ["LSE 병합 (m, l, o) → context vector → FFN"], fill=GF2, line=G7F, size=8, bold=True)
arrow(s, x1 + 2.3, y0 + 2.4, x1 + 2.3, y0 + 2.6); arrow(s, x1 + 4.65, y0 + 2.55, x1 + 4.65, y0 + 2.6)
tbox(s, x1 - 0.1, y0 + 1.0, 1.35, 0.6, [("측정 기반", {"color": RED}), ("거절 시 KV 전체 이동", {"color": RED})], size=7.5)
tbox(s, 0.42, 6.95, 12.4, 0.26, ["주: 원안(분산 자율 중개자)에서 Shared Task Bus를 제거하고 결정 시점을 Turn 시작 1회로 재정의. 신규 Tier 확장성·오버헤드 우위 주장은 근거가 약해 제외. 점수·별점 없음 (평가 전)"], size=8, color=G7F)
pageno(s, 4)
notes(s, "핵심: A는 중앙에서 예측 Cost로, B′는 Tier 로컬 측정 신호로 Turn 시작 시 실행 위치를 정한다.\n"
         "근거: A는 전역 조정과 HBM 기회비용에서, B′는 측정 기반 강건성에서 이점이 있다는 가설이다. 모두 평가 전이다.\n"
         "선택의 의미: 별점을 쓰지 않는 것은 정량 구간이 정의되지 않았기 때문이다. 우열은 예상 방향이다.")

# ================================================================ 5 A 상세
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 설계", 2)
band(s, "[A안] 전역 Cost 기반 중앙 오케스트레이터  [2/3]", [("Q1", CQ), ("Q2", CQ), ("Q3", CQ), ("C1", CC), ("C2", CC)])
dashed(s, 0.38, 1.45, 7.6, 5.4, "구조 및 호출 흐름", 2.4)
rect(s, 0.6, 1.95, 1.7, 0.7, ["vLLM Scheduler", ("Turn 시작: 요청, KV 구간", {"size": 8})], fill=GF0, line=G7F, size=10, bold=True)
container(s, 2.8, 1.85, 4.95, 2.35, "DP2 Planner (중앙, 동기 · 단일 패스)", LBLUE, BLUE, 10)
rect(s, 2.95, 2.25, 2.2, 0.85, ["Global State Table", ("Tier 점유, 링크 잔여 BW,", {"size": 8}), ("HBM 예산", {"size": 8})], size=10, bold=True)
rect(s, 5.4, 2.25, 2.2, 0.85, ["Cost Evaluator", ("argmin Cost(target)", {"size": 8})], size=10, bold=True)
rect(s, 2.95, 3.35, 4.65, 0.7, [("ExecutionPlan", {"bold": True, "size": 10}), ("[(req, seg) → GPU_STAGE | IN_SITU(tier)]", {"size": 9})], fill=C("FFF2CC"), line=C("BF9000"))
arrow(s, 2.3, 2.3, 2.8, 2.3, G59, 1.2); arrow(s, 5.15, 2.67, 5.4, 2.67, G59, 1.2); arrow(s, 6.5, 3.1, 6.5, 3.35, G59, 1.2)
rect(s, 2.8, 4.7, 2.3, 0.8, ["GPU Worker", ("QKV 투영, GPU_STAGE 구간", {"size": 8}), ("staging 후 attention", {"size": 8})], fill=GF0, line=G7F, size=10, bold=True)
rect(s, 5.45, 4.7, 2.3, 0.8, ["Tier Worker", ("CXL-PNM / ScHBM:", {"size": 8}), ("q 수신 → fused attention", {"size": 8})], fill=LGRN, line=GRN, size=10, bold=True)
arrow(s, 3.9, 4.2, 3.9, 4.7, BLUE, 1.5); arrow(s, 6.6, 4.2, 6.6, 4.7, BLUE, 1.5)
tbox(s, 4.0, 4.3, 1.0, 0.25, ["plan 하달"], size=8, color=G7F)
rect(s, 2.8, 6.05, 4.95, 0.5, ["LSE 병합: (m, l, o) → 최종 context vector → FFN"], fill=GF2, line=G7F, size=10, bold=True)
arrow(s, 3.9, 5.5, 3.9, 6.05, G59, 1.2); arrow(s, 6.6, 5.5, 6.6, 6.05, G59, 1.2)
tbox(s, 6.7, 5.65, 0.9, 0.25, ["(m, l, o)"], size=8, color=RED)
tbox(s, 0.55, 3.0, 2.1, 2.2, [("호출 경계", {"bold": True, "size": 10, "color": BLUE}), "Turn 시작 시 1회", "동기, 중앙 단일 패스", "Turn 중 재결정 없음"], size=9, color=G40)
dashed(s, 8.3, 1.45, 4.55, 1.55, "책임", 1.2)
tbox(s, 8.4, 1.7, 4.35, 1.3, bullets(["Cost = T_attn + T_rt + T_link + λ_HBM·ΔHBM점유", "λ_HBM: 부하에 따른 HBM 기회비용 가격 (저부하 ≈ 0)", "HBM 예산·staging 허용도 Planner가 소유"]), size=9.5, space=2)
dashed(s, 8.3, 3.3, 4.55, 1.45, "강점", 1.2)
tbox(s, 8.4, 3.55, 4.35, 1.2, bullets(["전역 자원(링크, HBM)을 함께 최적화해 충돌 감소 [가설]", "HBM 기회비용을 목적함수에 명시 → Q1·Q3 유리 [가설]"]), size=9.5, space=2)
dashed(s, 8.3, 5.05, 4.55, 1.8, "비용 · 약점", 1.5)
tbox(s, 8.4, 5.3, 4.35, 1.55, bullets([("Cost Model 예측 오차에 의존 (DP2-C-1: 정확도는 범위 밖)", {"color": RED}), "Tier jitter·큐 적체를 Turn 중 반영 못 함", "λ_HBM 보정, T_rt(layer당 q/o 왕복)는 미검증 [추정]"]), size=9.5, space=2)
pageno(s, 5)
notes(s, "핵심: A는 중앙 Planner가 전역 상태와 예측 Cost로 Turn 시작 시 모든 구간의 실행 위치를 한 번에 확정한다.\n"
         "근거: HBM 기회비용 항(λ_HBM)이 목적함수에 들어 있어 고동시성에서 끌어오기를 억제할 수 있다.\n"
         "비용: 예측에 의존하므로 Cost Model 오차와 Tier jitter에 취약하다. 모두 평가 전 가설이다.")

# ================================================================ 6 B' 상세
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 설계", 2)
band(s, "[B′안] 로컬 신호·규칙 기반 분산 admission  [3/3]", [("Q1", CQ), ("Q2", CQ), ("Q3", CQ), ("C1", CC), ("C2", CC)])
dashed(s, 0.38, 1.45, 7.6, 5.4, "구조 및 호출 흐름", 2.4)
rect(s, 0.6, 1.95, 1.7, 0.7, ["Thin Scheduler", ("Attention Task 발행", {"size": 8}), ("(req, KV_desc, Deadline)", {"size": 7.5})], fill=GF0, line=G7F, size=10, bold=True)
container(s, 2.8, 1.85, 4.95, 2.1, "Tier Admission Agent (CXL-PNM / ScHBM 앞단)", LGRN, GRN, 10)
rect(s, 2.95, 2.25, 4.65, 0.65, [("측정 신호 vs 임계치 (hysteresis)", {"bold": True, "size": 10}), ("큐 깊이, in-situ 점유 시간, 링크 사용률", {"size": 8})])
rect(s, 2.95, 3.15, 2.2, 0.6, ["접수 → IN_SITU"], fill=C("FFF2CC"), line=C("BF9000"), size=10, bold=True)
rect(s, 5.4, 3.15, 2.2, 0.6, ["거절 → GPU_STAGE"], fill=LORG, line=ORG, size=10, bold=True, color=RED)
arrow(s, 2.3, 2.3, 2.8, 2.3, G59, 1.2); arrow(s, 4.05, 2.9, 4.05, 3.15, G59, 1.2); arrow(s, 6.5, 2.9, 6.5, 3.15, G59, 1.2)
rect(s, 2.8, 4.45, 2.3, 0.8, ["Tier Worker", ("q 수신 → fused attention", {"size": 8})], fill=LGRN, line=GRN, size=10, bold=True)
rect(s, 5.45, 4.3, 2.3, 1.1, ["GPU HBM Budget Admission", ("예산 내 staging 허용,", {"size": 8}), ("없으면 대기 또는 IN_SITU 강제", {"size": 8})], fill=C("FFF2CC"), line=C("BF9000"), size=9.5, bold=True)
arrow(s, 3.9, 3.75, 3.9, 4.45, GRN, 1.5); arrow(s, 6.5, 3.75, 6.5, 4.3, ORG, 1.5)
rect(s, 2.8, 6.05, 4.95, 0.5, ["LSE 병합: (m, l, o) → 최종 context vector → FFN"], fill=GF2, line=G7F, size=10, bold=True)
arrow(s, 3.9, 5.25, 3.9, 6.05, G59, 1.2); arrow(s, 6.6, 5.4, 6.6, 6.05, G59, 1.2)
tbox(s, 0.55, 3.0, 2.2, 3.4, [("호출 경계", {"bold": True, "size": 10, "color": GRN}), "Turn 시작 시 1회", "Tier 로컬 접수/거절", "디코드 중 우회 없음",
                              ("우회 비대칭: 거절 시 KV 전체 이동 (128K ≈ 43GB, ≈ 0.67 s) [추정] → 임계치 hysteresis", {"color": RED, "size": 8.5})], size=9, color=G40, space=3)
dashed(s, 8.3, 1.45, 4.55, 1.55, "책임", 1.2)
tbox(s, 8.4, 1.7, 4.35, 1.3, bullets(["중앙은 Cost 계산 없이 Task만 발행", "원안의 Task Bus 제거: KV가 있는 Tier만 claim 가능 → 경쟁 없음, hop만 증가", "폴백 staging은 GPU 쪽 HBM 예산 admission이 허용"]), size=9.5, space=2)
dashed(s, 8.3, 3.3, 4.55, 1.45, "강점", 1.2)
tbox(s, 8.4, 3.55, 4.35, 1.2, bullets(["측정값 기반이라 Cost 오차·jitter에 강건할 수 있음 [가설]", "중앙 결정 비용 거의 없음, Tier별 책임 분리"]), size=9.5, space=2)
dashed(s, 8.3, 5.05, 4.55, 1.8, "비용 · 약점", 1.5)
tbox(s, 8.4, 5.3, 4.35, 1.55, bullets([("로컬 신호만 보므로 전역 최적 아님", {"color": RED}), "동시 거절 시 링크 스파이크·HBM 압박 가능", "Tier별 임계치 보정 필요. HBM Budget Admission은 사실상 중앙 요소"]), size=9.5, space=2)
pageno(s, 6)
notes(s, "핵심: B′는 Tier마다 로컬 측정 신호로 접수·거절을 정하는 분산 admission이다. 원안의 Task Bus는 경쟁이 없어 제거했다.\n"
         "근거: 측정값을 쓰므로 예측 오차에 강건할 수 있다는 가설이 이 후보의 고유 가치다.\n"
         "주의: 원안의 '즉시 우회'는 디코드 중 전환이라 제약에 어긋나므로 Turn 시작 1회 결정으로 바꿨다. 거절은 KV 전체 이동이라 비싸다.")

# ================================================================ 7 선택 방법 / 검증 계획
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 구현/검증 계획", 3)
band(s, "하나를 선택하는 방법: 사전 등록한 규칙과 시나리오", [("Q1", CQ), ("Q2", CQ), ("Q3", CQ)])
dashed(s, 0.38, 1.45, 4.9, 5.4, "선택 규칙 (사전 등록)", 2.4)
tbox(s, 0.5, 1.75, 4.7, 5.1, [
    ("1. 우선순위", {"bold": True, "size": 11, "color": BLUE}), "Q1 (고동시성 Max SLO Goodput) > Q3 (iso-load HBM 점유) > Q2 (꼬리). 임계값·동률 판정은 통합 단계에서 정함",
    ("2. 판정 방식", {"bold": True, "size": 11, "color": BLUE}), "평균이 아니라 시나리오별 승패로 판단. 어떤 시나리오에서도 크게 지지 않는 구조를 우세로 판정",
    ("3. 단일 변수", {"bold": True, "size": 11, "color": BLUE}), "같은 시뮬레이터, 같은 DP1 배치, 같은 ExecTarget 출력, 같은 병합 가정. 바꾸는 것은 결정 주체·정보 범위뿐",
    ("4. 근거 수준", {"bold": True, "size": 11, "color": BLUE}), "모두 시뮬레이션 [B+C]. 실측 [A] 아님 (GC-1)"], size=10, space=4)
table(s, 5.5, 1.3, [3.55, 3.8], [
    ["시나리오", "보려는 것"],
    ["동시성 sweep (저 → 포화)", "Q1 교차점(끌어오기가 불리해지는 부하), Q3"],
    ["Cost Model 오차 주입 (ε = 0 / 0.2 / 0.4 / 0.6)", "A의 예측 의존성 대 B′의 측정 강건성 (기존 평가 변수 재사용)"],
    ["Tier 연산 jitter, 일시적 큐 적체", "Q2 꼬리, 판단 오류의 영향"],
    ["폭주 부하 (동시 거절)", "B′의 폴백 폭주, 링크·HBM 압박"],
    ["긴 컨텍스트 세션 재개 (64K ~ 256K)", "우회 비대칭 비용, staging 경로"],
    ["결정 비용 민감도 (0.1 / 1 / 10 ms)", "결정 오버헤드의 Q1 영향"]],
    [0.35, 0.62, 0.45, 0.45, 0.45, 0.45, 0.45][:7], size=10, first_color=BLACK, align_first=PP_ALIGN.LEFT)
dashed(s, 5.5, 5.15, 7.35, 1.7, "시뮬레이터 확장 필요 [추정]", 3.0)
tbox(s, 5.6, 5.4, 7.15, 1.45, bullets([
    "현재: Turn 단위 (n_p, n_d) 모델, staging과 연산은 순차(overlap 없음)",
    "필요: KV 구간 단위 배치, Tier 큐 모델, 로컬 임계치 정책. HBM 풀 용량 제약(A08, A24)은 이미 모델링됨",
    "확인: 고동시성 병목이 HBM 풀인지 max_num_seqs=256 / token budget인지 (sim-extension-scope.md 미확인)"]), size=9.5, space=2)
pageno(s, 7)
notes(s, "핵심: 평가 전에 선택 규칙과 시나리오를 먼저 정한다.\n"
         "근거: ε 주입과 jitter 시나리오는 A의 예측 의존성과 B′의 측정 강건성을 직접 가른다.\n"
         "주의: 임계값은 통합 단계에서 정한다. 시뮬레이터 확장 범위는 sim-extension-scope.md를 읽고 확정해야 한다.")

# ================================================================ 8 하이브리드 C
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 설계 (보완)", 2)
band(s, "보완 설계: 하이브리드 C (선택된 구조 + 상대 강점 이식)", [("Q1", CQ), ("Q2", CQ), ("Q3", CQ), ("C1", CC), ("C2", CC)])
table(s, 0.27, 1.25, [0.8, 5.0, 3.3], [
    ["구조", "실패 양상", "보완하는 쪽"],
    ["A", "Cost 오차, Tier jitter, 큐 적체를 Turn 시작 시 못 봄", "B′의 측정 기반 거부권"],
    ["B′", "전역 HBM·링크를 못 봄, 폴백 폭주", "A의 전역 예산·Cost 계획"]], [0.35, 0.45, 0.45], size=10)
dashed(s, 0.38, 3.0, 9.1, 2.35, "C의 구조 (기준이 A일 때) [Claude 제안, 평가 전]", 5.2)
y = 3.45
rect(s, 0.6, y, 1.95, 1.2, [("Planner (A)", {"bold": True, "size": 10}), ("Cost로 계획 P 생성,", {"size": 8}), ("구간별 예측값 첨부", {"size": 8})], fill=LBLUE, line=BLUE)
rect(s, 2.95, y, 1.95, 1.2, [("Tier Veto Agent", {"bold": True, "size": 10}), ("측정 > 예측×(1+tol)이면", {"size": 8}), ("VETO(측정값)", {"size": 8})], fill=LGRN, line=GRN)
rect(s, 5.3, y, 1.85, 1.2, [("재계획 (≤ 1 round)", {"bold": True, "size": 10}), ("VETO 구간만", {"size": 8}), ("측정값 반영", {"size": 8})], fill=LBLUE, line=BLUE)
rect(s, 7.5, y, 1.8, 1.2, [("HBM Budget", {"bold": True, "size": 10}), ("Admission", {"bold": True, "size": 10}), ("예산 내 staging", {"size": 8})], fill=C("FFF2CC"), line=C("BF9000"))
for xa, xb in ((2.55, 2.95), (4.9, 5.3), (7.15, 7.5)):
    arrow(s, xa, y + 0.6, xb, y + 0.6, G59, 1.5)
tbox(s, 0.6, y + 1.22, 2.4, 0.3, ["Turn 시작 직후, commit 전"], size=8, color=G7F)
tbox(s, 2.9, y + 1.22, 2.4, 0.3, ["VETO 없으면 P 그대로 commit"], size=8, color=G7F)
tbox(s, 5.3, y + 1.22, 4.0, 0.3, ["재계획 실패 / 예산 없음 → IN_SITU 유지 또는 대기"], size=8, color=G7F)
dashed(s, 9.75, 1.5, 3.1, 3.85, "C 채택 조건", 1.6)
tbox(s, 9.85, 1.75, 2.95, 3.6, bullets(["선택된 순수 구조 대비 적어도 하나의 상위 QA에서 개선되고 다른 QA를 악화시키지 않을 때만 채택", "개선이 크지 않으면 순수 구조 유지", ("하이브리드가 항상 우월하다는 보장은 없음 (결합 비용이 이득을 넘는 영역 가능)", {"color": RED})]), size=9.5, space=4)
dashed(s, 0.38, 5.75, 6.1, 1.2, "권한 분리 · 루프 방지", 2.6)
tbox(s, 0.5, 5.98, 5.9, 0.95, bullets(["계획 생성: Planner / 거부권: Tier (거부만 가능) / 예산: Planner·Admission", "재계획 1회, Turn 중 전환 없음, 거절 임계치 hysteresis"]), size=9, space=2)
dashed(s, 6.75, 5.75, 6.1, 1.2, "기존 자산 · 추가 비용", 2.6)
tbox(s, 6.87, 5.98, 5.9, 0.95, bullets(["기존 C2 Late Validation과 같은 계열 (Tier 로컬 측정 기반): hard/soft 조건, 백업 → 재계획 → 기본 경로 재사용", ("Tier별 Veto Agent, 로컬 왕복 1회, 결정 주체 2개 → plan 일관성 검증", {"color": RED})]), size=9, space=2)
pageno(s, 8)
notes(s, "핵심: C는 선택된 구조의 약점을 상대의 강점으로 보완한다. A 기준이면 Tier 로컬 측정 거부권을 얹는다.\n"
         "근거: A의 약점(예측 오차, jitter)과 B′의 약점(전역 시야 부재)이 상보적이다.\n"
         "결정: 채택 조건을 두어 개선이 작으면 순수 구조를 유지한다. 기준이 B′일 때는 중앙 예측을 임계치 보정용으로 이식하는 방향이다.")

# ================================================================ 9 왜 처음부터 C가 아닌가
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 설계 근거", 2)
band(s, "Q. 왜 처음부터 하이브리드 C를 제안하지 않았나?")
reasons = [("1  효과의 귀속", "C는 A의 전역 Cost와 B′의 로컬 거부권의 결합이다. 처음부터 C만 설계하면 어느 요소가 효과를 냈는지 분리할 수 없다. 순수형으로 각 요소의 한계 기여를 먼저 측정한다."),
           ("2  설계 변수의 출처", "veto 임계(tol), hysteresis 폭, 재계획 횟수, 폴백 비용은 평가 전에는 근거 없는 숫자다. A / B′ 평가가 거부 빈도, 우회 비용, 상태 신선도를 제공한다."),
           ("3  복잡도와 책임 이중화", "결정 주체가 둘이면 plan 일관성과 루프·진동 위험이 생긴다. 순수형의 한계가 실증된 뒤에만 이 비용이 정당화된다."),
           ("4  기준선 확보", "A는 기존 C1과 연결되는 기준선이고 B′는 새 변수(정보의 국소성)다. 순수형이 있어야 C의 개선이 무엇 대비인지 말할 수 있다."),
           ("5  C가 항상 이기는 것은 아님", "하이브리드가 순수형보다 나쁜 영역이 있을 수 있다. 그래서 C는 선택 후 보완 설계이며 채택 조건을 둔다.")]
yy = 1.3
for t, d in reasons:
    rect(s, 0.27, yy, 1.95, 1.0, [t], fill=GF2, line=GBF, size=11, bold=True, color=BLACK)
    rect(s, 2.22, yy, 5.4, 1.0, [d], fill=WHITE, line=GBF, size=9.5, align=PP_ALIGN.LEFT)
    yy += 1.1
table(s, 7.85, 1.3, [1.95, 3.3], [
    ["예상 반론", "대응"],
    ["\"처음부터 C를 알았으면 비교는 형식적이다\"", "솔직히 상보성 때문에 C 방향은 예견되었다. 비교의 목적은 C의 설계 변수와 요소별 기여를 얻는 측정이다. 결과에 따라 C의 모양이 바뀌거나 불필요하다는 결론도 가능하다"],
    ["\"B′는 A를 돋보이게 하는 허수아비다\"", "B′의 고유 가설(측정 기반 강건성)을 ε 주입, jitter 시나리오에서 검증한다. 원안의 근거 약한 장점은 제외했다"],
    ["\"C가 복잡하기만 하고 이득이 작으면?\"", "채택 조건에 따라 순수 구조를 유지한다"]],
    [0.35, 1.9, 1.0, 0.8][:4] if False else [0.35, 2.2, 1.45, 0.7], size=9.5, first_color=BLACK, align_first=PP_ALIGN.LEFT)
pageno(s, 9)
notes(s, "핵심: 순수형을 먼저 평가하는 이유는 효과 귀속, 설계 변수 확보, 복잡도 정당화, 기준선 확보다.\n"
         "근거: 하이브리드의 veto 임계·hysteresis는 순수형 평가에서 나오는 거부 빈도, 우회 비용이 있어야 근거를 가진다.\n"
         "선택의 의미: C는 결론이 아니라 측정 결과에 따라 모양이 정해지는 보완 설계이며, 이득이 작으면 채택하지 않는다.")

# ================================================================ 10 결론
s = new_slide("DP2. 노드 내 Attention 실행 위치 결정 구조 - 결론 및 다음 단계", 4)
band(s, "전략: 순수형 평가 → 선택 → 보완 설계 C")
steps = [("① A / B′ 평가", "같은 조건, 단일 변수\n시뮬레이션 [B+C]"), ("② 선택", "사전 등록 규칙\nQ1 > Q3 > Q2"), ("③ 약점 분석", "ε, jitter, 폭주 부하로\n실패 양상 확인"),
         ("④ C 설계 변수", "tol, hysteresis,\n재계획 횟수, 폴백 비용"), ("⑤ C 평가", "선택 구조 대비\n채택 조건 확인")]
for i, (t, d) in enumerate(steps):
    sh = shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, 0.27 + i * 2.52, 1.3, 2.62, 1.2, LBLUE if i != 1 else C("FFF2CC"), BLUE if i != 1 else C("BF9000"), 1.0)
    write(sh, [(t, {"bold": True, "size": 11}), (d, {"size": 9})], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.22)
table(s, 0.27, 2.85, [0.5, 9.0, 3.35], [
    ["#", "미결 질문 / 위험", "영향"],
    ["1", "평가 시나리오와 시뮬레이터 확장 범위 확정 (Tier 큐 모델, 로컬 임계치 정책, KV 구간 단위 배치)", "평가 가능성"],
    ["2", "λ_HBM(HBM 기회비용 가격)을 정적 상수 / 부하 비례 / 점유율 비례 중 무엇으로 할지", "A의 핵심 파라미터"],
    ["3", "부분 결과 병합 kernel 지원 가정(GC-5)을 인정할지. 불인정 시 한 시퀀스의 KV는 한 곳에서만 실행 → 후보 공간 축소", "전 후보 공통"],
    ["4", "T_rt(layer당 q 송신·o 수신 왕복)가 IN_SITU TPOT를 지배할 가능성. PCIe/CXL 왕복 지연 확인 필요 [추정]", "결정 기준 자체"],
    ["5", "통합 단계의 QA 임계값·별점 구간 확정 후 선택 규칙에 반영", "선택 규칙"],
    ["6", "DP4 확장: 노드 역할(P형/D형)을 DP2 후보 공간의 제한으로 정식화 (별도 논의)", "DP4 경계"]],
    [0.35, 0.55, 0.55, 0.55, 0.55, 0.55, 0.55], size=10, first_color=BLACK, align_first=PP_ALIGN.LEFT)
pageno(s, 10)
notes(s, "핵심: 순수형 평가 → 선택 → 약점 분석 → C 설계 변수 도출 → C 평가의 순서로 진행한다.\n"
         "근거: 각 단계의 산출물(거부 빈도, 우회 비용 등)이 다음 단계의 설계 입력이 된다.\n"
         "미결: 시뮬레이터 확장 범위, λ_HBM, 병합 kernel 가정, T_rt가 가장 큰 불확실성이다.")

# ================================================================ 11 appendix
s = new_slide("Appendix. 수치와 가정의 근거 수준", 4)
band(s, "검증 상태 구분: 추정 / 시뮬레이션 / 가정 / 미검증")
table(s, 0.27, 1.25, [2.3, 4.2, 1.9, 4.45], [
    ["항목", "값 / 내용", "근거 수준", "가정"],
    ["KV 크기", "약 320KB/token, 128K에서 약 43GB", "추정 (계산)", "Llama-3.1-70B, GQA 8, 80 layer, BF16"],
    ["KV 이동 시간", "약 0.67 s", "추정 (계산)", "PCIe 64 GB/s (GC-9), 전 layer 일괄 이동, 겹침 없음"],
    ["attention 연산 강도", "약 Tq × 8 FLOP/byte (H100 ridge ≈ 300)", "추정 (계산)", "GQA 8, KV BF16. Tq > 약 40에서 GPU compute-bound"],
    ["제자리 연산 break-even", "Tq ≲ 약 390 (ScHBM 198 TFLOPS), ≲ 약 6 (CXL-PNM 3.3 TFLOPS)", "추정 (계산)", "128K History, 링크 경합·병합 비용·겹침 미반영. 방향성 참고용"],
    ["기존 C1/C2 Q1", "954 / 956 tok/s (Baseline 575)", "시뮬레이션 [B+C]", "DP2-requirements.md QS-1, 기존 Turn 단위 모델"],
    ["T_rt, 결정 비용, λ_HBM", "미정", "미검증", "구현·평가로 확인. 결정 비용은 0.1/1/10 ms 민감도"],
    ["부분 결과 병합 kernel", "주어진 것으로 가정", "가정 (GC-5)", "확인 필요. 불인정 시 후보 공간 축소"]],
    [0.35, 0.55, 0.55, 0.55, 0.55, 0.55, 0.55, 0.55][:8], size=10, first_color=BLACK, align_first=PP_ALIGN.LEFT)
tbox(s, 0.42, 6.3, 12.4, 0.8, [("출처", {"bold": True, "size": 10}),
                                "doc-mk/DP2/dp2-decision-structure-candidates-A-B-C.md · Requirements/memo-dp2-redesign-ideation.md · memo-dp2-dp4-boundary.md · DP2-requirements.md · Evaluation/DP2/m0-spec.md. 모든 평가는 시뮬레이션 [B+C]이며 실측 [A]가 아님"], size=9, color=G59)
pageno(s, 11)
notes(s, "핵심: 이 발표의 수치는 모두 가정이 붙은 추정이거나 기존 시뮬레이션 값이며 실측이 아니다.\n"
         "근거: 계산식과 가정을 표에 명시했다.\n"
         "사용: 질의 시 수치의 근거 수준을 구분해 답한다.")

prs.save(OUT)
print("saved", OUT)
