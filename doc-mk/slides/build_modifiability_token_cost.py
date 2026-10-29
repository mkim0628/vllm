"""변경 용이성 평가 기준 장표(토큰 비용은 보조 지표) 생성 스크립트.

build_dp_deck_example.py의 헬퍼를 재사용하고 슬라이드 정의부만 바꾼 것.
실행: python slides/build_modifiability_token_cost.py slides/reference/DP-PPT-style-reference.pptx slides/modifiability-token-cost-criteria.pptx
"""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

REF, OUT = sys.argv[1], sys.argv[2]
FONT = "맑은 고딕"
C = lambda h: RGBColor.from_string(h)
BLACK, WHITE = C("000000"), C("FFFFFF")
G40, G59, G7F, GBF, GD9, GF2, GF0 = C("404040"), C("595959"), C("7F7F7F"), C("BFBFBF"), C("D9D9D9"), C("F2F2F2"), C("F0F1F3")
BLUE, LBLUE = C("2E75B6"), C("DEEBF7")
ORG, LORG = C("C55A11"), C("FBE5D6")
GRN, LGRN = C("548235"), C("E2F0D9")
KEY = C("0066FF")
RED = C("FF0000")
CF, CQ, CC = C("203864"), C("833C0B"), C("375623")   # FR / QA / 제약 chips
TOTAL = 1

prs = Presentation(REF)
# ---- drop the reference slides (keep master/layout/theme)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
LAY = [l for l in prs.slide_layouts if l.name == "1_제목 및 내용"][0]
for sh in list(LAY.shapes):
    if sh.shape_type == 6:          # nav group -> drawn per slide
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "/ 23":
        r = sh.text_frame.paragraphs[0].runs[0]
        r.text = f"/ {TOTAL}"
        for extra in sh.text_frame.paragraphs[0].runs[1:]:
            extra.text = ""
NAV = ["과제 소개", "요구사항", "설계", "구현/검증", "결론"]
NX = [8.227, 9.203, 10.186, 11.162, 12.145]
NW = [1.128, 1.128, 1.128, 1.128, 0.976]


def set_font(run, size, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=0.8, dash=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
        if dash:
            ln = s.line._get_or_add_ln()
            d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    s.shadow.inherit = False
    return s


def write(s, paras, size=10, color=BLACK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.05, space=2):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for p in paras:
        t, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = o.get("align", align)
        para.space_after = Pt(o.get("space", space))
        for st, so in (o.get("segs") or [(t, {})]):
            r = para.add_run(); r.text = st
            set_font(r, o.get("size", size), so.get("bold", o.get("bold", bold)), so.get("color", o.get("color", color)))
    return s


def tbox(slide, x, y, w, h, paras, **kw):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return write(s, paras, **kw)


def rect(slide, x, y, w, h, paras, fill=WHITE, line=G59, lw=0.8, size=9, bold=False, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, color=BLACK, dash=False, kind=MSO_SHAPE.RECTANGLE, **kw):
    s = shape(slide, kind, x, y, w, h, fill, line, lw, dash)
    return write(s, paras, size=size, bold=bold, align=align, anchor=anchor, color=color, **kw)


def arrow(slide, x1, y1, x2, y2, color=G7F, w=1.0, dash=False, both=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    if both:
        h = etree.SubElement(ln, qn("a:headEnd")); h.set("type", "triangle")
    t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle")
    return c


def new_slide(title, nav):
    s = prs.slides.add_slide(LAY)
    s.shapes.title.text = title
    for i, (n, x, w) in enumerate(zip(NAV, NX, NW)):
        cur = i == nav
        sh = shape(s, MSO_SHAPE.CHEVRON, x, 0.15, w, 0.391, G59 if cur else GBF)
        write(sh, [n], size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    shape(s, MSO_SHAPE.RECTANGLE, 8.234, 0.15, 0.232, 0.391, G59 if nav == 0 else GBF)
    shape(s, MSO_SHAPE.RECTANGLE, 12.893, 0.15, 0.232, 0.391, G59 if nav == 4 else GBF)
    return s


def pageno(s, n):
    tbox(s, 5.85, 7.2, 0.68, 0.26, [(str(n), {"align": PP_ALIGN.RIGHT})], size=10, color=G7F)


def band(s, text, chips=()):
    shape(s, MSO_SHAPE.RECTANGLE, 0.26, 0.65, 12.85, 0.43, G7F)
    shape(s, MSO_SHAPE.RECTANGLE, 0.44, 0.74, 0.07, 0.24, WHITE)
    tbox(s, 0.51, 0.68, 8.3, 0.36, [text], size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    x = 13.01
    for name, col in reversed(chips):
        w = 0.78
        x -= w + 0.04
        b = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.69, w, 0.35, col)
        write(b, [name], size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)


def dashed(s, x, y, w, h, label, lw=2.6):
    shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, None, GBF, 1.0, dash=True)
    r = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + (w - lw) / 2, y - 0.17, lw, 0.34, GF2, G7F, 0.8)
    write(r, [label], size=12, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)


def container(s, x, y, w, h, title, fill, line, tsize=9):
    shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, 1.0)
    r = shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.24, line, line)
    write(r, [title], size=tsize, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)


def keymsg(s, y, text, issues):
    ar = shape(s, MSO_SHAPE.RIGHT_ARROW, 0.31, y, 0.45, 0.3, BLUE, None)
    tbox(s, 0.8, y - 0.02, 12.2, 0.34, [text], size=12, bold=True, color=KEY, anchor=MSO_ANCHOR.MIDDLE)
    tbox(s, 0.42, y + 0.34, 12.5, 0.6, issues, size=11, bold=True, space=1)


def table(s, x, y, colw, rows, heights, hdr=True, size=10.5, fills=None, first_color=G7F, align_first=PP_ALIGN.CENTER):
    nr, nc = len(rows), len(colw)
    gs = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)), Inches(sum(heights)))
    tb = gs.table
    tbl = gs._element.graphic.graphicData.tbl
    sid = tbl.tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # no style, no grid
    tbl.tblPr.set("bandRow", "0"); tbl.tblPr.set("firstRow", "0")
    for i, w in enumerate(colw):
        tb.columns[i].width = Inches(w)
    for r in range(nr):
        tb.rows[r].height = Inches(heights[r])
        for c in range(nc):
            cell = tb.cell(r, c)
            val = rows[r][c]
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            items = val if isinstance(val, list) else [val]
            for k, it in enumerate(items):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                text, o = (it, {}) if isinstance(it, str) else it
                para.space_after = Pt(1)
                head = hdr and r == 0
                if c == 0 or head:
                    para.alignment = o.get("align", PP_ALIGN.CENTER if (head or align_first == PP_ALIGN.CENTER) else PP_ALIGN.LEFT)
                for st, so in (o.get("segs") or [(text, {})]):
                    run = para.add_run(); run.text = st
                    col = first_color if (c == 0 or head) else so.get("color", o.get("color", BLACK))
                    set_font(run, o.get("size", size + (1 if head else 0)), so.get("bold", o.get("bold", head and c > 0)), col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (fills or {}).get((r, c), WHITE)
            tcPr = cell._tc.get_or_add_tcPr()
            fill_el = tcPr.find(qn("a:solidFill"))
            if fill_el is not None:
                tcPr.remove(fill_el)
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(tag)); ln.set("w", "6350")
                sf = etree.SubElement(ln, qn("a:solidFill")); cl = etree.SubElement(sf, qn("a:srgbClr")); cl.set("val", "BFBFBF")
            if fill_el is not None:
                tcPr.append(fill_el)
    return gs


def bullets(items, size=10.5):
    return [("• " + t if isinstance(t, str) else ("• " + t[0], t[1])) for t in items]


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t



# ================================================================ 변경 용이성 평가 기준: 토큰 비용은 보조 지표
s = new_slide("공통. 변경 용이성(Modifiability) 평가 기준 - 요구사항", 1)
band(s, "변경 용이성은 '변경 범위'로 판정하고, 에이전트 토큰 비용은 보조 지표로만 쓴다", [("Modif.", CQ)])

# ---- 좌: 평가 기준
dashed(s, 0.38, 1.5, 6.1, 4.6, "평가 기준 [제안]", 2.2)
table(s, 0.5, 1.8, [1.0, 2.5, 2.0], [
    ["구분", "지표 (시나리오 S1~S6별 측정)", "역할"],
    [[("주 지표", {"bold": True, "color": BLUE})],
     [("변경 module 수", {"bold": True, "size": 10}), ("외부(OSS) 코드 수정 여부, 재배포 필요 여부", {"size": 9.5})],
     "구조가 변경을 얼마나 국소화하는가. 별점의 근거"],
    [[("보조 1", {"bold": True, "color": G59})],
     [("개발 공수 (man-day)", {"bold": True, "size": 10}), ("사람 리뷰·빌드·upstream 재검증 포함", {"size": 9.5})],
     "사람 쪽 비용을 반영"],
    [[("보조 2", {"bold": True, "color": ORG})],
     [("에이전트 토큰 비용 ($)", {"bold": True, "size": 10, "color": ORG}), ("입력(코드 읽기 × 턴 수) + 출력(작성 × 재작업)", {"size": 9.5})],
     "방향 교차 확인과 민감도용. 별점 결정에는 단독 사용 금지"]],
    [0.36, 0.8, 0.65, 0.85], size=10, align_first=PP_ALIGN.CENTER,
    fills={(3, 0): LORG, (3, 1): LORG, (3, 2): LORG, (0, 0): GF2, (0, 1): GF2, (0, 2): GF2})

tbox(s, 0.5, 4.72, 5.9, 0.25, [("보조 지표 보고 규칙 [제안]", {"bold": True, "size": 10.5, "color": ORG})], size=10)
tbox(s, 0.5, 5.0, 5.9, 1.1, bullets([
    "후보 간 상대 비교로만 쓰고 절대 $는 근거로 쓰지 않는다 (가격·모델 가정)",
    "E(OSS 탐색 토큰)=0, 시나리오 가중 같게 두는 민감도를 항상 병기한다",
    "구조 지표와 반대 방향이면 원인 분석 전에는 결론을 내지 않는다"]), size=10, color=G40, space=3)

# ---- 우: 제약
dashed(s, 6.75, 1.5, 6.1, 4.6, "토큰 비용만으로는 무리인 이유 (제약)", 3.9)
LIM = [
    ("1", "독립 증거가 아니다",
     "토큰·공수·module 수가 같은 시나리오 사실(M, L, E)에서 파생된다. 6개 시나리오 모두 방향이 일치하는 것은 구성상 당연한 결과이지 교차 검증이 아니다 [모델]"),
    ("2", "가정 하나에 우열이 뒤집힌다",
     "균등 가중에서 후보 1 $1.02 대 후보 2 $0.48이 E=0이면 $0.45 대 $0.48로 동점이 된다. 시나리오 가중에 따라서도 $1.28 대 $0.38까지 벌어진다 [가정]"),
    ("3", "사람 리뷰·이해 비용을 못 담는다",
     "토큰은 에이전트의 작성 비용이다. 변경이 머지되는 병목은 diff 리뷰, 설계 이해, 롤아웃이며 코드 줄 수가 작아도 OSS 내부 동작에 기대면 리뷰 부담이 크다"),
    ("4", "변경이 올바른지·유지되는지 못 잰다",
     "토큰이 적다고 올바른 변경이 아니다. 업스트림 변경으로 조용히 깨지는 wrapper 의존, 회귀는 재작업 횟수(고정 ×2)로 가정했을 뿐 실제 에이전트 세션은 실행하지 않았다 [추정]"),
    ("5", "시점·모델에 묶이고 단건만 본다",
     "가격(2026-10 가정)과 모델 tier·성능이 바뀌면 절대값이 바뀐다. 한 번의 변경 비용만 보며 변경이 쌓이는 설계 침식은 반영하지 못한다"),
]
for i, (n, t, d) in enumerate(LIM):
    y = 1.72 + i * 0.86
    b = shape(s, MSO_SHAPE.OVAL, 6.88, y + 0.04, 0.3, 0.3, CC)
    write(b, [n], size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    tbox(s, 7.25, y - 0.02, 5.55, 0.84, [(t, {"bold": True, "size": 10.5}), (d, {"size": 9, "color": G40})], size=9, space=1)

# ---- 하단 핵심 메시지
keymsg(s, 6.27, "토큰 비용은 '변경이 얼마나 비싼가'를 가리키는 보조 근거이며, 변경 용이성의 판정은 구조(변경 범위)와 사람 리뷰가 맡는다",
       ["근거 수준 [C] 모델·가정 기반이며 실측 아님. 수치 출처: DP4/qa-model/modifiability_model.py (S1~S6, 가정 상수), Evaluation/DP1/qa4-preregistration.md"])
pageno(s, 1)
notes(s, "핵심: 변경 용이성은 변경 module 수·외부 코드 수정 여부 같은 구조 지표로 판정하고, 에이전트 토큰 비용은 방향을 교차 확인하는 보조 지표로만 쓴다.\n"
         "근거: (1) 토큰·공수·module은 같은 시나리오 사실에서 계산되므로 독립 증거가 아니다. (2) OSS 탐색 토큰 E 하나를 0으로 두면 $1.02 대 $0.48이 $0.45 대 $0.48 동점으로 바뀐다. "
         "(3) 사람 리뷰·이해 비용, (4) 변경의 정확성·유지 가능성, (5) 가격·모델 의존과 단건 측정은 토큰으로 담기지 않는다.\n"
         "선택의 의미: 토큰 비용을 별점의 단독 근거로 쓰면 가정 하나로 결론이 바뀌므로, 병기와 민감도 보고로 제한한다. 질문이 오면 '이미 검토했고 보조로만 둔 이유'로 답한다.\n"
         "출처: DP4/qa-model/modifiability_model.py, Evaluation/DP1/qa4-preregistration.md §4, DP4 설계 문서 §7.3~7.4.")

prs.save(OUT)
print("saved", OUT)
