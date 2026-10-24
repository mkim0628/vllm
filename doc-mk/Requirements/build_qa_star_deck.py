"""QA 별 기준 근거 슬라이드 제작 (doc-mk/slides/README.md 스타일).

구성: 1장 품질 요구사항 표(Scenario 칸 채움) + 2~6장 QA-01~05 별 기준과 근거(1 QA = 1장).
별 규칙: ★ = 시나리오 충족(하한), ★★★ = 상한(문헌·환산 근거), ★★ = 두 값의 등분 중점. 하한 미달은 미충족(☆).

실행(저장소 루트에서):
  .venv/bin/python doc-mk/Requirements/build_qa_star_deck.py \
      doc-mk/slides/reference/DP-PPT-style-reference.pptx doc-mk/Requirements/qa-star-evidence-slides.pptx
근거와 검증 상태: doc-mk/Requirements/qa-star-criteria-evidence.md
"""
import sys, glob, os, statistics
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

REF, OUT = sys.argv[1], sys.argv[2]
FONT = "맑은 고딕"
C = lambda h: RGBColor.from_string(h)
BLACK, WHITE = C("000000"), C("FFFFFF")
G40, G59, G7F, GBF, GD9, GF2, GF0 = C("404040"), C("595959"), C("7F7F7F"), C("BFBFBF"), C("D9D9D9"), C("F2F2F2"), C("F0F1F3")
BLUE, LBLUE = C("2E75B6"), C("DEEBF7")
ORG, LORG = C("C55A11"), C("FBE5D6")
GRN, LGRN = C("548235"), C("E2F0D9")
YEL, LYEL = C("BF9000"), C("FFF2CC")
KEY, RED = C("0066FF"), C("C00000")
CQ = C("833C0B")
TOTAL = 6
ROOT = os.path.dirname(os.path.abspath(OUT)) + "/../.."

prs = Presentation(REF)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
LAY = [l for l in prs.slide_layouts if l.name == "1_제목 및 내용"][0]
for sh in list(LAY.shapes):
    if sh.shape_type == 6:
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "/ 23":
        r = sh.text_frame.paragraphs[0].runs[0]
        r.text = f"/ {TOTAL}"
        for extra in sh.text_frame.paragraphs[0].runs[1:]:
            extra.text = ""
NAV = ["과제 소개", "요구사항", "설계", "구현/검증", "결론"]
NX = [8.227, 9.203, 10.186, 11.162, 12.145]
NW = [1.128, 1.128, 1.128, 1.128, 0.976]


def set_font(run, size, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=0.8, dash=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
        if dash:
            ln = s.line._get_or_add_ln()
            d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    s.shadow.inherit = False
    return s


def write(s, paras, size=10, color=BLACK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.05, space=2):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for p in paras:
        t, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = o.get("align", align)
        para.space_after = Pt(o.get("space", space))
        for st, so in (o.get("segs") or [(t, {})]):
            r = para.add_run(); r.text = st
            set_font(r, o.get("size", size), so.get("bold", o.get("bold", bold)), so.get("color", o.get("color", color)))
    return s


def tbox(slide, x, y, w, h, paras, **kw):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return write(s, paras, **kw)


def rect(slide, x, y, w, h, paras, fill=WHITE, line=G59, lw=0.8, size=9, bold=False, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, color=BLACK, dash=False, kind=MSO_SHAPE.RECTANGLE, **kw):
    s = shape(slide, kind, x, y, w, h, fill, line, lw, dash)
    return write(s, paras, size=size, bold=bold, align=align, anchor=anchor, color=color, **kw)


def vline(slide, x, y1, y2, color=G7F, w=1.0, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y1), Inches(x), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    if dash:
        ln = c.line._get_or_add_ln()
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    return c


def hline(slide, x1, x2, y, color=G7F, w=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    c.line.color.rgb = color; c.line.width = Pt(w)
    return c


def new_slide(title, nav):
    s = prs.slides.add_slide(LAY)
    s.shapes.title.text = title
    for i, (n, x, w) in enumerate(zip(NAV, NX, NW)):
        cur = i == nav
        sh = shape(s, MSO_SHAPE.CHEVRON, x, 0.15, w, 0.391, G59 if cur else GBF)
        write(sh, [n], size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    shape(s, MSO_SHAPE.RECTANGLE, 8.234, 0.15, 0.232, 0.391, G59 if nav == 0 else GBF)
    shape(s, MSO_SHAPE.RECTANGLE, 12.893, 0.15, 0.232, 0.391, G59 if nav == 4 else GBF)
    return s


def pageno(s, n):
    tbox(s, 5.85, 7.2, 0.68, 0.26, [(str(n), {"align": PP_ALIGN.RIGHT})], size=10, color=G7F)


def band(s, text, chips=()):
    shape(s, MSO_SHAPE.RECTANGLE, 0.26, 0.65, 12.85, 0.43, G7F)
    shape(s, MSO_SHAPE.RECTANGLE, 0.44, 0.74, 0.07, 0.24, WHITE)
    tbox(s, 0.51, 0.68, 10.6, 0.36, [text], size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    x = 13.01
    for name, col in reversed(chips):
        w = 0.86
        x -= w + 0.04
        b = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 0.69, w, 0.35, col)
        write(b, [name], size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)


def table(s, x, y, colw, rows, heights, hdr=True, size=10.5, fills=None, first_color=G7F, align_first=PP_ALIGN.CENTER,
          hdr_fill=GF2):
    nr, nc = len(rows), len(colw)
    gs = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)), Inches(sum(heights)))
    tb = gs.table
    tbl = gs._element.graphic.graphicData.tbl
    sid = tbl.tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    tbl.tblPr.set("bandRow", "0"); tbl.tblPr.set("firstRow", "0")
    for i, w in enumerate(colw):
        tb.columns[i].width = Inches(w)
    for r in range(nr):
        tb.rows[r].height = Inches(heights[r])
        for c in range(nc):
            cell = tb.cell(r, c)
            val = rows[r][c]
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            items = val if isinstance(val, list) else [val]
            for k, it in enumerate(items):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                text, o = (it, {}) if isinstance(it, str) else it
                para.space_after = Pt(1)
                head = hdr and r == 0
                if c == 0 or head:
                    para.alignment = o.get("align", PP_ALIGN.CENTER)
                for st, so in (o.get("segs") or [(text, {})]):
                    run = para.add_run(); run.text = st
                    col = first_color if (c == 0 or head) else so.get("color", o.get("color", BLACK))
                    if o.get("color") and c == 0:
                        col = o["color"]
                    set_font(run, o.get("size", size), so.get("bold", o.get("bold", head and c > 0)), col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (fills or {}).get((r, c), hdr_fill if (hdr and r == 0) else WHITE)
            tcPr = cell._tc.get_or_add_tcPr()
            fill_el = tcPr.find(qn("a:solidFill"))
            if fill_el is not None:
                tcPr.remove(fill_el)
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(tag)); ln.set("w", "6350")
                sf = etree.SubElement(ln, qn("a:solidFill")); cl = etree.SubElement(sf, qn("a:srgbClr")); cl.set("val", "BFBFBF")
            if fill_el is not None:
                tcPr.append(fill_el)
    return gs


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def keymsg(s, y, text):
    shape(s, MSO_SHAPE.RIGHT_ARROW, 0.31, y + 0.02, 0.4, 0.26, BLUE, None)
    tbox(s, 0.76, y - 0.02, 12.3, 0.34, [text], size=11.5, bold=True, color=KEY, anchor=MSO_ANCHOR.MIDDLE)


# ----------------------------------------------------------------------------------------------
# 근거 수준 색: 실측·실트래픽 = BLUE, 시뮬레이션/조건부 최대 = ORG, 내부 시뮬레이션 = GRN, 환산·원리 = G7F
KIND = {"real": (BLUE, "문헌: 실GPU·실트래픽"), "sim": (ORG, "문헌: 시뮬레이션·조건부 최대"),
        "int": (GRN, "내부 시뮬레이션 [B+C]"), "calc": (G7F, "환산·원리")}


def strip(s, x, y, w, vmin, vmax, ticks, thr, items, title, unit_fmt, rowh=0.2, lw=2.35):
    """가로축 위에 문헌 값(막대/점)과 별 경계(점선)를 그린다. thr=[(값, 라벨)], items=[(라벨, lo, hi, kind, 값표기)]. 끝 y를 반환."""
    pw = w - lw
    px = x + lw
    pos = lambda v: px + (min(max(v, vmin), vmax) - vmin) / (vmax - vmin) * pw
    n = len(items)
    top = y + 0.74
    bot = top + n * rowh
    tbox(s, x, y - 0.02, w, 0.26, [title], size=10, bold=True)
    for k, (v, lab) in enumerate(thr):
        ly = y + 0.24 + (0.2 if k % 2 else 0.0)
        tbox(s, pos(v) - 0.6, ly, 1.2, 0.2, [(lab, {"align": PP_ALIGN.CENTER})], size=8.5, bold=True, color=RED, margin=0)
        vline(s, pos(v), ly + 0.2, bot, RED, 1.0, dash=True)
    hline(s, px, px + pw, bot + 0.02, G7F, 0.75)
    for t_ in ticks:
        vline(s, pos(t_), bot + 0.02, bot + 0.07, G7F, 0.75)
        tbox(s, pos(t_) - 0.3, bot + 0.05, 0.6, 0.2, [(unit_fmt(t_), {"align": PP_ALIGN.CENTER})], size=8, color=G7F, margin=0)
    for i, (lab, lo, hi, kind, val) in enumerate(items):
        yy = top + i * rowh
        col = KIND[kind][0]
        tbox(s, x, yy - 0.01, lw - 0.05, rowh, [(lab, {"align": PP_ALIGN.RIGHT})], size=8, color=G40, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
        if hi > vmax or lo < vmin:       # 축 밖으로 나가는 값: 축 끝까지 막대 + 화살표
            x0, x1 = pos(lo), px + pw
            shape(s, MSO_SHAPE.RECTANGLE, x0, yy + 0.04, max(x1 - x0, 0.05), rowh - 0.08, col)
            tbox(s, x1 + 0.01, yy - 0.01, 0.2, rowh, ["▶"], size=8, color=col, anchor=MSO_ANCHOR.MIDDLE, margin=0)
            if x1 - x0 > 0.9:
                tbox(s, x0 + 0.05, yy - 0.01, 1.2, rowh, [val], size=8, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0)
            else:
                tbox(s, x0 - 1.25, yy - 0.01, 1.2, rowh, [(val, {"align": PP_ALIGN.RIGHT})], size=8, color=col, bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0)
            continue
        if abs(hi - lo) < 1e-9:
            d = 0.12
            shape(s, MSO_SHAPE.OVAL, pos(lo) - d / 2, yy + (rowh - d) / 2, d, d, col)
            tx = pos(lo) + 0.1
        else:
            shape(s, MSO_SHAPE.RECTANGLE, pos(lo), yy + 0.04, max(pos(hi) - pos(lo), 0.05), rowh - 0.08, col)
            tx = pos(hi) + 0.05
        if tx > px + pw - 1.0:           # 오른쪽 끝이면 막대·점 왼쪽에 값 표기
            tbox(s, pos(lo) - 1.15, yy - 0.01, 1.05, rowh, [(val, {"align": PP_ALIGN.RIGHT})], size=8, color=col, bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0)
        else:
            tbox(s, tx + 0.04, yy - 0.01, 1.2, rowh, [val], size=8, color=col, bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    lx = x
    for kk, (col, name) in KIND.items():
        if not any(it[3] == kk for it in items):
            continue
        shape(s, MSO_SHAPE.RECTANGLE, lx, bot + 0.34, 0.12, 0.1, col)
        tbox(s, lx + 0.15, bot + 0.29, 1.8, 0.2, [name], size=7.5, color=G59, margin=0)
        lx += 0.15 + len(name) * 0.1 + 0.15
    return bot + 0.56


def ladder(s, x, y, colw, header, rows, heights, size=9.5):
    """별 사다리 표. rows[i][0]은 별 표기 문자열."""
    data = [header] + rows
    fills = {}
    for r in range(1, len(data)):
        fills[(r, 0)] = [GF2, LYEL, LBLUE, LGRN][min(r - 1, 3)] if len(data) == 5 else GF2
    star_cols = {1: G7F, 2: YEL, 3: BLUE, 4: GRN}
    table(s, x, y, colw, data, [0.3] + heights, size=size, fills=fills, first_color=BLACK)


def source_line(s, text, y=6.92):
    tbox(s, 0.3, y, 12.7, 0.3, [text], size=7.5, color=G7F, space=0)


# ================================================================== 슬라이드 정의
# ---- 문서 표(1장)에 들어갈 Scenario 텍스트: (문장, 측정 줄). 측정 줄은 굵게.
ROWS = [
    ("QA-01", "Performance\n(Time behavior)", "서비스 수준 목표(SLO)를\n만족하는 토큰 처리량",
     "다수 agent 세션 동시 요청 시 지정 SLO*를 만족하는 최대 요청량이 Baseline보다 커야 한다 (*SLO: Service Level Objective)",
     "[측정: SLO 만족 Max Goodput 증가율 ★1.5x / ★★2.25x / ★★★3.0x (Pensieve·KVFlow 등 문헌), SLO: TTFT 2s, TPOT 2ms]",
     "H", "H", "1"),
    ("QA-02", "Resource\nEfficiency", "HBM 점유율",
     "동일한 성능 조건에서 HBM 점유율이 Baseline보다 낮아야 한다.",
     "[측정: HBM 점유율 감소율(Baseline 대비) ★20% / ★★47.5% / ★★★75% (KV 압축·offload 문헌 상한 3~4x, KVzip)]",
     "H", "H", "2"),
    ("QA-03", "Modifiability", "변경 용이성",
     "새 메모리 tier·이동/배치 정책·데이터 객체 유형을 서빙 경로 수정 없이 반영하고, vLLM upstream 갱신에도 최소 수정으로 따라갈 수 있어야 한다.",
     "[측정: 토큰 비용 ★<0.7M ★★<0.4M ★★★<0.17M개, 공수 ★<1.0 ★★<0.75 ★★★<0.5 MM (SWE-bench 평균 1.7 file, vLLM connector 중앙값 356 LOC)]",
     "H", "M", "3"),
    ("QA-04", "Scalability", "확장성",
     "실장 환경을 넘는 규모로 노드·메모리 type이 늘어도 결정이 병목이 되지 않고 성능이 자원 증가에 비례해야 한다 (시뮬레이션 평가).",
     "[측정: Scaling Efficiency(N=16) ★>80% / ★★>85% / ★★★>90% (대규모 병렬 효율 보고 80%, 조율 직렬 비중 환산)]",
     "H", "M", "4"),
    ("QA-05", "Functional\nCorrectness", "모델 출력 정확도",
     "LLM 모델 출력의 정확도는 보장되어야 한다.",
     "[측정: LLM 출력 정확도(F1) Baseline 대비 감소율 ★<3% / ★★<2% / ★★★<1% (소유자 허용 한도 1%, KV 압축 문헌)]",
     "H", "M", "5"),
    ("QA-06", "Co-existence", "이동 트래픽 공존",
     "데이터 이동·복원 트래픽과 서빙 요청이 같은 자원을 동시에 쓰는 상황에서도 서빙 성능이 유지되어야 한다.",
     "[측정: 서빙 TTFT·TPOT 증가율 < 10% (SLO 대비 변동 마진, 가정)]",
     "M", "H", "6"),
    ("QA-07", "Fault tolerance", "전송·offload 실패 대체",
     "KV 전송이나 offload가 실패하거나 시간 초과해도 재계산 등 대체 경로로 요청 처리를 이어가야 한다.",
     "[측정: fallback 성공률 > 99.9% (three nines, 월 다운타임 약 43분, 가정)]",
     "M", "M", "7"),
    ("QA-08", "Adaptability", "워크로드 변화 적응",
     "입력 분포나 접근 패턴이 바뀌어도 접근 통계를 반영해 배치를 다시 맞추고 성능을 회복해야 한다.",
     "[측정: 변화 감지 후 성능 회복 시간 < 5분 (변화 전 성능의 90%, 가정)]",
     "M", "M", "8"),
    ("QA-09", "Interoperability", "외부 구성요소 연동",
     "vLLM, llm-d, NIXL 같은 외부 구성요소와 연동할 때 해당 구성요소의 소스를 수정하지 않아야 한다.",
     "[측정: 무수정 연동 인터페이스 수 ≥ 3개 (대상 3종 전부, 가정)]",
     "M", "L", "9"),
    ("QA-10", "Analysability", "결정 근거 진단",
     "이동·배치·라우팅 결정이 기대와 다를 때 결정 입력과 근거 기록으로 원인을 찾을 수 있어야 한다.",
     "[측정: 원인 파악 시간 < 1시간 (개발자 1인, 결정 1건 기준, 가정)]",
     "L", "M", "10"),
]

# ================================================================== 1장: 품질 요구사항 표
s1 = new_slide("품질 요구사항", 1)
tbox(s1, 0.27, 0.66, 12.8, 0.45, [("QAW에서 도출된 품질 요구사항을 Utility Tree를 통해 핵심 품질 속성(H/H or H/M 이상) 선정", {"size": 14})], size=14, anchor=MSO_ANCHOR.MIDDLE)
shape(s1, MSO_SHAPE.RECTANGLE, 0.1, 0.8, 0.14, 0.14, BLACK)
COLW = [0.69, 1.18, 1.73, 7.82, 0.63, 0.65, 0.49]
HDR = ["번호", "QA", "Refinement", "Scenario", "중요도", "난이도", "우선\n순위"]
TOP, HH = 1.15, 0.4
rh = [0.55] * 5 + [0.44] * 5
data = [HDR]
for r in ROWS:
    sent, meas = r[3], r[4]
    data.append([r[0], r[1].split("\n"), r[2].split("\n"),
                 [(sent, {"size": 9}), (meas, {"size": 9, "bold": True})],
                 r[5], r[6], r[7]])
fills = {}
for i in range(1, 6):
    for c in range(7):
        fills[(i, c)] = C("FFF2CC")
gt = table(s1, 0.14, TOP, COLW, data, [HH] + rh, size=9.5, fills=fills, first_color=BLACK, hdr_fill=C("404040"))
# 헤더 글자색 흰색, 본문 셀 정렬 보정
tb = gt.table
for ri in range(1, 6):
    tb.cell(ri, 3).margin_right = Inches(0.5)
for c in range(7):
    for p in tb.cell(0, c).text_frame.paragraphs:
        for r_ in p.runs:
            r_.font.color.rgb = WHITE
            r_.font.bold = True
for ri in range(1, 11):
    for c in (4, 5, 6):
        for p in tb.cell(ri, c).text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r_ in p.runs:
                r_.font.color.rgb = BLACK
    for c in (1, 2):
        for p in tb.cell(ri, c).text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r_ in p.runs:
                r_.font.color.rgb = BLACK
    for p in tb.cell(ri, 0).text_frame.paragraphs:
        for r_ in p.runs:
            r_.font.color.rgb = BLACK
    for p in tb.cell(ri, 3).text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r_ in p.runs:
            r_.font.color.rgb = BLACK
# 상세 슬라이드 링크 아이콘(QA-01~05)
LINK_ICONS = []
yy = TOP + HH
for i in range(5):
    ic = shape(s1, MSO_SHAPE.PENTAGON, 0.14 + sum(COLW[:4]) - 0.42, yy + rh[i] / 2 - 0.09, 0.32, 0.2, BLUE, C("1F4E79"), 0.75)
    LINK_ICONS.append(ic)
    yy += rh[i]
tbox(s1, 0.14, TOP + HH + sum(rh) + 0.12, 12.9, 0.5, [
    "* ★ = 시나리오 충족(하한), ★★ = ★과 ★★★의 등분 중점, ★★★ = 상한(문헌·환산 근거). 하한 미달은 미충족(☆). 파란 화살표는 해당 QA의 근거 슬라이드로 이동.",
    "* QA-01~05 근거는 2~6장, 출처·검증 상태는 qa-star-criteria-evidence.md. QA-06~10은 임시 근거(가정)이며 QA-08 측정(복사 오류 정정)·QA-09 부등호(> 3 → ≥ 3)를 수정함."],
    size=8, color=G59, space=0)
pageno(s1, 1)
notes(s1, "핵심: QA-01~05의 [측정]에 별 3단계 조건(★ 하한 / ★★ 등분 / ★★★ 상한)과 근거 문헌을 채웠다.\n"
          "근거: 각 QA 근거는 2~6장. ★ 값은 소유자가 정한 시나리오 목표, ★★★은 문헌·환산 상한, ★★는 등분 중점이다.\n"
          "주의: QA-06~10은 문헌 확인 전 임시 근거. TPOT 2ms는 원문 그대로이며 기존 문서는 50ms 오타로 해석(확인 필요).")

# ================================================================== QA 근거 슬라이드 공통 함수
def qa_slide(no, qa_id, title, question, ladder_cfg, right_builder, defn, key, caution, sources, notes_text):
    s = new_slide(f"{qa_id}. {title} - 별 기준과 근거", 1)
    band(s, question, [(qa_id, CQ)])
    DEF_H = 0.95
    avail = 6.3 - 1.22 - 0.3 - 0.1 - DEF_H
    k = avail / sum(ladder_cfg["heights"])
    hs = [h_ * k for h_ in ladder_cfg["heights"]]
    ladder(s, 0.27, 1.22, ladder_cfg["colw"], ladder_cfg["header"], ladder_cfg["rows"], hs, size=ladder_cfg.get("size", 10.5))
    # 측정 정의 박스
    ytop = 1.22 + 0.3 + sum(hs) + 0.1
    rect(s, 0.27, ytop, sum(ladder_cfg["colw"]), 6.3 - ytop, [(defn, {"size": 8.5})], fill=GF0, line=GBF, size=8.5, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, color=G40)
    right_builder(s)
    keymsg(s, 6.42, key)
    tbox(s, 0.27, 6.7, 12.8, 0.25, [("주의  ", {"segs": [("주의  ", {"bold": True, "color": RED}), (caution, {})]})], size=8, color=G40, space=0)
    source_line(s, sources, 6.98)
    pageno(s, no)
    notes(s, notes_text)
    return s


def lit_table(s, x, y, colw, rows, heights, size=7.5):
    data = [["문헌", "지표·조건", "값", "별 반영"]] + rows
    table(s, x, y, colw, data, [0.26] + heights, size=size, first_color=BLACK, align_first=PP_ALIGN.LEFT)


STAR = {0: "☆ 미충족", 1: "★", 2: "★★", 3: "★★★"}

# ================================================================== QA-01
def r01(s):
    items = [("InferCept (2402.01869)", 1.6, 2.0, "real", "1.6~2x"),
             ("Mooncake 실트래픽 (2407.00079)", 1.75, 1.75, "real", "+75%"),
             ("KVFlow (2507.07400)", 1.83, 2.19, "real", "1.83~2.19x"),
             ("Pensieve (2312.05516)", 1.14, 3.0, "real", "1.14~3.0x"),
             ("Mooncake 시뮬레이션", 6.25, 6.25, "sim", "+525% (6.25x)"),
             ("DistServe (2401.09670)", 7.4, 7.4, "sim", "7.4x 요청 수"),
             ("LMCache (2510.09665)", 2.3, 14, "sim", "2.3~14x"),
             ("내부 DP1~4 후보 R", 0.99, 1.51, "int", "0.99~1.51x")]
    yb = strip(s, 6.85, 1.2, 6.25, 0.8, 8.0, [1, 2, 3, 4, 5, 6, 7, 8],
               [(1.5, "★ 1.5x"), (2.25, "★★ 2.25x"), (3.0, "★★★ 3.0x")], items,
               "Goodput·처리량 개선 배수 R (Baseline 대비)", lambda v: f"{v:g}x", rowh=0.19)
    lit_table(s, 6.85, yb, [1.45, 2.65, 0.95, 1.2], [
        ["InferCept\narXiv 2402.01869", "tool 호출(augmented LLM) 서빙, 최신 서빙 시스템 대비 처리량", "1.6~2x", "★ 하한 / ★★ 상단"],
        ["Mooncake FAST'25\narXiv 2407.00079", "KVCache 중심 분리 서빙, SLO 하 실트래픽(Kimi) 요청 처리", "+75%", "★ 하한대"],
        ["KVFlow\narXiv 2507.07400", "멀티에이전트 워크플로 prefix cache, SGLang(HiCache) 대비 최대", "1.83x / 2.19x", "★★ 근거"],
        ["Pensieve\narXiv 2312.05516", "다중 턴 stateful 서빙, vLLM·TensorRT-LLM 대비 처리량", "1.14~3.0x", "★★★ 상한"],
        ["시뮬·조건부 최대\n(위 3종)", "Mooncake 시뮬 +525%, DistServe 7.4x, LMCache 최대 14x는 특정 trace·baseline·시뮬 조건", "6.25~14x", "상한 제외"]],
        [0.4, 0.4, 0.4, 0.4, 0.44])

qa_slide(2, "QA-01", "Performance (Time behavior)",
         "SLO를 만족하는 Max Goodput이 Baseline의 몇 배일 때 ★ / ★★ / ★★★ 인가",
         {"colw": [0.85, 1.55, 4.0], "header": ["별", "조건 (R = 후보 ÷ Baseline)", "근거"],
          "rows": [["☆ 미충족", "R < 1.5x", "시나리오 목표 미달. 같은 수요에 서버를 2/3로 줄이지 못함"],
                   ["★", "R ≥ 1.5x", "시나리오 목표(하한). 서버 3대 중 1대 절감 환산 [C]. 문헌 하한대: InferCept 1.6x, Mooncake 실트래픽 +75%"],
                   ["★★", "R ≥ 2.25x", "★(1.5x)과 ★★★(3.0x)의 등분 중점. KVFlow 2.19x, InferCept 상단 2x에 대응"],
                   ["★★★", "R ≥ 3.0x", "상한. vLLM·TensorRT-LLM 대비 실GPU 다중 턴 서빙의 보고 상한(Pensieve 3.0x). 6~14x는 시뮬·조건부 최대로 제외"]],
          "heights": [0.5, 0.78, 0.7, 0.88]},
         r01,
         "지표: SLO(TTFT ≤ 2 s, TPOT ≤ 2 ms*)를 만족한 요청의 Max Goodput을 같은 DP Baseline과 비교한 배수 R. 시나리오·시스템 쌍별 값의 기하평균(공통 룰). *TPOT 2 ms는 원문 그대로, 기존 문서는 50 ms 오타로 해석(확인 필요).",
         "★ 1.5x는 실GPU·실트래픽 보고의 하한대, ★★★ 3.0x는 실GPU 보고 상한. 시뮬·조건부 최대치(6~14x)는 상한에서 제외한다.",
         "문헌 지표(처리량·요청 수·goodput)와 baseline·workload가 달라 직접 비교 불가 [문헌 인용]. 내부 시뮬 R 0.99~1.51 중 1.5x 이상은 DP3 C1(1.514)뿐. 기존 v3(★★★ = 4/3)와 기준이 다름.",
         "출처: arXiv 2402.01869, 2407.00079(FAST'25), 2507.07400, 2312.05516, 2401.09670(OSDI'24), 2510.09665(검색 요약 기준, 원문 미확인). 내부 값: doc-mk/Evaluation/unified_stars_check.md",
         "핵심: ★ 1.5x(소유자 목표) / ★★ 2.25x(등분) / ★★★ 3.0x(상한)이다.\n"
         "근거: InferCept 1.6~2x, Mooncake 실트래픽 +75%, KVFlow 최대 2.19x, Pensieve 최대 3.0x. 시뮬·특정 조건 최대치는 상한에서 제외했다.\n"
         "의미: 현재 내부 시뮬레이션 R은 대부분 1.5x 미만이라 이 기준에서는 미충족이 나온다. 기준 변경 영향은 별도 검토가 필요하다.")

# ================================================================== QA-02
def r02(s):
    items = [("KVzip (2505.23416) KV 3~4x 축소", 66.7, 75.0, "real", "67~75%"),
             ("KVQuant 3-bit (2401.18079) 환산", 81.3, 81.3, "calc", "81% (16→3bit)"),
             ("KIVI 2-bit (2402.02750) 환산", 87.5, 87.5, "calc", "88% (16→2bit)"),
             ("내부 DP3 C1 (시뮬)", 71.7, 71.7, "int", "71.7%"),
             ("내부 DP3 C2 (시뮬)", 20.5, 20.5, "int", "20.5%"),
             ("내부 DP1·DP2 후보 (시뮬)", -21.4, 6.0, "int", "-21.4~+6.0%")]
    yb = strip(s, 6.85, 1.2, 6.25, -25, 100, [0, 25, 50, 75, 100],
               [(20, "★ 20%"), (47.5, "★★ 47.5%"), (75, "★★★ 75%")], items,
               "HBM 점유 감소율 (Baseline 대비, %)", lambda v: f"{v:g}%", rowh=0.2)
    lit_table(s, 6.85, yb, [1.45, 2.65, 0.95, 1.2], [
        ["KVzip\narXiv 2505.23416", "KV 크기 3~4x 축소, 디코딩 지연 약 2x 감소, 성능 저하 \"negligible\"", "3~4x\n(67~75%)", "★★★ 상한"],
        ["KVQuant\narXiv 2401.18079", "3-bit KV 양자화, Wikitext-2·C4 perplexity 저하 < 0.1", "bit 환산\n81%", "상한 초과(정확도 확인)"],
        ["KIVI\narXiv 2402.02750", "2-bit KV 양자화, peak memory 2.6x↓(가중치 포함), batch 4x↑, 품질 \"거의 동일\"", "2.6x\n(총메모리)", "참고(정의 상이)"],
        ["vLLM PagedAttention\narXiv 2309.06180", "기존 시스템은 KV 메모리 중 실제 토큰 상태 20.4~38.2%만 사용(단편화). vLLM이 이를 제거", "낭비\n62~80%", "Baseline이 이미 제거 → 20%는 offload·압축 몫"],
        ["InferCept\narXiv 2402.01869", "tool 대기 중 KV 폐기·재계산이 forward 시간의 37~40%", "37~40%\n(시간)", "UC-1: 유휴 KV 내리기 수요"]],
        [0.38, 0.42, 0.42, 0.46, 0.42])

qa_slide(3, "QA-02", "Resource Efficiency",
         "동일 성능 조건에서 HBM 점유율이 Baseline보다 몇 % 줄면 ★ / ★★ / ★★★ 인가",
         {"colw": [0.85, 1.55, 4.0], "header": ["별", "조건 (감소율 = (Baseline − 후보) ÷ Baseline)", "근거"],
          "rows": [["☆ 미충족", "감소율 < 20%", "시나리오 목표 미달. 음수는 후보가 HBM을 더 쓴 경우"],
                   ["★", "감소율 ≥ 20%", "시나리오 목표(하한). DP1 평균 146.6 GiB 중 약 29 GiB, Baseline 점유 잡음(±1.6%)의 12배 이상 [측정+환산]"],
                   ["★★", "감소율 ≥ 47.5%", "★(20%)과 ★★★(75%)의 등분 중점. KV 약 1/2 수준"],
                   ["★★★", "감소율 ≥ 75%", "상한. KV 3~4x 축소가 성능 저하 없이 보고된 문헌 상한(KVzip). 그 이상(81~88%)은 저비트 환산이며 정확도(QA-05)와 trade-off"]],
          "heights": [0.5, 0.8, 0.6, 0.95]},
         r02,
         "지표: 측정 구간 동안 HBM에 올라와 있는 AI Data(KV·LoRA·expert 등, 모델 가중치·활성값 제외)의 시간 평균 크기(GiB)를 같은 DP Baseline과 비교한 감소율. 시나리오·시스템 쌍별 기하평균. 성능은 섞지 않고 SLO 통과를 전제로 한다.",
         "★ 20%는 정책 목표, ★★★ 75%는 \"성능 저하 없는 KV 압축·축출\" 문헌 상한(4x). 그 이상은 정확도 손실과 맞바꾼 영역이다.",
         "상한 75%는 KV 바이트 기준 문헌 값이라 본 지표(AI Data 전체 평균 GiB)로 직접 환산해야 한다. 저비트 값은 스케일·outlier 오버헤드를 뺀 이론 환산. 기존 v3는 ★★★ = 20%로 정의했음.",
         "출처: arXiv 2505.23416(KVzip, NeurIPS'25), 2401.18079(KVQuant), 2402.02750(KIVI), 2309.06180(vLLM), 2402.01869(InferCept). 내부 값: doc-mk/Evaluation/unified_stars_check.md",
         "핵심: ★ 20% / ★★ 47.5% / ★★★ 75%. 상한은 KV를 3~4x 줄여도 성능 저하가 negligible이라는 KVzip 보고에 근거한다.\n"
         "근거: PagedAttention이 단편화 낭비를 이미 제거했으므로, Baseline(vLLM) 대비 추가 감소는 offload·압축에서 나와야 한다. InferCept는 유휴 KV 처리의 수요를 보여 준다.\n"
         "의미: 75%를 넘으면 정확도 손실과 교환하는 영역이라 QA-05와 같이 봐야 한다.")

# ================================================================== QA-03
def r03(s):
    base = os.path.join(ROOT, "vllm/distributed/kv_transfer/kv_connector/v1")
    names = [("offloading_connector", "offloading_connector.py"), ("simple_cpu_offload", "simple_cpu_offload_connector.py"),
             ("flexkv_connector", "flexkv_connector.py"), ("lmcache_connector", "lmcache_connector.py"),
             ("hidden_states(예시)", "example_hidden_states_connector.py"), ("decode_bench", "decode_bench_connector.py"),
             ("example_connector", "example_connector.py"), ("multi_connector", "multi_connector.py"),
             ("lmcache_mp_connector", "lmcache_mp_connector.py")]
    files = [(n, len(open(os.path.join(base, f), encoding="utf8").read().splitlines())) for n, f in names]
    pk = [(d, sum(len(open(p, encoding="utf8").read().splitlines()) for p in glob.glob(os.path.join(base, d, "**/*.py"), recursive=True)))
          for d in ("p2p", "mooncake", "offloading", "lmcache_integration", "moriio", "hf3fs", "nixl")]
    pk.sort(key=lambda t: t[1])
    med = statistics.median(v for _, v in files)
    x, y, w, lw, pw = 6.85, 1.2, 6.25, 1.75, 3.6
    mx = 4000.0
    pos = lambda v: x + lw + v / mx * pw
    tbox(s, x, y - 0.02, w, 0.26, ["vLLM 저장소의 KV connector 코드량 (LOC, 이 저장소에서 직접 계수)"], size=10, bold=True)
    rows = files + [None] + pk
    top = y + 0.3
    rh_ = 0.168
    for i, it in enumerate(rows):
        if it is None:
            continue
        n, v = it
        yy = top + i * rh_
        pkg = i > len(files)
        tbox(s, x, yy - 0.01, lw - 0.05, rh_, [(n + (" (패키지)" if pkg else ""), {"align": PP_ALIGN.RIGHT})], size=7.5, color=G40, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
        shape(s, MSO_SHAPE.RECTANGLE, pos(0), yy + 0.025, max(pos(v) - pos(0), 0.03), rh_ - 0.05, ORG if pkg else BLUE)
        tbox(s, pos(v) + 0.04, yy - 0.01, 0.8, rh_, [f"{v:,}"], size=7.5, bold=True, color=ORG if pkg else BLUE, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    bot = top + len(rows) * rh_
    vline(s, pos(med), top - 0.04, top + len(files) * rh_, RED, 1.0, dash=True)
    tbox(s, pos(med) + 0.05, top + len(files) * rh_ - 0.02, 2.4, rh_, [f"최상위 connector 파일 중앙값 {int(med)} LOC"], size=7.5, bold=True, color=RED, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    hline(s, pos(0), pos(mx), bot + 0.02, G7F, 0.75)
    tbox(s, x, bot + 0.05, w, 0.22, ["파랑 = 최상위 connector 파일, 주황 = backend 패키지(하위 디렉터리 합). 파일 단위 LOC는 변경량이 아닌 규모 참고치"], size=7.5, color=G59, margin=0)
    lit_table(s, 6.85, bot + 0.32, [1.45, 2.65, 0.95, 1.2], [
        ["SWE-bench\narXiv 2310.06770", "실제 GitHub issue 해결 패치(2,294건): 평균 수정 파일 1.7개, 수정 줄 32.8", "1.7 file\n32.8 line", "module ≤ 2 (★★★)"],
        ["Claude Code 비용 문서\ncode.claude.com", "기업 평균 약 $13/개발자·활동일, 90%가 $30/일 미만", "$13/일", "$2 ≈ 일 평균의 15%"]],
        [0.4, 0.4])

qa_slide(4, "QA-03", "Modifiability",
         "새 tier·정책·객체를 추가하는 변경 1건의 비용이 얼마 이하일 때 ★ / ★★ / ★★★ 인가 (AND 조건)",
         {"colw": [0.85, 1.05, 0.9, 0.9, 2.7], "header": ["별", "토큰 비용*", "공수(MM)", "변경 module(보조)", "근거"],
          "rows": [["☆ 미충족", "> 0.7M개\n(> $2)", "> 1.0", "≥ 6 또는\nmajor IF 변경", "한 항목이라도 ★ 기준을 넘으면 미충족"],
                   ["★", "≤ 0.7M개\n($2)", "≤ 1.0", "≤ 5", "시나리오 목표(하한). 1 MM = 1인 1개월(21 man-day) 로드맵 항목 수준. $2 ≈ 개발자 일 평균 비용의 15%"],
                   ["★★", "≤ 0.4M개\n($1.25)", "≤ 0.75", "≤ 3\n(3.5 내림)", "★과 ★★★의 등분 중점"],
                   ["★★★", "≤ 0.17M개\n($0.5)", "≤ 0.5", "≤ 2", "상한. 0.5 MM = 1인 2주 sprint. 일상 변경 규모(SWE-bench 평균 1.7 file)와 같은 범위. $0.5 ≈ 일 평균의 4%"]],
          "heights": [0.55, 0.85, 0.5, 0.85], "size": 9.5},
         r03,
         "지표: 변경 시나리오 4종(신규 tier·정책 교체·객체 유형 추가·vLLM upstream 갱신 대응)을 구현해 시나리오마다 (M3) 코드 에이전트 토큰 비용, (M2) 개발 공수 MM(1 MM = 21 man-day), (M1) 변경 module 수를 측정한다. 세 값이 모두 그 별의 기준을 만족해야 그 별(AND). *토큰 환산: 입력 단가 $3/MTok 가정(2차 자료).",
         "★★★은 \"일상적인 변경 한 건\" 규모(≤ 2 module, 2주)로 새 tier·정책을 붙이는 것. 현재 vLLM에서 새 backend는 수백~수천 LOC다.",
         "공수·비용은 구현 측정 × 가정 상수 [B+C]이며 실제 에이전트 세션 측정이 아니다. LOC는 변경 난이도의 대리 지표일 뿐 Baseline 변경량이 아님. module은 셀(표)에는 표기하지 않은 보조 지표.",
         "출처: arXiv 2310.06770(SWE-bench, 통계는 논문 Table), code.claude.com/docs/en/costs, 코드량은 vllm/distributed/kv_transfer/kv_connector/v1 직접 계수. 경계값 이력: doc-mk/Evaluation/qa-star-criteria-unified.md §3.4",
         "핵심: 토큰 비용·공수·module 세 지표를 모두 만족해야 그 별이다. ★ $2·1 MM·≤5 / ★★ 등분 / ★★★ $0.5·0.5 MM·≤2.\n"
         "근거: SWE-bench의 평균 변경은 파일 1.7개·32.8줄이라 ≤2 module이 일상 변경의 하한에 가깝다. Claude Code 일 평균 비용 $13 대비 $0.5~2는 4~15%다.\n"
         "의미: 현재 vLLM의 신규 backend는 1.4~3.8천 LOC 패키지이므로 plugin 인터페이스로 수백 LOC 수준에 들어오는 것이 설계 목표다.")

# ================================================================== QA-04
def r04(s):
    items = [("arXiv 2409.16053, 65,536 GPU", 80, 80, "real", "80% (학습)"),
             ("내부 DP4 후보1 (시뮬, N=16)", 93, 93, "int", "93%"),
             ("내부 DP4 후보2 (시뮬, N=16)", 64, 64, "int", "64%"),
             ("USL 환산 σ=1.67% (N=16)", 80, 80, "calc", "σ 1.67%"),
             ("USL 환산 σ=1.18%", 85, 85, "calc", "σ 1.18%"),
             ("USL 환산 σ=0.74%", 90, 90, "calc", "σ 0.74%")]
    yb = strip(s, 6.85, 1.2, 6.25, 50, 100, [50, 60, 70, 80, 90, 100],
               [(80, "★ 80%"), (85, "★★ 85%"), (90, "★★★ 90%")], items,
               "Scaling Efficiency SE(N=16) (%)", lambda v: f"{v:g}%", rowh=0.2)
    lit_table(s, 6.85, yb, [1.45, 2.65, 0.95, 1.2], [
        ["arXiv 2409.16053", "AMD Frontier 65,536 GPU 병렬 효율(학습, weak scaling)", "80%", "★ 근거(유추)"],
        ["Amdahl(1967)\nGunther USL(2008)", "SE(N) = 1 / (1 + σ(N−1)), 조율·직렬 비중 σ. N=16에서 80% ⇒ σ ≤ 1.67%, 90% ⇒ σ ≤ 0.74%", "σ 환산", "★★★ 상한 환산"],
        ["HPC 보고 관행\n(정책 선택)", "90% 이상이면 \"효율적\"으로 보고하는 관행. 이론 상한은 100% (요청 독립 replica는 조율 비용만 손실)", "90%", "★★★ 정책"]],
        [0.4, 0.6, 0.6])

qa_slide(5, "QA-04", "Scalability",
         "노드를 N배로 늘릴 때 처리량이 선형 대비 몇 % 유지되면 ★ / ★★ / ★★★ 인가",
         {"colw": [0.85, 1.55, 4.0], "header": ["별", "조건 (SE(N=16) = Goodput(N) ÷ ((N/N0) · Goodput(N0)), N0=2)", "근거"],
          "rows": [["☆ 미충족", "SE ≤ 80%", "증설 비용의 20% 초과가 조율 계층에서 손실. 시나리오 목표 미달"],
                   ["★", "SE > 80%", "시나리오 목표(하한). 조율 직렬 비중 σ ≤ 1.67%. 65,536 GPU 학습 병렬 효율 80% 보고와 같은 수준"],
                   ["★★", "SE > 85%", "★(80%)과 ★★★(90%)의 등분 중점. σ ≤ 1.18%"],
                   ["★★★", "SE > 90%", "상한. σ ≤ 0.74%(결정·라우팅 계층이 요청 처리 시간의 1% 미만). 병렬 효율 \"효율적\" 보고 관행. 이론 상한 100%"]],
          "heights": [0.5, 0.78, 0.55, 0.92], "size": 10},
         r04,
         "지표: N0 = 2 노드 기준 SE(N) = Goodput(N) ÷ ((N / N0) × Goodput(N0)), N_max = 16. N > 2는 모델 외삽(시뮬레이션, 제약 C04). 별은 보완 설계(샤딩 등) 적용 전 기준. 확장성은 5개 증가 축(S1~S5)으로 정의되며 별 경계는 S2(노드 수) 축 SE에만 있고 나머지는 값 보고(Evaluation/qa-star-criteria-unified.md §3.6, D9).",
         "★ 80%는 대규모 병렬 보고 수준, ★★★ 90%는 조율 직렬 비중 0.74% 이하의 환산 값이다. 이론 상한은 100%.",
         "문헌 80%는 학습(weak scaling) 보고라 서빙 SE와 의미가 다름 [유추]. 90%는 문헌 직접 확인 실패 → 정책 선택. 이전 문서가 인용한 arXiv 1704.03329는 분자동역학 코드 생성 논문으로 근거 부적합(교체).",
         "출처: arXiv 2409.16053(abstract), Amdahl 1967, Gunther 2008(USL, 원리). Cornell CVW 페이지는 90% 기준을 확인하지 못함. 내부 값: doc-mk/Evaluation/unified_stars_check.md",
         "핵심: ★ 80% / ★★ 85% / ★★★ 90%. 같은 구간을 조율 직렬 비중 σ 1.67% / 1.18% / 0.74%로 환산해 심사에서 \"결정 계층의 비용 예산\"으로 설명한다.\n"
         "근거: 65,536 GPU 학습 병렬 효율 80% 보고가 하한의 유추 근거다. 90%는 정책 선택임을 명시한다.\n"
         "주의: 이전 문서의 arXiv 1704.03329 인용은 부적합해 제거한다.")

# ================================================================== QA-05
def r05(s):
    items = [("무손실 경로 (offload·전송·prefix 재사용)", 0, 0, "calc", "0% (bit-exact)"),
             ("소유자 허용 한도 (2026-10-09)", 1, 1, "calc", "1%"),
             ("내부 DP3 C2 (실측 매핑)", 0, 0, "int", "-1.99% (개선)"),
             ("내부 DP3 C1 (손실 압축 설정)", 5, 16.98, "int", "16.98%")]
    yb = strip(s, 6.85, 1.2, 6.25, 0, 5, [0, 1, 2, 3, 4, 5],
               [(1, "★★★ 1%"), (2, "★★ 2%"), (3, "★ 3%")], items,
               "F1 상대 하락률 (압축·재사용 전 full recompute 대비, %)", lambda v: f"{v:g}%", rowh=0.2)
    lit_table(s, 6.85, yb, [1.45, 2.65, 0.95, 1.2], [
        ["KVzip\narXiv 2505.23416", "KV 3~4x 축소, QA·검색·추론·코드 이해에서 \"negligible performance loss\"", "정성", "★★★ 가능성"],
        ["KVQuant\narXiv 2401.18079", "3-bit KV 양자화, perplexity 저하 < 0.1 (Wikitext-2, C4)", "ppl\n< 0.1", "F1 아님(참고)"],
        ["KIVI\narXiv 2402.02750", "2-bit KV 양자화, 품질 \"거의 동일\" (수치는 본문 확인 필요)", "정성", "참고"],
        ["소유자 결정\n(내부)", "압축 전 + full recompute 대비 F1 1% 이내(상대 비율) 허용 한도 확정", "1%", "★★★ 상한"],
        ["부재", "KV 압축의 F1 하락에 대한 일반 임계 문헌 없음 → 3%는 한도의 3배(정책)", "—", "★ 하한"]],
        [0.4, 0.4, 0.4, 0.4, 0.4])

qa_slide(6, "QA-05", "Functional Correctness",
         "LLM 출력 F1이 Baseline 대비 몇 % 이내로 유지되면 ★ / ★★ / ★★★ 인가",
         {"colw": [0.85, 1.55, 4.0], "header": ["별", "조건 (F1 상대 하락률 = (Baseline − 후보) ÷ Baseline)", "근거"],
          "rows": [["☆ 미충족", "하락 ≥ 3%", "시나리오 목표 미달. 허용 한도의 3배 이상"],
                   ["★", "하락 < 3%", "시나리오 목표(하한). 허용 한도 1%의 3배, 눈에 띄는 열화가 시작되는 구간 [정책]"],
                   ["★★", "하락 < 2%", "★(3%)과 ★★★(1%)의 등분 중점"],
                   ["★★★", "하락 < 1%", "상한. 소유자가 정한 허용 정확도 저하 한도. 손실 압축의 문헌 \"negligible\"과 같은 방향. 무손실 경로는 0%(bit-exact)"]],
          "heights": [0.5, 0.72, 0.55, 0.85]},
         r05,
         "지표: 같은 task·dataset·모델·precision에서 압축·재사용 전 + full recompute의 F1 대비 후보 F1의 상대 하락률. 데이터 이동·복원·재사용만 거친 경로(무손실)는 출력 bit-exact 확인으로 갈음한다.",
         "★★★ 1%는 소유자 허용 한도(상한), ★ 3%는 그 3배. 무손실 경로(offload·전송·prefix 재사용)는 이론 상한 0%다.",
         "F1 하락의 일반 임계 문헌이 없고, 문헌은 perplexity·정성 표현이라 F1로 직접 환산 불가 → 소유자 한도 + 정책. 현재 DP3 C1(손실 압축) F1 하락 16.98%는 3% 초과로 미충족. 3% 경계는 task별 run-to-run F1 변동으로 확인 필요.",
         "출처: arXiv 2505.23416(KVzip), 2401.18079(KVQuant), 2402.02750(KIVI). 내부 값: doc-mk/Evaluation/unified_stars_check.md, 소유자 결정: doc-mk/Requirements/qa-status.md §4",
         "핵심: ★ 3% / ★★ 2% / ★★★ 1%. 1%는 소유자 확정 한도라 상한으로 쓰고 3%는 그 3배로 둔다.\n"
         "근거: KVzip·KVQuant·KIVI는 손실 압축에서도 품질 저하가 작다고 보고하지만 F1 정량값은 없다.\n"
         "주의: 문헌 수치로 3%를 뒷받침하지 못한다. 이 경계는 정책 선택이며 소유자 확인이 필요하다.")

# ---- 1장 링크 아이콘 → 상세 슬라이드
for i, ic in enumerate(LINK_ICONS):
    ic.click_action.target_slide = prs.slides[i + 1]

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
