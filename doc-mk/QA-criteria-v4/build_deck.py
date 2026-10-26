"""첨부 덱(base/DP-QA-attached.pptx)을 베이스로 QA 평가 기준 덱 v4를 만든다 (판정 Baseline = B1 정적 offload).

기존 자료는 건드리지 않고 이 폴더에만 쓴다.  python3 build_deck.py  →  DP-QA-criteria-v4.pptx
"""
import copy
import math
import re

from pptx import Presentation
from pptx.util import Inches, Pt

import criteria_data as cd
import theory_model as tm
from deck_lib import *  # noqa: F401,F403  (헬퍼: tbox, rect, add_table, band_chart, criteria_page ...)
from deck_lib import _font, BAND_FILL, NAVY, BASE, OUT


_l0 = tm.max_lambda(tm.B0)
_n0, _h0 = [tm.latencies(tm.B0, f * _l0, tm.max_lambda(tm.B0))[0] for f in (0.5, 0.9)]
_n1, _h1 = [tm.latencies(tm.B1, f * _l0, tm.max_lambda(tm.B1))[0] for f in (0.5, 0.9)]
_LOADROW = (_n0, _h0, _h0 / _n0, _n1, _h1)


def q4_spec():
    r = cd.BAND_TEXT["Q4"]
    return dict(
        title="시스템 확장성 - 평가 기준",
        subtitle="시스템 및 자원의 증가분 대비 처리량의 증가분 비율 — ★2는 경쟁 기술(노드 독립 단순 offloading)이 보이는 '효율적' 수준",
        src="출처: AthenaK (arXiv 2409.16053, 학습 workload) · USL 환산. 90% '효율적' 기준은 관례(정책)이며 인용 논문을 확인하지 못해 출처로 삼지 않았다",
        metric=["[Metric] SE(N) = Goodput 증가 배율 ÷ N 증가 배율  (N = 노드 수, 평가는 N=16 시뮬레이션)",
                "USL 환산: SE(N) = 1 ÷ (1 + σ(N−1)), σ = 조율 계층의 직렬 비중. N=16에서 SE 80%↔σ 1.67%, 85%↔1.18%, 90%↔0.74%, 92.5%↔0.54%, 95%↔0.35%",
                "Baseline 절대값(이론): B0·B1 모두 노드 독립이라 조율 비용이 없다 → SE 98%(부하 불균형 2% [가정]). 기준은 이상적 선형 확장 대비라 Baseline이 B0든 경쟁 기술(B1)이든 구간은 같다. 단 노드 독립 구조는 노드 간 공유 이득이 없다",
                "구간 규칙: ★1 = Stakeholder 80% · ★2 = 90%(경쟁 기술 수준: 병렬 효율에서 '효율적'이라 부르는 관례, 손실 10% 이내) · ★3 = 95%(조율 계층을 두고도 손실 5% 이내) · 반 단계는 중점"],
        metric_h=1.05,
        table_y=2.3,
        rows=[
            (3, r[3], "경쟁 기술 위의 우리 목표: 노드 간 KV 공유·조율 계층을 두고도 처리량 손실 5% 이내(16노드에서 15.2노드분 효과). 조율 계층 직렬 비중 σ ≤ 0.35%(USL 환산, 정책 선택)", "USL 환산 / 정책"),
            (2.5, r[2.5], "★2와 ★3의 중점(손실 7.5%)", "—"),
            (2, r[2], "경쟁 기술 수준: 처리량 손실 10% 이내(16노드에서 14.4노드분 효과), '효율적'이라 부르는 관례. 노드 독립 단순 offloading은 조율 비용이 없어 SE 98%지만 노드 간 공유 이득이 없다", "정책 선택, 이론 계산"),
            (1.5, r[1.5], "★1과 ★2의 중점(손실 15%)", "—"),
            (1, r[1], "Stakeholder 요구: SE > 80%. 문헌: 65,536 GPU에서 병렬 효율 80% 보고(AthenaK, 학습 workload라 서빙 SE와 의미가 다름)", "Stakeholder 요구사항, arXiv 2409.16053"),
            (0, r[0], "확장은 되지만 조율 계층이 병목. 시나리오 요구(80%) 미충족", "—"),
        ],
        chart=dict(
            title="구간 대비 근거 위치 (SE %, N=16)",
            axis=dict(min=50, max=100, log=False, ticks=[50, 60, 70, 80, 90, 100], unit="%"),
            bands=[(50, 80, 0), (80, 85, 1), (85, 90, 1.5), (90, 92.5, 2), (92.5, 95, 2.5), (95, 100, 3)],
            markers=[
                dict(kind="point", v=80, style="req", label="Stakeholder 요구: SE > 80%"),
                dict(kind="point", v=80, style="lit", label="AthenaK: 65,536 GPU에서 병렬 효율 80% (학습)"),
                dict(kind="point", v=90, style="calc", label="'효율적' 보고 관례 90% (정책, 손실 10% 이내)"),
                dict(kind="range", v=(80, 95), style="calc", label="USL 환산 σ: 80%↔1.67% · 90%↔0.74% · 95%↔0.35%"),
            ]),
    )


def q5_rows():
    mm = cd.mm_of_modules
    return [
        (3, "≤ 2", "≤ %.1f" % mm(2), "component 12~14개 중 2개 이하(약 1/7)만 바뀜: 변경이 한두 곳에서 끝나는 최선의 경우. 2개 ≈ 0.4 MM(1인 2주 미만). 참고: SWE-bench 평균 패치 1.7 file(단위가 file이라 module과 다름)", "설계 정책값 · SWE-bench (arXiv 2310.06770, 참고용)"),
        (2.5, "≤ 2.5", "≤ %.1f" % mm(2.5), "★2와 ★3의 중점", "—"),
        (2, "≤ 3", "≤ %.1f" % mm(3), "component 약 1/4(3개)까지: 변경이 한 Plane 또는 한 Layer 안(예: 정책 계층만)에서 끝남. 3개 ≈ 0.6 MM(1인 약 2.5주)", "설계 정책값"),
        (1.5, "≤ 4", "≤ %.1f" % mm(4), "★1과 ★2의 중점", "—"),
        (1, "≤ 5", "≤ %.1f" % mm(5),
         "최소 충족선. component 약 1/3(5개)까지만 바뀌면 변경이 설계 일부에 국한. 6개 이상이면 설계 전반으로 퍼져 확장 지점이 없다고 판단. 5개 ≈ 1.0 MM(1인 1개월 상한)",
         "설계 정책값 (DP1 component 목록)"),
        (0, "≥ 6 또는 major interface 변경", "> 1.0", "6개 이상이면 미충족(설계 전반으로 번짐). 기존 필드 의미 변경(Backend I/F, Event schema)은 module 수와 무관하게 미충족", "—"),
    ]


def q5_page(slide):
    page_chrome(slide, "변경 용이성 - 평가 기준", "별 조건 = 변경 module 수 하나 (시나리오 평균) — 개발 공수 M/M은 module 수에서 환산한 값으로 병기",
                "출처: SWE-bench (arXiv 2310.06770) · vLLM 저장소 kv_connector/v1 직접 계수(최상위 파일 192~1,190 LOC, 중앙값 356) · DP1 qa4-preregistration(공식·상수 ASSUMED)")
    for nm in ("tbl", "formula", "l"):
        for sh in [s for s in slide.shapes if s.name == nm]:
            remove(sh)
    # 변경 시나리오 박스
    tbox(slide, 0.4, 1.17, 12.5, 0.5,
         [[("[변경 시나리오 3종, KV 캐시 관리 범위] ", {"bold": True, "color": NAVY, "size": 10}),
           ("S1 새 메모리 tier 추가 · S2 이동·배치·재사용 정책 교체 · S3 vLLM upstream 갱신 대응(block pool, kv_offload, CUDA graph, async scheduling 접점 수정). "
            "시나리오 단순 평균으로 판정. module = 설계 문서의 component 단위(한 줄이라도 바꾸면 1개).", {"size": 9.5})]],
         size=9.5, line_spacing=1.1, name="metric")
    rows = [["별", "변경 module 수 (평균) = 판정 기준", "환산 개발 공수 (M/M, 참고)", "근거", "출처"]]
    stars = []
    for st, m, mm, ev, src in q5_rows():
        rows.append([[[(star_label(st), {"bold": True})]], m, mm, ev, src])
        stars.append(BAND_FILL[st])
    colw = [0.7, 1.5, 1.5, 6.9, 1.93]
    rowh = [0.3] + [0.32] * 6
    add_table(slide, 0.4, 1.7, colw, rows, rowh, fills=stars, size=8.3, aligns=["c", "c", "c", "l", "l"])
    # 환산 박스
    rect(slide, 0.4, 4.12, 12.53, 0.86, fill="F2F6FC", line="2F5597", line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="formula")
    tbox(slide, 0.5, 4.14, 12.3, 0.84,
         [[("module 1개 = 고정 3일 + 24 LOC ÷ 40 LOC/일 × 테스트·리뷰 overhead 2.0 = 4.2일 = 0.2 MM  →  M/M = 0.2 × module 수  (21일/MM)", {"bold": True, "color": NAVY, "size": 9.5})],
          [("24 LOC/module: SWE-bench 평균 패치 32.8 line ÷ 1.7 file ≈ 19 LOC/file을 올림. 실측 대조: DP1 후보 구현 측정 1.75 module → 0.38 MM(환산 0.35), 2.5 module → 0.51 MM(환산 0.50). "
            "LOC는 실제 vLLM 통합 기준(시뮬레이터 proxy LOC는 ×5 환산).", {"size": 8.5, "color": "404040"})],
          [("출처: 고정 3일(interface 파악·설계 리뷰·통합 테스트)과 40 LOC/일(순 생산성)은 문헌값이 아니라 DP1 사전 등록 가정(ASSUMED, 범위 2~5일 · 25~80 LOC/일 · overhead 1.5~3.0). "
            "범위 양 끝이면 module 1개 = 0.12~0.38 MM, 5개 = 0.6~1.9 MM이라 M/M 환산은 약 ±50% 불확실하다(module 수 판정은 영향 없음).", {"size": 8.3, "color": "C00000"})]],
         size=9, line_spacing=1.1, name="formula_text")
    # 왜 하나만 보는가
    tbox(slide, 0.4, 5.02, 12.5, 0.3,
         [[("왜 module 수만 판정하나: ", {"bold": True, "color": NAVY, "size": 9}),
           ("M/M이 module 수에서 계산되므로 둘은 독립이 아니다. 둘을 AND로 걸면 같은 정보를 두 번 세는 것이다. 판정은 module 수 하나로 하고, 각 경계가 몇 M/M인지는 환산해 보인다. "
            "토큰 비용은 별점이 아니라 별도 참고 페이지(다음 장)에 둔다.", {"size": 8.5})]],
         size=8.5, line_spacing=1.05)
    return 0



# ============================================================ SLO 페이지 (슬라이드 9)
def budget_bar(slide, x, y, w, label, segs, slo, unit, scale_max, note, stacked=True):
    """segs=[(값, 색, 라벨)]. stacked면 누적 막대 1줄, 아니면 항목별 막대. SLO 기준선은 붉은 세로선."""
    tbox(slide, x, y, w, 0.2, [[(label, {"bold": True, "color": NAVY})]], size=9, margins=(0.02, 0, 0, 0))
    tx, tw, th = x + 0.02, w - 0.04, 0.24
    n = 1 if stacked else len(segs)
    top = y + 0.28
    for i in range(n):
        rect(slide, tx, top + i * (th + 0.05), tw, th, fill="F2F2F2", line="BFBFBF", line_w=0.5)
    cur = 0.0
    for i, (val, col, lab) in enumerate(segs):
        row = 0 if stacked else i
        sx = tx + (cur if stacked else 0) / scale_max * tw
        sw = val / scale_max * tw
        ty = top + row * (th + 0.05)
        rect(slide, sx, ty, sw, th, fill=col, line="FFFFFF", line_w=0.75)
        tbox(slide, sx, ty, sw, th, lab, size=7.5, color="FFFFFF" if col in ("2F5597", "7F7F7F") else "000000",
             align="c", anchor="m", margins=(0, 0, 0, 0), wrap=False)
        if stacked:
            cur += val
    bottom = top + n * (th + 0.05)
    lx = tx + slo / scale_max * tw
    hline(slide, lx, top - 0.04, lx, bottom, color="C00000", w=1.75)
    tbox(slide, lx - 1.4, y, 1.38, 0.2, "SLO %g %s" % (slo, unit), size=8, color="C00000", align="r",
         margins=(0, 0, 0, 0), wrap=False)
    tbox(slide, tx, bottom + 0.02, tw, 0.34, note, size=7.5, color="404040", margins=(0, 0, 0, 0))


def slo_page(slide):
    page_chrome(slide, "SLO 기준", "TTFT·TPOT의 측정 정의, 절대 기준, 근거 계산과 적용 규칙 (모든 DP·후보·Baseline에 동일 적용)",
                "출처: 공통 SLO 기준 · MLPerf Inference(2차) · Brysbaert 2019. 계산: 가중치 141 GB, H100×8 3.35 TB/s·989 TFLOPS, B200×8 8 TB/s, MFU 0.45. 하한은 attention 2차 항·통신·커널 제외")
    for sh in [s for s in slide.shapes if s.name == "tbl"]:
        remove(sh)
    rows = [["Metric", "측정 정의", "절대 SLO", "근거 (계산)"],
            ["TTFT\n(첫 토큰까지의 시간)",
             "요청 도착 → 첫 출력 토큰 수신. 큐잉 + prefill + (필요 시) KV 복원 시간 포함. 요청별 값의 P99로 판정",
             [[("TTFT P99 ≤ 2 s", {"bold": True})]],
             ["8K prefill 이론 0.32 s = (2 × 70.6B × 8,192 = 1.16 PFLOP) ÷ (H100×8 989 TFLOPS×8 × MFU 0.45 = 3.56 PFLOP/s). 2 s는 이의 약 6배(큐잉 여유).",
              "Agent KV Fleet의 16K 컨텍스트 재prefill은 0.65 s라 약 3배 여유. MLPerf Llama-2-70B 표준 시나리오와 같은 자릿수(2차 자료)"]],
            ["TPOT\n(출력 토큰당 시간)",
             "(요청 E2E − TTFT) ÷ (출력 토큰 수 − 1). 토큰 간 지연의 요청별 평균에 대한 P99로 판정",
             [[("TPOT P99 ≤ 50 ms", {"bold": True})]],
             ["50 ms = 20 tok/s = 성인 묵독(5~6 tok/s)의 약 3.5배. MLPerf interactive 40 ms와 표준 200 ms 사이(2차 자료).",
              "이론 하한: 가중치 141 GB를 step마다 읽으면 H100×8 5.3 ms, B200×8 2.2 ms. 16K KV(5 GiB/세션) batch 64를 더하면 H100 약 18 ms. 50 ms는 이의 약 2.8배 여유"]]]
    add_table(slide, 0.4, 1.25, [1.5, 3.3, 1.5, 6.23], rows, [0.26, 0.82, 0.82], fills=[None, None], size=8.3,
              aligns=["l", "l", "c", "l"])
    # 예산 막대 (좌: TTFT, 우: TPOT)
    tbox(slide, 0.4, 3.3, 12.5, 0.24, [[("SLO 대비 이론 예산 (대역폭·연산 한계 기준 하한이며 실측이 아니다)", {"bold": True, "color": NAVY})]],
         size=10, margins=(0.02, 0, 0, 0))
    budget_bar(slide, 0.4, 3.62, 6.0, "TTFT 예산 (s)", [(0.32, "2F5597", "8K 0.32 s"), (0.65, "7F7F7F", "16K 0.65 s")],
               2.0, "s", 2.4, "8K·16K 입력의 재계산 prefill 시간. 남는 시간이 큐잉·KV 복원·결정 지연의 예산이다(16K에서 1.35 s).", stacked=False)
    budget_bar(slide, 6.9, 3.62, 6.0, "TPOT 예산 (ms, H100×8)", [(5.3, "2F5597", "가중치 5.3"), (12.8, "7F7F7F", "KV 12.8 (batch 64, 16K)")],
               50, "ms", 60, "B200×8은 가중치 하한이 2.2 ms. 오타였던 TPOT 2 ms는 이 모델에서 가중치 읽기 하한보다 낮아 달성 불가.", stacked=True)
    # 적용 규칙
    tbox(slide, 0.4, 4.95, 12.5, 2.2,
         [[("적용 규칙", {"bold": True, "color": NAVY, "size": 10})],
          [("• Q1 SLO goodput: 요청별로 TTFT ≤ 2 s 그리고 TPOT ≤ 50 ms를 함께 만족한 요청의 초당 처리 수(query/s)로 센다(1 query = 요청 1건 = 1 turn, 출력 256 token). 부하 sweep의 최댓값이 Max SLO goodput이다.", {"size": 9})],
          [("• Q3 고부하 지연: 부하는 Baseline Max SLO goodput의 90%로 고정하고 TTFT·TPOT P99로 r을 계산한다. 후보가 SLO(P99)를 넘으면 개선 폭과 무관하게 ☆로 본다.", {"size": 9})],
          [("• 같은 값을 모든 DP 후보와 Baseline에 적용하고, 결과를 본 뒤 바꾸지 않는다. 바꾸면 새 버전으로 취급해 이전 결과와 비교하지 않는다.", {"size": 9})],
          [("• Baseline이 SLO를 만족하지 못하는 시나리오는 조용히 제외하지 않고 infeasible로 표시한다. P99 판정에는 구간당 요청 수 ≥ 1,000(seed별)을 확보한다 [가정].", {"size": 9})],
          [("• 이전 요구에 적힌 TPOT 2 ms는 50 ms의 오타로 확정해 이 덱의 모든 페이지에 50 ms를 쓴다 (소유자 확인 필요).", {"size": 9, "color": "C00000"})]],
         size=9, line_spacing=1.15)



# ============================================================ 이론 값 가져오기
_R = tm.all_results()
SGL = _R["single"]                    # 단일 노드(B0, B1, DP1, DP3)
B0_QPS, B1_QPS = SGL["B0"]["qps"], SGL["B1"]["qps"]
B0_TTFT, B0_TPOT = SGL["B0"]["ttft"], SGL["B0"]["tpot"]
B1_TTFT, B1_TPOT = SGL["B1"]["ttft"], SGL["B1"]["tpot"]
HBM_B0 = SGL["B0"]["hbm"]
CEIL_RATIO = (1 / tm.G0) / SGL["B1"]["lam"]                   # GPU 한계 ÷ 경쟁 기술(B1)
HBM_FLOOR_GIB = 53                                             # decode 상주분 + 버퍼(이론 하한)
COMP = cd.competitor_stars()


def n0(x):
    return format(round(x), ",")


def q1_spec():
    e = cd.EDGES["Q1"]
    ab = [cd.fmt_q(x * B0_QPS) for x in e]
    r = cd.BAND_TEXT["Q1"]
    vals = sorted(v for _, v in cd.Q1_LIT)
    import statistics as _st
    med = _st.median(vals)

    def rng(star, i):
        return [[(r[star], {})], [("(≥ %s query/s)" % ab[i], {"size": 7.5, "color": "595959"})]]
    ceil_r0 = B1_QPS / B0_QPS * CEIL_RATIO
    return dict(
        title="처리량 - 평가 기준",
        subtitle="Baseline(GPU-only) 대비, SLO를 만족하는 조건에서의 처리량(Max값) 증가율 — ★2는 경쟁 기술(KV offload·재사용) 문헌 대표값",
        src="출처: LMCache 2510.09665 · Mooncake 2407.00079 · InferCept 2402.01869 · Pensieve 2312.05516 · AttentionStore 2403.19708 · DistServe 2401.09670 · Sarathi-Serve 2403.02310. 'up to' 보고값이며 지표·baseline이 달라 직접 비교 아님",
        metric=["[Metric] R = Max SLO goodput(후보) ÷ Max SLO goodput(B0 GPU-only)",
                "B0 = vLLM 기본(HBM만, 부족하면 recompute). 경쟁 기술 = KV offload·재사용·상태 보존 시스템(LMCache, Pensieve, InferCept, Mooncake, AttentionStore 등). 이론 대리 B1 = 단순 offloading(증분 write-through, 동기 복원, prefetch 없음)",
                "Baseline 절대값(이론, 1노드): B0 %s query/s → ★1 ≥ %s · ★2 ≥ %s · ★3 ≥ %s query/s. 이 시나리오의 이론값: B1 ×%.2f(%s query/s), 이론 상한 ×%.1f" % (
                    cd.fmt_q(B0_QPS), ab[0], ab[2], ab[4], B1_QPS / B0_QPS, cd.fmt_q(B1_QPS), ceil_r0),
                "구간 규칙: ★1 = Stakeholder 2x · ★2 = 문헌 8건 headline 개선의 중앙값 %.1fx · ★3 = 문헌 상위권 8x(★2의 2배) · ★ 1단계 = 2배(log 등간격), 반 단계는 기하 중점" % med],
        metric_h=1.05,
        table_y=2.2,
        row_h=0.31,
        rows=[
            (3, rng(3, 4), "문헌 상위권 ≈ 8x: AttentionStore prefill 처리량 7.8~8.2x(TTFT −87%), DistServe 요청 수 7.4x(SLO 내), Sarathi-Serve SLO 내 처리량 최대 6.9x(Falcon-180B). "
                           "★2(4x)의 2배. LMCache 최대 14x는 극단 조건이라 앵커에서 제외", "AttentionStore, DistServe, Sarathi-Serve"),
            (2.5, rng(2.5, 3), "★2와 ★3의 기하 중점(√(4×8)=5.7x). Mooncake 시뮬 +525%(6.25x), LMCache 실서비스 QPS 확장 4~6x의 상단과 Sarathi-Serve 6.9x 사이", "Mooncake, LMCache, Sarathi-Serve"),
            (2, rng(2, 2), "경쟁 기술 대표값 = KV offload·재사용 시스템 8건의 headline 개선(1.75, 2.0, 2.3, 3.0, 5.0, 6.25, 7.8, 14x)의 중앙값 %.1fx(기하평균 4.1x). LMCache 실서비스 QPS 확장 4~6x의 하한과도 일치" % med,
             "InferCept, Mooncake, LMCache, Pensieve, AttentionStore"),
            (1.5, rng(1.5, 1), "★1과 ★2의 기하 중점(√(2×4)=2.8x). Pensieve 최대 3.0x, Sarathi-Serve 2.6x 구간", "Pensieve, Sarathi-Serve"),
            (1, rng(1, 0), "Stakeholder 요구: \"처리량 개선 2배 이상 되어야 제품화 논의 가능\" (GPU-only 대비). 문헌 하위 사분위(≈2.2x: InferCept 2.0x, LMCache 단일 노드 하한 2.3x, Mooncake 실트래픽 1.75x)와 같은 수준",
             "Stakeholder 요구사항, InferCept, LMCache, Mooncake"),
            (0, [[(r[0], {})]], "Stakeholder 요구(2배) 미충족", "—"),
        ],
        chart=dict(
            title="구간 대비 근거 위치 (GPU-only 대비 R, log 축) — 문헌 7편 + 이론",
            axis=dict(min=1.5, max=16, log=True, ticks=[2, 3, 4, 6, 8, 16], unit="x"),
            bands=[(1.5, 2.0, 0), (2.0, 2.8, 1), (2.8, 4.0, 1.5), (4.0, 5.7, 2), (5.7, 8.0, 2.5), (8.0, 16, 3)],
            markers=[
                dict(kind="point", v=2.0, style="req", label="Stakeholder 요구: 처리량 2배 이상 (GPU-only 대비)"),
                dict(kind="range", v=(1.6, 2.0), style="lit", label="InferCept 1.6~2x · Mooncake 실트래픽 +75%(1.75x)"),
                dict(kind="range", v=(1.5, 3.0), style="lit", label="Pensieve 1.14~3.0x (vLLM·TensorRT-LLM 대비)"),
                dict(kind="point", v=2.3, style="lit", label="LMCache 단일 노드 로컬 CPU offload 하한 2.3x"),
                dict(kind="range", v=(4, 6), style="lit", label="LMCache 실서비스 고 QPS: 동일 TTFT 환산 QPS 확장 4~6x"),
                dict(kind="range", v=(2.6, 6.9), style="lit", label="Sarathi-Serve SLO 내 처리량 2.6x(Mistral-7B)~6.9x(Falcon-180B)"),
                dict(kind="point", v=7.4, style="lit", label="DistServe 요청 수 7.4x (SLO 내)"),
                dict(kind="range", v=(7.8, 8.2), style="lit", label="AttentionStore prefill 처리량 7.8~8.2x (TTFT −87~88%)"),
                dict(kind="point", v=6.25, style="excl", label="Mooncake 시뮬 +525% (6.25x, 시뮬)"),
                dict(kind="point", v=14, style="excl", label="LMCache 최대 14x (극단 조건)"),
                dict(kind="range", v=(B1_QPS / B0_QPS, ceil_r0), style="calc", label="이 시나리오 이론: 단순 offload ×%.2f ~ 상한 ×%.1f" % (B1_QPS / B0_QPS, ceil_r0)),
            ]),
        row_h_chart=0.158,
    )


def q2_spec():
    e = cd.EDGES["Q2"]
    ab = [n0(HBM_B0 * (1 - x / 100)) for x in e]
    r = cd.BAND_TEXT["Q2"]

    def rng(star, i):
        return [[(r[star], {})], [("(HBM ≤ %s GiB)" % ab[i], {"size": 7.5, "color": "595959"})]]
    ceil_red = (1 - HBM_FLOOR_GIB / HBM_B0) * 100
    return dict(
        title="자원(HBM) 효율성 - 평가 기준",
        subtitle="Baseline(GPU-only) 대비 HBM 사용량 감소율 (KV 캐시 관리, 압축 제외) — ★2는 오프로딩 문헌 수준",
        src="출처[1]: DirectKV OSDI'26 · HeadInfer (arXiv 2502.12574, 검색 요약 기준) · ShadowKV ICML'25 (arXiv 2410.21465). 40~80% 범위는 사용자 인용(원문 대조 필요). 절대값은 분석 모델",
        metric=["[Metric] 감소율 = (B0 HBM 사용 − 후보 HBM 사용) ÷ B0 HBM 사용  (HBM 사용 = HBM에 올라와 있는 KV 크기의 시간 평균, 같은 부하·SLO)",
                "Baseline 절대값(이론): B0 %s GiB(HBM KV 풀 = 640×0.9 − 가중치 131 − 활성값 20 전체) → ★1 ≤ %s · ★2 ≤ %s · ★3 ≤ %s GiB" % (n0(HBM_B0), ab[0], ab[2], ab[4]),
                "단순 LRU offload(B1)는 HBM을 캐시로 채워 %s GiB(감소 0%%, ☆)다. 문헌의 40%%는 HBM 절감을 목적으로 한 offload(DirectKV 등)의 보고값이라 ★2 앵커로 쓴다. 이론 상한 ≈ %.0f%%" % (n0(COMP_B1_HBM), ceil_red),
                "구간 규칙: ★1 = 시나리오 요구 20% · ★2 = 오프로딩 문헌 수준 40% · ★3 = 잔여 HBM(60%)을 1.5x 더 줄인 값 = 감소율 60% · 반 단계는 중점(10%p)"],
        metric_h=1.15,
        table_y=2.35,
        rows=[
            (3, rng(3, 4), "경쟁 기술이 줄이고 남은 HBM(60%%)을 1.5x 더 줄이면 40%%만 남는다 → 감소율 60%%. 이론 상한(≈%.0f%%)과 HeadInfer 92%%(1M ctx 극단)보다 보수적. ShadowKV(batch 6x)는 key 저랭크 압축이라 제외" % ceil_red, "DirectKV, HeadInfer [1], 이론 계산"),
            (2.5, rng(2.5, 3), "★2와 ★3의 중점", "—"),
            (2, rng(2, 2), "경쟁 기술(오프로딩 계열) 수준: DirectKV는 CPU 메모리 직접 접근으로 GPU 메모리 43% 감소(GH200, 기존 offload 대비). 사용자 인용 '배치·오프로딩 계열 40~80%'의 하한", "DirectKV [1]"),
            (1.5, rng(1.5, 1), "★1과 ★2의 중점", "—"),
            (1, rng(1, 0), "시나리오 요구 20%. 계산: HBM 20% 감소 = 같은 HBM에 KV 1/0.8 = 1.25배 수용(동시 처리 +25% 이상, 이론)", "시나리오 요구 / 이론 계산"),
            (0, [[(r[0], {})]], "시나리오 요구(20%) 미충족. 음수는 후보가 HBM을 더 쓴 경우", "—"),
        ],
        chart=dict(
            title="구간 대비 근거 위치 (HBM 감소율 %, GPU-only 대비)",
            axis=dict(min=0, max=100, log=False, ticks=[0, 20, 40, 60, 80, 100], unit="%"),
            bands=[(0, 20, 0), (20, 30, 1), (30, 40, 1.5), (40, 50, 2), (50, 60, 2.5), (60, 100, 3)],
            markers=[
                dict(kind="point", v=0, style="calc", label="이론 B1(정적 LRU offload) = 0% (HBM을 캐시로 채움)"),
                dict(kind="point", v=20, style="req", label="시나리오 요구 20% (≈ 동시 수용 +25%)"),
                dict(kind="point", v=43, style="lit", label="DirectKV (OSDI'26): 기존 offload 대비 GPU 메모리 43% 감소 = 경쟁 기술 수준"),
                dict(kind="range", v=(40, 80), style="lit", label="오프로딩 계열 보고 범위 40~80% [1]"),
                dict(kind="point", v=ceil_red, style="calc", label="이론 상한 ≈ %.0f%% (decode 상주분만 남김)" % ceil_red),
                dict(kind="point", v=83, style="excl", label="ShadowKV: batch 6x ≈ 83%/seq (key 저랭크 압축 포함)"),
                dict(kind="point", v=92, style="excl", label="HeadInfer: 총 GPU 메모리 92% (1M ctx 단일 요청, 극단)"),
            ]),
        left_note=True,
    )


COMP_B1_HBM = SGL["B1"]["hbm"]


def q3_spec():
    e = cd.EDGES["Q3"]
    tt = [B0_TTFT / (x * x) for x in e]
    r = cd.BAND_TEXT["Q3"]
    reds = [cd.red_of_r(x) for x in e]
    rb1 = COMP["Q3"][0]

    def rng(star, i):
        return [[(r[star] + " (감소율 %.0f%%)" % reds[i], {})], [("(TPOT 동일 시 TTFT ≤ %.2f s)" % tt[i], {"size": 7.5, "color": "595959"})]]
    return dict(
        title="고부하 지연시간(Tail Latency) - 평가 기준",
        subtitle="Baseline(GPU-only) 대비 TTFT·TPOT P99 개선 r (고부하 = B0 포화 부하의 90%, SLO 충족 전제) — ★2는 경쟁 기술 수준",
        src="출처: LMCache (arXiv 2510.09665, 검색 요약 기준) · KVFlow (arXiv 2507.07400) · DirectKV OSDI'26. LMCache E2E 50~75%는 사용자 인용(원문 대조 필요). 절대값은 분석 모델",
        metric=["[Metric] r = √[ (B0 TTFT P99 ÷ 후보 TTFT P99) × (B0 TPOT P99 ÷ 후보 TPOT P99) ],  감소율 = 1 − 1/r",
                "고부하: 부하를 B0 포화 부하(%.2f turn/s)의 90%%(%.2f)에 고정. Baseline 절대값(이론): B0 TTFT P99 %.2f s · TPOT P99 %.1f ms (이 부하에서 B0는 TTFT SLO 초과). TPOT이 같으면 TTFT가 ★1 ≤ %.2f · ★2 ≤ %.2f · ★3 ≤ %.2f s" % (
                    SGL["B0"]["lam"], cd.LAM_TEST, B0_TTFT, B0_TPOT, tt[0], tt[2], tt[4]),
                "경쟁 기술의 이론값(B1): TTFT %.2f s · TPOT %.1f ms → r %.2f (★2 구간). 후보는 SLO(TTFT ≤ 2 s, TPOT ≤ 50 ms) 충족 필수" % (B1_TTFT, B1_TPOT, rb1),
                "구간 규칙: ★1 = Stakeholder 체감 30%(r 1.43) · ★2 = 경쟁 기술 수준 50%(r 2.0) · ★3 = ★2 × 1.5 = r 3.0(감소율 67%) · 반 단계는 기하 중점"],
        metric_h=1.15,
        table_y=2.3,
        rows=[
            (3, rng(3, 4), "경쟁 기술(r 2.0)의 1.5x(offload baseline 대비 문헌 개선 1.2~1.83x의 중간). LMCache 고 QPS 실트래픽 TTFT 4.4~6.6x·ITL 34~58%를 r로 환산하면 2.6~4.0이라 그 중간~상단에 해당", "LMCache, KVFlow, DirectKV"),
            (2.5, rng(2.5, 3), "★2와 ★3의 기하 중점(√(2.0×3.0)=2.45)", "—"),
            (2, rng(2, 2), "경쟁 기술 수준: LMCache Context 벤치마크 E2E 지연 약 50~75%% 감소의 하한(r 2.0). 이론 B1 r %.2f가 이 구간에 놓인다" % rb1, "LMCache"),
            (1.5, rng(1.5, 1), "★1과 ★2의 기하 중점(√(1.43×2.0)=1.7)", "—"),
            (1, rng(1, 0), "Stakeholder 요구: 사용자 체감 성능 개선 기준 30% 감소(r 1.43, GPU-only 대비)", "Stakeholder 요구사항"),
            (0, [[(r[0], {})]], "r < 1.43 또는 SLO 위반. 음수는 B0보다 나빠진 경우", "—"),
        ],
        chart=dict(
            title="구간 대비 근거 위치 (P99 개선 r, GPU-only 대비)",
            axis=dict(min=1.0, max=4.5, log=False, ticks=[1.0, 1.5, 2.0, 2.5, 3.0, 3.5, 4.0, 4.5], unit="x"),
            bands=[(1.0, 1.43, 0), (1.43, 1.7, 1), (1.7, 2.0, 1.5), (2.0, 2.45, 2), (2.45, 3.0, 2.5), (3.0, 4.5, 3)],
            markers=[
                dict(kind="point", v=1.43, style="req", label="Stakeholder: 체감 30% 감소 (r 1.43)"),
                dict(kind="point", v=rb1, style="calc", label="이론 B1(단순 offload) r %.2f" % rb1),
                dict(kind="range", v=(2.0, 4.0), style="lit", label="LMCache E2E 지연 50~75% 감소 (r 2.0~4.0) = 경쟁 기술 수준"),
                dict(kind="range", v=(2.59, 3.96), style="lit", label="LMCache 고 QPS TTFT 4.4~6.6x·ITL 34~58% → r 2.6~4.0"),
                dict(kind="range", v=(2.0 * 1.2, 2.0 * 1.83), style="calc", label="★3 근거: 경쟁 r 2.0 × offload 대비 문헌 개선 1.2~1.83x = 2.4~3.7"),
            ]),
        keep_picture=True,
    )


# ============================================================ 페이지 6~11 구성
def build_criteria_slides(prs):
    sl = prs.slides
    specs = {5: q1_spec(), 6: q2_spec(), 7: q3_spec(), 9: q4_spec()}
    for idx, spec in specs.items():
        s = sl[idx]
        ty = criteria_page(s, spec)
        ch = spec["chart"]
        if spec.get("keep_picture") or spec.get("left_note"):
            cx, cw = 4.85, 8.1
        else:
            cx, cw = 0.4, 12.5
        band_chart(s, cx, ty + 0.02, cw, ch["axis"], ch["bands"], ch["markers"], title=ch["title"],
                   row_h=spec.get("row_h_chart", 0.205), label_size=6.8 if spec.get("row_h_chart") else 7.2)
    s = sl[7]                                   # Q3: P99 설명 + 그림
    pics = [p for p in s.shapes if p.shape_type == 13]
    tbox(s, 0.4, 4.68, 4.3, 0.5,
         [[("P99의 의미", {"bold": True, "color": NAVY, "size": 9.5})],
          [("상위 1%의 극단적 지연을 포착해 고부하 서비스 끊김을 검증", {"size": 8.5})]], size=8.5, line_spacing=1.05)
    for p in pics:
        p.left, p.top = Inches(0.45), Inches(5.25)
        scale = 1.75 / (p.height / 914400)
        p.width, p.height = int(p.width * scale), int(p.height * scale)
    s = sl[6]                                   # Q2: HBM 사용률 정의
    tbox(s, 0.4, 5.0, 4.3, 1.9,
         [[("HBM 사용량과 전체 Data 대비 사용률", {"bold": True, "color": NAVY, "size": 9.5})],
          [("• HBM 사용량: HBM에 올라와 있는 KV의 시간 평균 크기", {"size": 8.5})],
          [("  예) 100초 중 40초 100 GiB, 60초 50 GiB → 70 GiB", {"size": 8.5, "color": "404040"})],
          [("• 요청이 늘면 HBM 크기 내 점유율은 포화되어 이동 효율을 가리지 못한다 → 전체 KV 대비 사용률을 병기", {"size": 8.5})],
          [("  예) HBM 점유 40→64 GiB, KV 120 GiB: 크기 내 50→80%, 전체 대비 66→53%", {"size": 8.5, "color": "404040"})],
          [("• 압축이 아니므로 정확도(출력)는 변하지 않는다.", {"size": 8.5})]],
         size=8.5, line_spacing=1.1)
    slo_page(sl[8])
    q5_page(sl[10])
    s = sl[10]
    ax_m = dict(min=0, max=7, log=False, ticks=[0, 1, 2, 3, 4, 5, 6, 7], unit="", reverse=True)   # 왼쪽 = 미충족, 오른쪽 = 높은 별
    bands_m = [(0, 2, 3), (2, 2.5, 2.5), (2.5, 3, 2), (3, 4, 1.5), (4, 5, 1), (5, 7, 0)]
    mk_m = [dict(kind="point", v=1.7, style="lit", label="참고: SWE-bench 평균 패치 1.7 file(단위 file ≠ module) · ★3 경계 2 module ≈ 0.4 MM"),
            dict(kind="point", v=3.0, style="calc", label="★2 경계 3 module ≈ 0.6 MM(약 2.5주) · 설계 component의 약 1/4, 한 계층 안"),
            dict(kind="point", v=5, style="req", label="★1 경계 5 module ≈ 1.0 MM(1인 1개월) · 설계 component 12~14개의 약 1/3")]
    band_chart(s, 0.4, 5.5, 12.5, ax_m, bands_m, mk_m, title="변경 module 수 (시나리오 평균) — 오른쪽일수록 별이 높다", row_h=0.2, label_size=7.2, legend=True)


# ============================================================ DP 페이지(2~5): 절대값·개선율 병기
DP_SLIDE = {"DP1": 1, "DP2": 2, "DP3": 3, "DP4": 4}
# 같은 후보 영역(병합된 폭) 안에서 [라벨, (빈 칸), 별] 열 너비를 다시 나눈다 (별 칸을 넓혀 절대값을 넣는다)
COL_PLAN = {
    "DP1": [0.96, 2.86, 1.33, 1.75, 3.05, 1.15, 1.75],
    "DP2": [0.96, 4.20, 1.75, 4.20, 1.75],
    "DP3": [0.96, 4.20, 1.75, 4.20, 1.75],
    "DP4": [0.96, 4.19, 1.75, 4.20, 1.75],
}


def _auto_comment(dp, k, q):
    o = cd.OWNER[(dp, k)].get(q)
    t = cd.theory_star(dp, k, q)
    if o is None or t is None:
        return None
    other = 2 if k == 1 else 1
    if (dp, other) in cd.THEORY and q in cd.THEORY[(dp, other)]:
        a, b = cd.THEORY[(dp, k)][q], cd.THEORY[(dp, other)][q]
        key = {"Q1": "c", "Q4": "c"}.get(q)
        if key and abs(a[key] - b[key]) / max(abs(a[key]), 1e-9) < 0.01:
            return "%d안과 동일" % other
    if t - o >= 1.0:
        return "예상보다 높음"
    if o - t >= 1.0:
        return "예상보다 낮음"
    return None


def annotate_dp(slide, dp):
    tbl_shape = next(s for s in slide.shapes if s.has_table)
    tbl = tbl_shape.table
    for i, w in enumerate(COL_PLAN[dp]):
        tbl.columns[i].width = Inches(w)
    row = tbl.rows[len(tbl.rows) - 1]
    cells = row.cells
    star_idx = [i for i, c in enumerate(cells) if "●" in c.text]
    for k, si in enumerate(star_idx, start=1):
        li = si - 1
        blanks = []
        while li > 0 and not cells[li].text.strip().startswith("Q"):
            blanks.append(li)
            li -= 1
        label_cell, star_cell = cells[li], cells[si]
        labels = {}
        for p in label_cell.text_frame.paragraphs:
            m = re.match(r"\s*Q(\d)\.", p.text)
            if m:
                labels["Q" + m.group(1)] = p
        if "Q2" in labels:                       # 사용자 덱의 Q2 라벨을 Capacity로 교체
            for r in labels["Q2"].runs:
                if "Resource Efficiency" in r.text:
                    r.text = r.text.replace("Resource Efficiency", "Capacity")
                elif r.text.strip() in ("(HBM", "사용량", "GB)") or r.text in ("(HBM ", "사용량 ", "GB)"):
                    r.text = ""
        q_order = list(labels.keys())
        star_pars = [p for p in star_cell.text_frame.paragraphs if p.text.strip()]
        star_cell._tc.get_or_add_tcPr().set("marR", "0")
        blank_cell = cells[blanks[0]] if blanks else None
        if blank_cell is not None:
            btf = blank_cell.text_frame
            while len(btf.paragraphs) < len(q_order):
                btf.add_paragraph()
        for qi, q in enumerate(q_order):
            qq = "QC" if q == "Q2" else q          # 이 덱의 Q2는 Capacity로 교체(기존 자원 효율은 평가 기준 페이지에만 유지)
            txt = cd.cell_text(dp, k, qq) if (dp, k) in cd.THEORY else "—"
            r = star_pars[qi].add_run()
            r.text = " " + txt
            _font(r, 6.3, False, "404040")
            cm = _auto_comment(dp, k, q) if (dp, k) in cd.THEORY else None
            if q == "Q2":
                cm = "Capacity 별 재평가 필요"
            if blank_cell is not None:
                bp = blank_cell.text_frame.paragraphs[qi]
                sp0 = bp.add_run()
                sp0.text = " "
                _font(sp0, 12, False, "FFFFFF")
                if cm:
                    rr = bp.add_run()
                    rr.text = "▸ " + cm
                    _font(rr, 7.5, False, "0000FF")
            elif cm:
                rr = labels[q].add_run()
                rr.text = "   ▸ " + cm
                _font(rr, 7.5, False, "0000FF")
    for sh in slide.shapes:                     # 상단 QA 태그
        if sh.has_text_frame and sh.text_frame.text.strip() == "Q2 HBM 점유율":
            rr = [r for p in sh.text_frame.paragraphs for r in p.runs]
            rr[0].text = "Q2 Capacity"
            for r in rr[1:]:
                r.text = ""
    tbox(slide, 0.62, 7.2, 4.3, 0.28,
         [[("작은 수치 = B0(GPU-only) → 후보의 이론 예상값(분석 모델, 시뮬레이션 아님)과 B0 대비 개선. ", {"size": 6.3, "color": "595959"}),
           ("파란 글씨 = 소유자 별 대비 코멘트.", {"size": 6.3, "color": "0000FF"})]],
         size=6.3, line_spacing=0.95, margins=(0.02, 0.0, 0.02, 0.0), name="legend_v4")


# ============================================================ 요구사항 페이지(1)
def _replace_run(par, old, new):
    for r in par.runs:
        if r.text == old:
            r.text = new
            return True
    return False


def edit_requirements(slide):
    """Baseline = GPU-only. 요구값(2x, 20%, 30%)은 원문 그대로 두고 Baseline 표기만 명확히 한다."""
    tbl = next(s for s in slide.shapes if s.has_table).table
    for r in tbl.rows:
        n = r.cells[0].text.strip()
        c = r.cells[3]
        pars = c.text_frame.paragraphs
        if n == "QA-01":
            _replace_run(pars[0], "(Baseline)", "(Baseline: GPU-only)")
        elif n == "QA-02":
            _replace_run(pars[0], "Baseline", "Baseline(GPU-only)")
            _replace_run(pars[0], "20% ] ]", "20% ]")
        elif n == "QA-03":
            _replace_run(pars[0], "baseline ", "Baseline(GPU-only) ")
        elif n == "QA-05":
            p0 = pars[0]
            sz = 9
            for run in p0.runs:
                if run.font.size:
                    sz = run.font.size.pt
                break
            for p in pars[1:]:
                p._p.getparent().remove(p._p)
            for run in list(p0.runs):
                p0._p.remove(run._r)
            add_runs(p0, "새로운 메모리 tier 추가, 이동·배치·재사용 정책 교체, vLLM upstream 갱신 시 서빙 경로를 수정하지 않고 반영할 수 있어야 한다.", sz)
            p1 = c.text_frame.add_paragraph()
            add_runs(p1, [("[측정: 변경 모듈 수 ≤ 5개 (변경 시나리오 3종 평균, 환산 개발 공수 ≈ 1.0 M/M 이하)]", {"color": "000000", "bold": True})], sz)


# ============================================================ 신규 페이지
def clone_chrome(prs, ref_idx, title, subtitle, src):
    ref = prs.slides[ref_idx]
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    for nm in ("title", "rule", "src"):
        sh = shape_by_name(ref, nm)
        s.shapes._spTree.append(copy.deepcopy(sh._element))
    sub = None
    for sh in ref.shapes:
        if sh.name == "subtitle" and sh.top < Inches(0.9):
            sub = sh
    s.shapes._spTree.append(copy.deepcopy(sub._element))
    set_text_keep_first(shape_by_name(s, "title"), [[(title, {"size": 22, "bold": True, "color": NAVY})]])
    for sh in s.shapes:
        if sh.name == "subtitle":
            set_text_keep_first(sh, [[(subtitle, {"size": 11, "color": "404040"})]])
    set_text_keep_first(shape_by_name(s, "src"), [[(src, {"size": 7.5, "color": "404040"})]])
    return s


def scenario_slides(prs):
    s = clone_chrome(prs, 5, "평가 시나리오 - Agent KV Fleet (DP1~4 공통)",
                     "KV 캐시 관리만 대상(RAG·LoRA·MoE 제외). KV가 HBM·DRAM 용량을 넘고 재사용할 history가 있는 구간을 겨냥한 통합 시나리오",
                     "근거: Llama-3.1-70B KV 327,680 B/token · H100×8 HBM 640 GiB, 가중치 약 131 GiB · prefill ≈ 25k tok/s(MFU 0.45) · DRAM 1 TiB, NVMe 16 GB/s(system-specs). 부하·idle 분포는 설계값[가정]")
    rows = [["항목", "값", "근거 / 계산"],
            ["모델·HW", "Llama-3.1-70B BF16, 8-GPU 노드(SYS-H100): HBM 640 GiB, DRAM 1 TiB(호스트 링크 64 GB/s), NVMe 16 GB/s. DP2·DP4는 4~16 노드",
             "기존 평가 환경과 동일(비교 가능성). 노드 간 링크는 RDMA 50 GB/s [ASSUMED]"],
            ["요청 모델", "open-loop. 세션 idle 시간 지수분포 평균 60 s(tool·사용자 대기, 0.2 s~수 분), 컨텍스트 평균 16K, 턴당 +1K 입력·출력 256, 세션당 10~20턴",
             "16K 세션 KV = 327,680 B × 16,384 ≈ 5.0 GiB. 동시 세션 수 = λ × (idle + 응답) ≈ 250~450"],
            ["메모리 압박", "HBM 640 GiB 중 KV 풀(가중치·활성값 제외) ≈ 425 GiB(≈85 세션분) + DRAM 922 GiB(≈184 세션분). 동시 세션 KV가 HBM+DRAM(1.35 TiB)을 넘으면 SSD가 쓰인다",
             "반환 세션이 HBM/DRAM/SSD에서 발견될 확률은 LRU stack distance로 계산(idle ~ Exp(60 s))"],
            ["재사용", "세션 간 공유 문서 청크 30%, 그중 50%는 non-prefix 위치. 선택 재계산 15% 토큰 [가정]",
             "DP3 재사용·선택 재계산이 겨냥하는 조건. B0·경쟁 기술은 세션마다 복제 저장"],
            ["부하 구간·SLO", "Q1: 부하 sweep의 Max SLO goodput(이용률 상한 0.8 [가정]). Q3: 일반 부하 = B0 포화 부하의 50%, 고부하 = 90% 두 점에서 TTFT·TPOT P99를 잰다(별점은 고부하의 B0 대비 r, 일반 부하는 열화 확인용). SLO: TTFT P99 ≤ 2 s, TPOT P99 ≤ 50 ms",
             "이론값 TTFT P99(일반→고부하): B0 %.2f→%.2f s(×%.0f, 고부하에서 SLO 초과), B1 %.2f→%.2f s. Q4: 노드 2→16 sweep, SE(16)" % _LOADROW],
            ["Baseline: B0 GPU-only", "vLLM 기본: HBM prefix cache + preemption recompute, DRAM·SSD 사용 없음. 모든 별 기준은 B0 대비 변화량",
             "HBM에 없으면 16K 재prefill ≈ 0.65 s. '스케줄링도 offload도 없는' 하한"],
            ["경쟁 기술: 단순 offloading (B1, 경쟁 기술의 이론 대리)", "HBM이 차면 LRU로 DRAM, DRAM이 차면 SSD로 이동(턴 종료 시 증분만 기록하는 write-through, 축출은 복사 없이 삭제). 요청 시 동기 복원, prefetch·예측·배치 정책 없음",
             "LMCache류 CPU offload를 SSD 계층까지 확장한 정의. 이론값 B1 = B0의 ×2.35: LMCache 단일 노드 하한(2.3x)과 같은 수준이지만 문헌 8건 중앙값(4x)보다는 낮다"],
            ["반복", "분석 모델은 결정적(deterministic). 시뮬레이션 전환 시 seed ≥ 5, median, 95% CI", "단일 seed로 별점을 매기지 않는다"]]
    add_table(s, 0.4, 1.2, [1.9, 6.0, 4.63], rows, [0.26] + [0.55] * 8, fills=[None] * 8, size=8.2, aligns=["l", "l", "l"])
    tbox(s, 0.4, 6.25, 12.5, 0.8,
         [[("조건부 별: ", {"bold": True, "color": "C00000", "size": 9.5}),
           ("이 시나리오는 KV가 HBM·DRAM을 넘고 재사용할 history가 있는 구간을 의도적으로 겨냥한다(문헌이 이득을 보고한 조건과 같은 영역). "
            "시나리오 한계: 단순 offloading의 이론값은 지연(Q3)에서 문헌 수준(r 2.0)에 놓이지만 처리량(Q1)에서는 ×2.35로 문헌 하단이다. 이 시나리오는 세션별 히스토리 재사용만 있어 처리량 이론 상한이 ×4.0이라, ★3(8x)는 공유 캐시·고 적중 regime(문헌)에서만 가능하다. 값은 분석 모델의 예상값이며 시뮬레이션 전이다.", {"size": 9})]],
         size=9, line_spacing=1.1)

    s2 = clone_chrome(prs, 5, "평가 시나리오 - 구간(Phase)과 DP·QA 대응",
                      "한 시나리오 안에서 DP별 메커니즘이 작동하는 구간을 나누고, QA별 측정 규칙을 정한다",
                      "N 노드 확장(P4)은 별도 run(N = 2, 4, 8, 16, 세션 수 ∝ N). 모든 값은 분석 모델 기반 설계값이다.")
    rows = [["Phase", "구간 내용", "겨냥하는 DP 메커니즘", "측정 QA"],
            ["P1 정상 multi-turn\n(0~10분)", "agent 세션이 idle을 두고 턴을 이어 감. 유휴 KV가 HBM을 점유",
             "DP1: 유휴 KV 하위 tier 이동(demote)과 재개 시 승격·prefetch. 1안 = 자원 상태 기반, 2안 = 데이터 접근 특성 기반", "Q1, Q2, Q3, Q5"],
            ["P2 burst\n(10~15분)", "신규 세션 3배 폭주 + 기존 세션 decode가 겹침(prefill burst + decode-heavy)",
             "DP2: Prefill/Decode 실행 위치 결정. 1안 = 스케줄링 시점 결정, 2안 = 사전 계획 결정(Re-plan 위험)", "Q1, Q2, Q3, Q4"],
            ["P3 공유 문서 wave\n(15~25분)", "같은 문서 청크를 여러 세션이 연달아 참조(non-prefix 50%)",
             "DP3: 재사용·재계산 대상 선정. 1안 = Pipe-and-Filter(요청마다 체인), 2안 = Shared Repository(분석 1회 + 조회), 둘 다 선택 재계산", "Q1, Q2, Q3"],
            ["P4 노드 확장\n(별도 run)", "노드 2→16, 세션 수 ∝ 노드 수. 노드 간 KV 블록 조회·소유권 이동 발생",
             "DP4: 서버 간 KV 공유. 1안 = 중앙 서버 중재, 2안 = 락 기반 직접 접근. CXL 공유 풀이 SSD를 대체", "Q1, Q3, Q4"],
            ["변경 시나리오\n(구현 측정)", "S1 새 메모리 tier · S2 정책 교체 · S3 vLLM upstream 갱신", "모든 DP 후보의 module 수 (시나리오 평균, M/M 환산 병기)", "Q5"]]
    add_table(s2, 0.4, 1.2, [1.6, 3.9, 5.5, 1.53], rows, [0.26] + [0.62] * 5, fills=[None] * 5, size=8.5, aligns=["l", "l", "l", "c"])
    tbox(s2, 0.4, 4.9, 12.5, 1.5,
         [[("측정 규칙", {"bold": True, "color": NAVY, "size": 10})],
          [("• Q1: 부하 sweep 최댓값끼리 R = 후보 ÷ B0 · Q2: HBM 상주 KV 시간 평균의 B0 대비 감소율 · Q3: B0 포화 부하의 90%에서 TTFT·TPOT P99로 r 계산(B0 대비)", {"size": 9})],
          [("• Q4: 노드 2→16 sweep, SE(16) = goodput 배율 ÷ 노드 배율 · Q5: 변경 시나리오 3종을 각 후보에 구현해 module 수 평균(M/M 환산 병기)", {"size": 9})],
          [("• 경쟁 기술(단순 offloading)은 판정 기준이 아니라 ★2 앵커와 비교용이다. 기준·시나리오 값은 결과를 보기 전에 고정한다. 기존 중립 벤치마크 시뮬 결과는 md의 참고 표에만 둔다.", {"size": 9})]],
         size=9, line_spacing=1.12)
    return s, s2


def baseline_slide(prs):
    s = clone_chrome(prs, 5, "평가 Baseline(GPU-only)과 경쟁 기술(단순 offloading) 절대 수치",
                     "모든 별 기준은 B0(GPU-only) 대비 변화량. 경쟁 기술(B1)의 이론값이 어느 별에 놓이는지 확인용으로 병기 (이론 계산, Agent KV Fleet, SYS-H100)",
                     "분석 모델: theory_model.py (GPU 한계 · DRAM 링크 51 GB/s · SSD 링크 12.8 GB/s 중 최솟값 × 이용률 0.8). 시뮬레이션이 아니며 후보 파라미터는 [가정]")
    c2, c4 = _R["dp2"], _R["dp4"]
    e1, e3 = cd.EDGES["Q1"], cd.EDGES["Q3"]
    t = lambda base, i: cd.fmt_q(base * e1[i])
    bs = lambda v: "%s (%s)" % (cd.fmt_q(v), "×%.2f" % (v / B0_QPS))
    rows = [["QA", "범위", "지표", "B0 GPU-only (Baseline)", "경쟁 기술 B1 (이론)", "★1", "★2", "★3"],
            ["Q1 처리량", "1 노드", "Max SLO goodput (query/s)", cd.fmt_q(B0_QPS), "%s (×%.2f, %s)" % (cd.fmt_q(B1_QPS), B1_QPS / B0_QPS, cd.fmt_star(COMP["Q1"][1])),
             "≥ " + t(B0_QPS, 0), "≥ " + t(B0_QPS, 2), "≥ " + t(B0_QPS, 4)],
            ["", "DP2 4노드 (P2+D2 고정)", "Max SLO goodput (query/s)", cd.fmt_q(c2["B0"]["qps"]), "%s (×%.2f)" % (cd.fmt_q(c2["B1"]["qps"]), c2["B1"]["qps"] / c2["B0"]["qps"]),
             "≥ " + t(c2["B0"]["qps"], 0), "≥ " + t(c2["B0"]["qps"], 2), "≥ " + t(c2["B0"]["qps"], 4)],
            ["", "DP4 4노드 (노드 독립)", "Max SLO goodput (query/s)", cd.fmt_q(c4["B0"]["qps"]), "%s (×%.2f)" % (cd.fmt_q(c4["B1"]["qps"]), c4["B1"]["qps"] / c4["B0"]["qps"]),
             "≥ " + t(c4["B0"]["qps"], 0), "≥ " + t(c4["B0"]["qps"], 2), "≥ " + t(c4["B0"]["qps"], 4)],
            ["Q2 Capacity", "1 노드", "KV 보존 동시 세션 수 (HBM 가득 채움, 16K 기준)", "%d (HBM만)" % round(tm.capacity_sessions(tm.B0)),
             "%d (×%.1f)" % (round(tm.capacity_sessions(tm.B1)), tm.capacity_sessions(tm.B1) / tm.capacity_sessions(tm.B0)),
             "≥ %d (2x)" % round(2 * tm.capacity_sessions(tm.B0)), "미정", "미정"],
            ["Q2' HBM 효율(참고)", "1 노드", "HBM 상주 KV (GiB, 낮을수록 좋음)", n0(HBM_B0), "%s (0%%, ☆: LRU가 HBM을 채움)" % n0(COMP_B1_HBM),
             "≤ %s" % n0(HBM_B0 * 0.8), "≤ %s" % n0(HBM_B0 * 0.6), "≤ %s" % n0(HBM_B0 * 0.4)],
            ["Q3 고부하 지연", "B0 포화 부하 90%", "TTFT P99 (s) / TPOT P99 (ms)", "%.2f s / %.1f ms (SLO 초과)" % (B0_TTFT, B0_TPOT),
             "%.2f s / %.1f ms (r %.2f, ★2)" % (B1_TTFT, B1_TPOT, COMP["Q3"][0]),
             "TTFT ≤ %.2f s" % (B0_TTFT / e3[0] ** 2), "≤ %.2f s" % (B0_TTFT / e3[2] ** 2), "≤ %.2f s" % (B0_TTFT / e3[4] ** 2)],
            ["Q4 확장성", "N=16", "SE (%)", "98 (노드 독립)", "98 (노드 독립, ★3)", "≥ 80", "≥ 90", "≥ 95"],
            ["Q5 변경 용이성", "후보 구현", "module 수 (환산 M/M)", "—", "—", "≤5 (≈1.0)", "≤3 (≈0.6)", "≤2 (≈0.4)"]]
    add_table(s, 0.4, 1.2, [1.3, 1.75, 2.35, 1.9, 2.3, 1.0, 1.0, 0.93], rows, [0.28] + [0.38] * 8, fills=[None] * 8, size=8,
              aligns=["l", "l", "l", "c", "c", "c", "c", "c"])
    tbox(s, 0.4, 4.7, 6.1, 2.4,
         [[("Baseline과 경쟁 기술 정의", {"bold": True, "color": NAVY, "size": 10})],
          [("• B0 GPU-only(Baseline): vLLM 기본. KV가 HBM에 없으면 16K 재prefill(≈0.65 s). 스케줄링도 offload도 없음", {"size": 8.8})],
          [("• 경쟁 기술 B1(단순 offloading): HBM이 차면 LRU로 DRAM, DRAM이 차면 SSD로 이동. 턴 종료 시 증분만 write-through, 요청 시 동기 복원, prefetch·예측 없음", {"size": 8.8})],
          [("• DP2의 B0는 P/D 고정 역할, DP4의 B0·B1은 노드 독립(교차 라우팅 25%는 재계산)", {"size": 8.8})]],
         size=8.8, line_spacing=1.12)
    tbox(s, 6.7, 4.7, 6.2, 2.4,
         [[("읽는 법", {"bold": True, "color": NAVY, "size": 10})],
          [("• ★1 = Stakeholder 요구(B0 대비), ★2 = 경쟁 기술 문헌 대표값, ★3 = 경쟁 기술 위 목표(Q1은 문헌 상위권 8x, Q2·Q3는 ★2 × 1.5). 수치는 모두 B0 대비", {"size": 8.8})],
          [("• 경쟁 기술(B1) 이론값: Q1 ×%.2f(%s, 문헌 중앙값 4x에는 못 미침), Q3 r %.2f(%s). Q1은 이 시나리오의 이론 상한(×%.1f)이 ★3(8x)보다 낮아 ★3에 도달할 수 없다" % (B1_QPS / B0_QPS, cd.fmt_star(COMP["Q1"][1]), COMP["Q3"][0], cd.fmt_star(COMP["Q3"][1]), CEIL_RATIO * B1_QPS / B0_QPS), {"size": 8.8})],
          [("• Q2는 예외: 정적 LRU offload는 HBM을 캐시로 채워 감소 0%다. ★2의 40%는 HBM 절감을 목적으로 한 offload 문헌(DirectKV 등)의 값이다", {"size": 8.8})],
          [("• 이론 상한: 처리량은 GPU 한계로 B1의 %.1fx(= B0의 ×%.1f), HBM은 ≈88%% 절감, TTFT는 prefill(≈41 ms)까지" % (CEIL_RATIO, CEIL_RATIO * B1_QPS / B0_QPS), {"size": 8.8})]],
         size=8.8, line_spacing=1.12)
    return s


def method_slide(prs):
    s = clone_chrome(prs, 5, "이론 계산 기반 예상값 산출 방식",
                     "B0·B1은 물리 한계에서 계산하고, 후보는 같은 식에 설계 파라미터만 바꿔 넣는다 (시뮬레이션 아님)",
                     "코드: doc-mk/QA-criteria-v4/theory_model.py. 후보 파라미터는 덱의 장단점과 DP1 요구 수치(prefetch 적중 ≥ 60~70%, 낭비 ≤ 30%)에서 방향만 가져온 [가정]이며 시뮬레이션으로 교체한다")
    tbox(s, 0.4, 1.17, 6.2, 3.0,
         [[("계산 구조", {"bold": True, "color": NAVY, "size": 10})],
          [("① 반환 세션 hit: LRU stack distance. idle ~ Exp(60 s), 그 사이 서로 다른 세션 λT개 접근 → HBM(85 세션분)·DRAM(184 세션분)에서 발견될 확률", {"size": 8.5})],
          [("② 처리 한계 λ = min(GPU 1/0.113 s, DRAM 링크 51 GB/s ÷ 이동 바이트, SSD 링크 12.8 GB/s ÷ 이동 바이트). B0는 miss마다 0.65 s 재계산. Max SLO goodput = 0.8 × 256 tok × λ", {"size": 8.5})],
          [("③ 이동 바이트 = miss × (1−π) 동기 복원 + miss × π × (1+ω) prefetch + w × miss write-back (π 적시율, ω 낭비, w write-back 비율; B1은 w=0.1 증분 기록)", {"size": 8.5})],
          [("④ 고부하 = B0 포화 부하의 90%. TTFT P99 = 새 입력 prefill 41 ms + 결정 지연 + 99% 복원 경로 ÷ (1 − 링크 이용률). TPOT P99 = 18.1 ms × (1 + 0.5 × 이용률)", {"size": 8.5})],
          [("⑥ Capacity = θ × (HBM 85 + DRAM 184 세션분) ÷ f′ (HBM 가득 채움, SSD 제외, 신규 메모리 미반영 [잠정]). B0는 HBM만 85세션", {"size": 8.5})],
          [("⑤ SE(16) = min(1, 결정 용량 ÷ 요구 결정률). 결정 비용은 후보 수에 선형(2 ms × N) [기존 평가의 ASSUMED]", {"size": 8.5})]],
         size=8.5, line_spacing=1.1)
    tbox(s, 6.8, 1.17, 6.1, 3.0,
         [[("한계(정직하게)", {"bold": True, "color": "C00000", "size": 10})],
          [("• 후보 파라미터(θ, π, ω, watermark, u_P …)는 가정이다. 값이 바뀌면 별이 바뀐다", {"size": 8.5})],
          [("• 시나리오는 단순 offloading 이론값이 Q3(r 2.0)에서 문헌 수준이 되도록 정했다. Q1은 문헌 하단(×2.35)이며 이 선택 자체가 설계 선택이다", {"size": 8.5})],
          [("• 큐잉은 링크 이용률 1/(1−ρ)로 근사했다. 연속 배칭·chunked prefill의 세부는 반영하지 않았다", {"size": 8.5})],
          [("• 이용률 상한 0.8, idle 평균 60 s, 공유 문서 30%는 설계값이다", {"size": 8.5})],
          [("• DP2·DP4는 4노드 요약 모델이다. 결정 용량·제어 평면 지연은 기존 평가의 상수를 썼다", {"size": 8.5})]],
         size=8.5, line_spacing=1.1)
    rows = [["후보", "θ 배치 효율", "w write-back", "π prefetch 적시", "ω 낭비", "추가 GPU·결정 지연", "HBM watermark", "근거(덱 장단점·요구 수치)"],
            ["B1 경쟁 기술", "1.0", "0.1", "0", "0", "—", "1.00", "정의상 정적 LRU, 증분 write-through, 동기 복원"],
            ["DP1 1안 자원 상태", "1.1", "0.1", "0", "0", "결정 3 ms", "0.80", "임계 trigger로 HBM 압박 구간 단축, 접근 특성 구분 한계"],
            ["DP1 2안 데이터 특성", "1.3", "0.1", "0.6", "0.3", "0", "0.95", "접근 특성 기반 이동·prefetch(요구 70%에서 낭비 차감), snapshot stale로 HBM 확보 지연"],
            ["DP2 1안 스케줄링 시점", "유휴 P 노드 활용 u_P 0.2", "—", "—", "—", "결정 8 ms(2 ms × N)", "0.85", "실시간 상태 반영, 결정 지연이 critical path"],
            ["DP2 2안 사전 계획", "u_P 0.5", "—", "—", "—", "Re-plan P99 +40 ms", "0.92", "결정은 critical path 밖, 고부하 시 Re-plan(31%), plan 예약으로 HBM 여유 필요"],
            ["DP3 1안 Pipe-and-Filter", "dedup 0.70", "0.1", "0", "0", "GPU +27 ms, 체인 tail 25 ms", "0.60", "관련 KV만 보유(HBM↓), 요청마다 체인 중복 계산, 고부하 tail"],
            ["DP3 2안 Shared Repository", "dedup 0.70", "0.1", "0", "0", "GPU +3 ms", "0.75", "분석 1회+조회라 처리량·tail 유리, 재사용 KV·위치 사전 보유로 점유↑"],
            ["DP4 1안 / 2안", "CXL 풀이 SSD 대체", "0.1", "0", "0", "RPC 7 µs / 락 P99 804 µs × 4 op", "—", "노드 유효 63 GB/s, 교차 라우팅 25% 재계산 제거, 제어 평면은 요청당 ms 이하"]]
    add_table(s, 0.4, 4.2, [1.9, 1.35, 0.85, 1.05, 0.7, 1.9, 1.0, 3.78], rows, [0.3] + [0.325] * 8, fills=[None] * 8, size=7.4,
              aligns=["l", "c", "c", "c", "c", "c", "c", "l"], header_size=8)
    return s



def sim_env_slide(prs):
    s = clone_chrome(prs, 5, "평가 환경과 시뮬레이션 정의",
                     "무엇이 실장 장비이고 무엇이 시뮬레이션인지, DP4 시뮬레이터가 무엇을 모델링하는지, 이번 수치가 어디서 나왔는지",
                     "출처: doc-mk/Evaluation/DP4/sim (DESIGN_NOTES.md, simulation-plan.md), system-specs.md, project-context GC-1·GC-9. 별점·표 수치는 theory_model.py 이론 예상값, 시뮬레이터 실행 결과는 DP4/results/2026-10-25_dp4-agent-kv-fleet.md")
    rows = [["DP", "평가 환경", "실장 / 시뮬레이션", "수치 산출 (이 덱)"],
            ["DP1", "H100×8 1노드, HBM 640 GiB · DRAM 1 TiB · NVMe 16 GB/s", "모두 실장 장비 구성", "이론 예상값(분석 모델)"],
            ["DP2", "H100×8 4노드(P 2 + D 2 고정 역할이 Baseline), 노드 간 링크 RDMA 50 GB/s [ASSUMED]", "노드·메모리는 실장 구성, 노드 간 링크는 가정", "이론 예상값(분석 모델)"],
            ["DP3", "H100×8 1노드(DP1과 같은 구성)", "모두 실장 장비 구성", "이론 예상값(분석 모델)"],
            ["DP4", "H100×8 4노드 + CXL 공유 메모리 풀 8 TiB", "노드는 실장, CXL 풀은 장비가 없어 시뮬레이션", "덱 수치는 이론 예상값. DP4 시뮬레이터로 Agent KV Fleet 실행 완료(별도 결과 문서, Baseline-RDMA 대비)"]]
    add_table(s, 0.4, 1.2, [0.7, 5.0, 3.4, 3.43], rows, [0.26] + [0.42] * 4, fills=[None] * 4, size=8.3, aligns=["c", "l", "l", "l"])
    tbox(s, 0.4, 3.35, 6.2, 3.7,
         [[("DP4 시뮬레이터(레포 doc-mk/Evaluation/DP4/sim)", {"bold": True, "color": NAVY, "size": 10})],
          [("• 방식: 이산 사건 시뮬레이터(DES). P/D 분리 클러스터에서 세 가지 구성(Baseline-RDMA, 중앙 직렬화 C1, 분산 락 C2)을 같은 요청 흐름으로 비교", {"size": 8.6})],
          [("• 링크: 선입선출 유체 서버(FIFO fluid server). 노드 egress/ingress, CXL 풀 집계 대역폭을 공유 자원으로 모델링, 활성 전송은 공정 분할", {"size": 8.6})],
          [("• decode: processor sharing. prefill·decode 단계 시간은 DP1 모델(`dp1_bridge`)을 그대로 가져와 DP 간 물리가 같다", {"size": 8.6})],
          [("• 제어 평면: lookup · publish · pin · unpin 연산별 지연. 중앙 RPC 약 2~7 µs, 분산 락 2.6~804 µs(P99)", {"size": 8.6})],
          [("• 파라미터: 전부 값·범위·출처(PAPER / SPEC / ASSUMED)를 기록. 결과를 본 뒤 바꾸지 않는다", {"size": 8.6})],
          [("• 반복: seed 5개(11, 23, 37, 53, 71), 부하 격자 ×0.5~×8, 95% 신뢰구간, 요청별 SLO로 goodput 계산", {"size": 8.6})]],
         size=8.6, line_spacing=1.1)
    tbox(s, 6.8, 3.35, 6.1, 3.7,
         [[("CXL 풀 모델과 가정", {"bold": True, "color": NAVY, "size": 10})],
          [("• 풀 8 TiB, 노드당 어댑터 2개(PCIe 5.0 x16, 각 63 GB/s), 풀 집계 1 TB/s, 유효 효율 0.5 [ASSUMED] → 노드 유효 약 63 GB/s", {"size": 8.6})],
          [("• 접근 지연은 제어 평면 비용과 데이터 이동 시간으로 반영. 풀 용량 한계와 HBM OOM은 미반영", {"size": 8.6})],
          [("• 미반영: CXL 스위치 혼잡, 메모리 뱅크 경합, 락 매니저 스레드의 CPU 스케줄링 간섭, 풀 장치 장애, 전력", {"size": 8.6})],
          [("이번 수치의 위치", {"bold": True, "color": "C00000", "size": 10})],
          [("• 이 덱의 별점·표 값은 theory_model.py 이론 예상값이다(GPU-only B0 기준). 시뮬레이터 Baseline-RDMA는 B0도 B1도 아니라(HBM 용량 압박·DRAM/SSD tier 미모델링) 직접 대체하지 않는다", {"size": 8.6})],
          [("• %s. 이론 모델과는 기준·유휴 모델(sim 8 s vs 60 s)이 달라 대조 불가" % "시뮬레이터 실행(2026-10-25, Agent KV Fleet 16K/10턴/노드 2~16, seeds 5): Baseline-RDMA 대비 goodput H100 x1.00(N=4)·x1.2(N=8)·x1.4(N=16), B200 x0.51(CXL ingress 포화). 상세 DP4/results/2026-10-25_dp4-agent-kv-fleet.md", {"size": 8.6})],
          [("• 노드·HBM·DRAM·SSD는 실장 장비라 DP1·DP3와 B0·B1 값은 실측으로 검증해 모델 파라미터를 보정할 수 있다", {"size": 8.6})]],
         size=8.6, line_spacing=1.1)
    return s


QLAB = {"Q1": ("Q1 처리량", "SLO(TTFT ≤ 2 s, TPOT ≤ 50 ms) 만족 Max goodput"),
        "Q2": ("Q2 HBM 효율", "HBM 상주 KV 시간 평균(GiB), 같은 부하·SLO"),
        "QC": ("Q2 Capacity", "같은 HBM에서 KV를 보존해 재개 가능한 최대 동시 세션 수(HBM 가득 채움, SSD 제외, 신규 메모리 미반영 [잠정])"),
        "Q3": ("Q3 고부하 지연", "B0 포화 부하의 90%에서 TTFT·TPOT P99"),
        "Q4": ("Q4 확장성", "SE(N=16) = goodput 배율 ÷ 노드 배율"),
        "Q5": ("Q5 변경 용이성", "변경 시나리오 평균 module 수 (기존 구현 측정, M/M은 환산 병기)")}

DPINFO = {
    "DP1": dict(title="DP1. Data Migration 결정 구조 — Trade-off 평가 결과", names=("1안 상태 관측+이동 결정 통합", "2안 관측·결정 분리"),
                qs=["Q1", "QC", "Q3", "Q5"], scope="single",
                env=["1 노드: H100×8(HBM 640 GiB), DRAM 1 TiB(호스트 링크 64 GB/s), SSD 16 TiB(NVMe 16 GB/s), 신규 메모리는 다음 페이지", "Llama-3.1-70B BF16, SYS-H100. Baseline B0 = GPU-only, 경쟁 기술 B1 = 단순 offloading"],
                scen=["Agent KV Fleet P1 정상 multi-turn: 컨텍스트 16K, 턴당 +1K·출력 256, idle 평균 60 s(open-loop)",
                      "동시 세션 ≈ 250~450 → KV 1.3~2.2 TiB(HBM 전체 640 GiB 중 KV 풀 425 GiB의 3~5배, HBM+DRAM도 초과)"]),
    "DP2": dict(title="DP2. 연산(Prefill/Decode) 실행 위치 결정 구조 — Trade-off 평가 결과", names=("1안 스케줄링 시점 결정", "2안 사전 계획 결정"),
                qs=["Q1", "QC", "Q3", "Q4"], scope="cluster",
                env=["H100×8 4노드 클러스터(실장 구성 기준): B0 = P 2대 + D 2대 고정 역할 GPU-only, B1 = 같은 구성 + 노드별 단순 offloading", "노드 간 링크 RDMA 50 GB/s [ASSUMED], 결정 비용은 후보 수에 선형(2 ms × N). 수치는 이론 예상값"],
                scen=["Agent KV Fleet P2 burst: 신규 세션 3배 폭주 + decode 중첩. 동적 배치가 유휴 P 노드의 tier·링크를 쓰는 정도 평가",
                      "확장성은 N=16 sweep(scheduler 결정 용량 대 요구 결정률)"]),
    "DP3": dict(title="DP3. KV 캐시 재사용 구조 — Trade-off 평가 결과", names=("1안 Pipe-and-Filter", "2안 Shared Repository"),
                qs=["Q1", "QC", "Q3"], scope="single",
                env=["1 노드(DP1과 같은 HW). 압축은 범위 밖, 선택 재계산만 사용", "Llama-3.1-70B BF16, SYS-H100. B0 = GPU-only, B1 = 단순 offloading"],
                scen=["Agent KV Fleet P3 공유 문서 wave: 컨텍스트의 30%가 세션 간 공유 문서(그중 50% non-prefix)",
                      "B0·B1은 세션마다 복제 저장, DP3는 1회 보관 + 선택 재계산(15% 토큰) [가정]"]),
    "DP4": dict(title="DP4. 서버 간 KV 캐시 공유 구조 — Trade-off 평가 결과", names=("1안 중앙 서버 중재", "2안 락 기반 직접 접근"),
                qs=["Q1", "Q3", "Q4"], scope="cluster",
                env=["H100×8 4노드(실장) + CXL 공유 풀 8 TiB(풀 장비가 없어 시뮬레이션 대상, 노드당 유효 63 GB/s). B0·B1은 노드 독립(SSD 16 GB/s)", "교차 라우팅 25%는 B0·B1에서 재계산, 후보는 풀에서 공유. 풀·제어 평면은 DP4 시뮬레이터 구조와 파라미터를 따르는 이론 모델"],
                scen=["Agent KV Fleet P4 노드 확장: 2→16 노드(세션 ∝ 노드), 노드 간 KV 블록 조회·소유권 이동",
                      "제어 평면: 중앙 RPC ≈ 7 µs vs 분산 락 P99 ≈ 804 µs (요청당 4 op)"]),
}


def _conclusion(dp):
    a, b = cd.THEORY[(dp, 1)], cd.THEORY[(dp, 2)]
    q1c = COMP["Q1"][0]
    if dp == "DP1":
        return ["GPU-only(B0) 대비 처리량 ×%.2f(1안)·×%.2f(2안), TTFT P99 r %.2f·%.2f로 경쟁 기술(단순 offloading, ×%.2f·r %.2f)과 같은 수준이다(처리량 ★1, 지연 ★2). 경쟁 기술 대비 처리량 %+.0f%%·%+.0f%%, TTFT %+.0f%%·%+.0f%%." % (
            a["Q1"]["r0"], b["Q1"]["r0"], a["Q3"]["r0"], b["Q3"]["r0"], q1c, COMP["Q3"][0],
            (a["Q1"]["r1"] - 1) * 100, (b["Q1"]["r1"] - 1) * 100, -a["Q3"]["ttft_red1"], -b["Q3"]["ttft_red1"]),
            "Capacity는 2안(데이터 특성 기반)이 배치 효율 θ가 높아 ×%.1f로 1안(×%.1f)보다 많은 세션을 보존한다. 지연은 prefetch(적시 60%%)를 쓰는 2안이 앞서지만 링크 병목 구간에서는 prefetch가 바이트를 줄이지 못해 처리량 이득이 작다." % (b["QC"]["r0"], a["QC"]["r0"])]
    if dp == "DP2":
        return ["GPU-only(P/D 고정, B0) 대비 처리량 ×%.2f(1안)·×%.2f(2안)로 경쟁 기술(×%.2f)을 넘지만 문헌 중앙값(4x)에는 못 미쳐 ★%g~★%g. 유휴 P 노드의 tier·링크를 더 쓰는 2안이 앞선다." % (
            a["Q1"]["r0"], b["Q1"]["r0"], b["Q1"]["b1"] / b["Q1"]["b0"], cd.theory_star(dp, 1, "Q1"), cd.theory_star(dp, 2, "Q1")),
            "확장성은 1안이 단일 scheduler 결정 용량에 막혀 N=16에서 SE %.0f%%, 2안은 worker 분산으로 %.0f%%. Capacity는 P 노드 tier를 보존 용량에 합산할지 정해야 해서 산정 보류." % (
                a["Q4"]["c"], b["Q4"]["c"])]
    if dp == "DP3":
        return ["두 안 모두 공유 문서를 1회만 보관(−30%%)해 GPU-only 대비 처리량 ×%.2f~×%.2f(★1.5), TTFT P99 r %.2f~%.2f(★2.5). 경쟁 기술(×%.2f) 대비 처리량 %+.0f%%~%+.0f%%." % (
            a["Q1"]["r0"], b["Q1"]["r0"], a["Q3"]["r0"], b["Q3"]["r0"], q1c, (a["Q1"]["r1"] - 1) * 100, (b["Q1"]["r1"] - 1) * 100),
            "Capacity는 두 안 모두 공유 문서 1회 보관(dedup 0.70)으로 ×%.1f(HBM 가득 채움 기준이라 두 안 차이가 사라진다). 안 간 처리량 차이는 링크 병목 구간이라 거의 없고, 1안의 고부하 tail 발산은 B0 포화 90%% 부하에서는 나타나지 않는다." % a["QC"]["r0"]]
    return ["CXL 공유 풀이 SSD(16 GB/s)를 대체하고 교차 라우팅 재계산을 없애 GPU-only 대비 처리량 ×%.2f(★1.5), TTFT P99 r %.2f(★3). 경쟁 기술 대비 처리량 %+.0f%%." % (
        a["Q1"]["r0"], a["Q3"]["r0"], (a["Q1"]["r1"] - 1) * 100),
        "두 안의 제어 평면 차이(중앙 RPC 7 µs vs 락 P99 804 µs)는 요청당 수 ms 이하라 성능이 거의 같고, N=16에서 SE %.0f%%로 동일하다. 소유자가 예상한 락 기반의 확장 우위는 이 규모에서는 드러나지 않는다." % a["Q4"]["c"]]


def dp_result_slide(prs, dp):
    info = DPINFO[dp]
    s = clone_chrome(prs, 5, info["title"],
                     "평가 환경·시나리오와 GPU-only(B0)·단순 offloading(경쟁 기술) 대비 개선 요약 (이론 계산 기반 예상값)",
                     "예상값은 theory_model.py 분석 모델(시뮬레이션 아님, 후보 파라미터 [가정]). 별은 GPU-only 대비 새 기준(5구간) 적용 결과이고 괄호는 소유자 trade-off 별. 구간·근거는 평가 기준 페이지 참조")
    tbox(s, 0.4, 1.18, 6.15, 1.1,
         [[("평가 환경", {"bold": True, "color": NAVY, "size": 9.5})]] + [[("• " + x, {"size": 8.6})] for x in info["env"]],
         size=8.6, line_spacing=1.1)
    tbox(s, 6.75, 1.18, 6.15, 1.1,
         [[("평가 시나리오", {"bold": True, "color": NAVY, "size": 9.5})]] + [[("• " + x, {"size": 8.6})] for x in info["scen"]],
         size=8.6, line_spacing=1.1)
    rows = [["QA", "무엇을 평가했나", "B0 GPU-only (Baseline)", "경쟁 기술 (단순 offloading)", "후보 " + info["names"][0], "후보 " + info["names"][1]]]
    nrow = len(info["qs"])
    for q in info["qs"]:
        lab, how = QLAB[q]
        t1, t2 = cd.THEORY[(dp, 1)].get(q), cd.THEORY[(dp, 2)].get(q)
        if t1 is None:
            hold = [[("산정 보류", {"bold": True, "size": 8.5})], [("보존 tier 정의(P 노드 합산 / CXL 풀 몫)가 정해져야 계산 가능", {"size": 7.6, "color": "404040"})]]
            rows.append([[[(lab, {"bold": True})]], how, "—", "—", hold, hold])
            continue
        if q == "Q1":
            b0c = cd.fmt_q(t1["b0"]) + " query/s"
            b1c = "%s query/s (×%.2f, %s)" % (cd.fmt_q(t1["b1"]), t1["b1"] / t1["b0"], cd.fmt_star(cd.star_q1(t1["b1"] / t1["b0"])))
        elif q == "QC":
            b0c = "%d 세션 (HBM만)" % round(t1["b0"])
            b1c = "%d 세션 (×%.1f)" % (round(t1["b1"]), t1["b1"] / t1["b0"])
        elif q == "Q2":
            b0c = n0(t1["b0"]) + " GiB"
            b1c = "%s GiB (0%%, ☆)" % n0(t1["b1"])
        elif q == "Q3":
            b0c = "TTFT %.2f s · TPOT %.1f ms (SLO 초과)" % (t1["b0_ttft"], t1["b0_tpot"])
            b1c = "TTFT %.2f s · TPOT %.1f ms (r %.2f, %s)" % (t1["b1_ttft"], t1["b1_tpot"], t1["r0"] / t1["r1"] if False else tm.r_reduction(t1["b0_ttft"], t1["b0_tpot"], t1["b1_ttft"], t1["b1_tpot"])[0],
                                                              cd.fmt_star(cd.star_q3(tm.r_reduction(t1["b0_ttft"], t1["b0_tpot"], t1["b1_ttft"], t1["b1_tpot"])[0])))
        elif q == "Q4":
            b0c, b1c = "SE %.0f%% (노드 독립)" % t1["b0"], "SE %.0f%% (노드 독립)" % t1["b1"]
        else:
            b0c, b1c = "—", "—"
        cells = []
        for k, t in ((1, t1), (2, t2)):
            st = cd.theory_star(dp, k, q)
            own = cd.OWNER[(dp, k)].get(q)
            if q == "Q1":
                l1 = "%s query/s" % cd.fmt_q(t["c"])
                l2 = "B0 대비 ×%.2f · 경쟁 기술 대비 %+.0f%%" % (t["r0"], (t["r1"] - 1) * 100)
            elif q == "QC":
                l1 = "%d 세션" % round(t["c"])
                l2 = "B0 대비 ×%.1f · 경쟁 기술 대비 ×%.2f" % (t["r0"], t["r1"])
            elif q == "Q2":
                l1 = "%s GiB" % n0(t["c"])
                l2 = "B0 대비 %+.0f%% · 경쟁 기술 대비 %+.0f%%" % (-t["red0"], -t["red1"])
            elif q == "Q3":
                l1 = "TTFT %.2f s · TPOT %.1f ms" % (t["c_ttft"], t["c_tpot"])
                l2 = "B0 대비 r %.2f(TTFT %+.0f%%) · 경쟁 기술 대비 TTFT %+.0f%%" % (t["r0"], -t["ttft_red0"], -t["ttft_red1"])
            elif q == "Q4":
                l1 = "SE %.0f%%" % t["c"]
                l2 = "B0·경쟁 기술 대비 %+.0f%%p" % (t["c"] - t["b0"])
            else:
                l1 = "%g module · %.2f MM" % (t["module"], t["mm"])
                l2 = "절대 기준(경쟁 기술 비교 없음)"
            if q == "QC":
                l3 = "★1 이상 (★2·★3 경계 미정, 소유자 별 미정)" if st >= 1 else "☆ (소유자 별 미정)"
            else:
                l3 = "%s (소유자 %s)" % (cd.fmt_star(st), cd.fmt_star(own))
            cells.append([[(l1, {"bold": True, "size": 8.5})], [(l2, {"size": 7.6, "color": "404040"})], [(l3, {"size": 8, "bold": True, "color": NAVY})]])
        rows.append([[[(lab, {"bold": True})]], how, b0c, b1c, cells[0], cells[1]])
    add_table(s, 0.4, 2.4, [1.2, 2.2, 1.7, 1.95, 2.74, 2.74], rows, [0.3] + [0.72] * nrow, fills=[None] * nrow, size=8.2,
              aligns=["l", "l", "c", "c", "l", "l"])
    y = 2.4 + 0.3 + 0.72 * nrow + 0.12
    con = _conclusion(dp)
    rect(s, 0.4, y, 12.53, 0.95, fill="F2F6FC", line="2F5597", line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tbox(s, 0.5, y + 0.03, 12.3, 0.9,
         [[("요약 (심사용)", {"bold": True, "color": NAVY, "size": 9.5})]] + [[("• " + x, {"size": 8.8})] for x in con],
         size=8.8, line_spacing=1.1)
    return s


def token_slide(prs):
    s = clone_chrome(prs, 5, "변경 용이성 - 에이전트 토큰 비용 (참고 지표, 별점 아님)",
                     "Claude Sonnet 5.5로 module n개를 변경할 때 코드 에이전트 1세션의 토큰·비용 추정 — 별점 조건이 될 수 없는 이유 포함",
                     "산식·상수: DP1 qa4-preregistration M3(상수는 ASSUMED, 가격만 현행 Sonnet 5.5로 교체) · 가격: Anthropic 공개 가격표 Sonnet 5.5 입력 $2 · 캐시 읽기 $0.20 · 출력 $10 / MTok. 실제 에이전트 실행 측정이 아니다")
    rows = [["module 수", "환산 M/M", "총 토큰 (입력+캐시+출력)", "비용 (Sonnet 5.5)", "재시도·탐색 ×3일 때", "대응 별 경계"]]
    marks = {2: "★3 경계", 3: "★2 경계", 5: "★1 경계"}
    for n in (1, 2, 3, 4, 5):
        tk, usd = cd.token_cost(n)
        rows.append([str(n), "%.1f" % cd.mm_of_modules(n), "%.2f M" % (tk / 1e6), "$%.2f" % usd, "$%.2f" % (3 * usd), marks.get(n, "—")])
    d1 = [(k, cd.Q5_VALUES[k]) for k in (("DP1", 1), ("DP1", 2))]
    for (dp, k), (m, mm) in d1:
        tk, usd = cd.token_cost(m)
        rows.append(["%g (DP1 %d안 측정)" % (m, k), "%.2f (측정)" % mm, "%.2f M" % (tk / 1e6), "$%.2f" % usd, "$%.2f" % (3 * usd), "—"])
    add_table(s, 0.4, 1.2, [2.1, 1.5, 2.9, 2.0, 2.2, 1.83], rows, [0.3] + [0.28] * 7, fills=[None] * 7, size=9, aligns=["c"] * 6)
    tbox(s, 0.4, 3.65, 12.5, 1.2,
         [[("추정 방법", {"bold": True, "color": NAVY, "size": 10})],
          [("• 비용 = 입력 컨텍스트 × $2 + 턴 수 × 컨텍스트 × $0.20(캐시 읽기) + 출력 × $10 (MTok 단위). 컨텍스트 = 시스템·도구 25k + 읽는 코드(1.5 × 12 token/LOC × module당 356 LOC × n).", {"size": 8.8})],
          [("• 턴 수 = 10 + 5n, 출력 = 작성(12 token/LOC × 24 LOC × n × 작성·수정 2회) + 턴당 300 token. 356 LOC = vLLM 최상위 KV connector 중앙값. 상수는 모두 가정(ASSUMED)이며 에이전트 실행으로 확인한 값이 아니다.", {"size": 8.8})]],
         size=8.8, line_spacing=1.1)
    rect(s, 0.4, 4.95, 12.53, 1.95, fill="FFF5F5", line="C00000", line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tbox(s, 0.55, 5.0, 12.25, 1.9,
         [[("토큰 비용을 별점 조건으로 쓰지 않는 이유", {"bold": True, "color": "C00000", "size": 10})],
          [("① 비용 ≠ 공수. 토큰 비용은 돈이고, 변경 용이성이 묻는 것은 사람이 들이는 시간이다. 에이전트가 코드를 써도 사람이 설계 적합성 리뷰·통합 테스트·승인을 해야 하므로 공수는 어떤 경우에도 0이 되지 않는다. 그 공수는 module 수에 따라 늘어난다(module 1개 = 0.2 MM).", {"size": 8.8})],
          [("② 후보를 가르지 못한다. 5개 module을 바꿔도 $%.2f, 2개면 $%.2f로 사람 공수(1.0 MM vs 0.4 MM)에 비해 사실상 무시할 수준이고, DP1 두 안의 차이는 $%.2f에 불과하다." % (cd.token_cost(5)[1], cd.token_cost(2)[1], abs(cd.token_cost(2.5)[1] - cd.token_cost(1.75)[1])), {"size": 8.8})],
          [("③ 가격에 종속된다. 모델·가격표가 바뀌면 같은 설계의 별이 바뀐다(후보 간 순위는 같은 모델 안에서는 그대로). 설계 품질이 아닌 시장 가격이 별을 정하게 된다.", {"size": 8.8})],
          [("④ 그래서 참고 지표로 둔다: 'AI 에이전트로 변경해도 비용 부담이 크지 않다'는 보조 근거로만 쓰고, 판정은 module 수(사람 리뷰 범위의 대리 지표)로 한다.", {"size": 8.8, "bold": True})]],
         size=8.8, line_spacing=1.1)
    return s


def capacity_slide(prs):
    s = clone_chrome(prs, 5, "Capacity (동시 세션 수용량) - 평가 기준",
                     "같은 HBM 용량에서 KV를 보존한 채 재개 가능한 최대 동시 세션 수 — HBM 사용량을 줄이는 것이 아니라 같은 HBM에서 더 많이 수용하는 것이 목표",
                     "★1 = Stakeholder(요구 문구: SLO 만족 최대 동시 세션 ≥ Baseline 대비 2x). ★1.5~★3은 문헌 근거 수집 후 확정(미정). 값은 theory_model.py 이론 계산(HBM·DRAM만, 신규 메모리 미반영) [잠정]")
    tbox(s, 0.4, 1.17, 12.5, 0.8,
         [[("[QA] Capacity  [Refinement] 동일 메모리 하 최대 동시 세션 수  ", {"bold": True, "color": NAVY, "size": 10})],
          [("[Scenario] 동일 HBM 용량에서 SLO를 만족하며 동시에 보존·처리 가능한 세션 수가 Baseline 대비 증가해야 한다.  [측정] SLO 만족 최대 동시 세션 수 ≥ Baseline 대비 2x", {"size": 9.5})],
          [("세션 = 컨텍스트 KV를 보존 중인 요청. HBM뿐 아니라 offload(DRAM 등)된 세션도 포함한다. 재계산이 필요한 세션은 세지 않는다.", {"size": 9, "color": "404040"})]],
         size=9.5, line_spacing=1.1)
    cb0, cb1 = tm.capacity_sessions(tm.B0), tm.capacity_sessions(tm.B1)
    rows = [["별", "기준 R_cap = 보존 세션(후보) ÷ 보존 세션(B0 GPU-only)", "근거 / 상태"],
            [[[("★ 3", {"bold": True})]], "미정", "경쟁 기술 위 목표. 제안: ★2 × 1.5 ≈ ×%.1f (문헌 근거 수집 후 확정)" % (1.5 * cb1 / cb0)],
            [[[("★ 2", {"bold": True})]], "미정", "경쟁 기술(KV offload) 수준. 이론 대리 B1 = ×%.1f (HBM+DRAM 단순 offloading). 문헌: 동시 batch·세션 증가 보고(FlexGen, InfiniGen, vLLM PagedAttention 등, 압축 제외)를 수집해 확정" % (cb1 / cb0)],
            [[[("★ 1", {"bold": True})]], "R_cap ≥ 2.0", "Stakeholder 요구: 동일 HBM에서 동시 세션 2배 이상"],
            [[[("미충족", {"bold": True})]], "R_cap < 2.0", "—"]]
    add_table(s, 0.4, 2.05, [0.9, 3.6, 8.03], rows, [0.28, 0.4, 0.55, 0.3, 0.28], fills=[None, "E2F0D9", "FFF2CC", "FCE4D6", "EDEDED"], size=8.8, aligns=["c", "c", "l"])
    tbox(s, 0.4, 3.95, 6.2, 3.0,
         [[("측정 조건 (확정)", {"bold": True, "color": NAVY, "size": 10})],
          [("• HBM을 가득 채웠을 때 측정한다. 설계의 watermark(HBM을 일부 비워 두는 정책)와 무관하다. 비운 자리를 다른 세션으로 채우면 HBM 사용량은 다시 100%가 되므로, 기존 HBM 사용량 평균(Q2')과는 다른 지표다", {"size": 8.6})],
          [("• 반환 세션의 KV 보존율 ≥ 95%(재계산 필요 ≤ 5%)이면서 재개 TTFT P99 ≤ SLO(2 s)일 때만 세션으로 센다. B0는 HBM에서 밀린 세션의 KV가 사라지므로 85세션이 한계다", {"size": 8.6})],
          [("• SSD는 보존 용량에 넣지 않는다(넘치는 KV의 backing store). 포함하면 세션 수가 SSD 링크 처리율로 정해져 Q1과 같아진다", {"size": 8.6})],
          [("• 신규 메모리(ScHBM·CXL-PNM·HBF)를 보존 tier로 셀지는 미정(시스템 구성 페이지 참고). 아래 값은 HBM·DRAM만 센 잠정값이다", {"size": 8.6, "color": "C00000"})]],
         size=8.6, line_spacing=1.1)
    cap = lambda dp, k: cd.THEORY[(dp, k)]["QC"]
    hdr = ["구성 (SYS-H100 1노드, 16K 세션 5 GiB)", "보존 세션", "B0 대비"]
    r2 = [hdr, ["B0 GPU-only (HBM %.0f GiB)" % 425, "%d" % round(cb0), "×1.0"],
          ["B1 단순 offloading (HBM+DRAM)", "%d" % round(cb1), "×%.1f" % (cb1 / cb0)]]
    for dp, k, nm in (("DP1", 1, "DP1 1안 (θ 1.1)"), ("DP1", 2, "DP1 2안 (θ 1.3)"), ("DP3", 1, "DP3 1안 (dedup 0.70)"), ("DP3", 2, "DP3 2안 (dedup 0.70)")):
        r2.append([nm, "%d" % round(cap(dp, k)["c"]), "×%.1f" % cap(dp, k)["r0"]])
    r2.append(["DP2 / DP4", "산정 보류", "보존 tier 정의 필요"])
    add_table(s, 6.8, 3.95, [3.4, 1.3, 1.43], r2, [0.28] + [0.3] * 7, fills=[None] * 7, size=8.6, aligns=["l", "c", "c"])
    tbox(s, 6.8, 6.4, 6.1, 0.6,
         [[("θ(배치 효율)·dedup은 가정 파라미터라 DP1 1안 대 2안 순서도 가정에 의존한다. 기존 Q2'(HBM 감소율) 평가 기준 페이지는 참고로 유지한다.", {"size": 8.2, "color": "595959"})]],
         size=8.2, line_spacing=1.05)
    return s


def system_slide(prs):
    s = clone_chrome(prs, 5, "평가 시나리오 - 시스템 구성 (노드 1대, SYS-H100, 메모리 6종 + SSD)",
                     "기존 HBM·DRAM·SSD에 신규 target 메모리(ScHBM, CXL-PNM, HBF)까지 포함한 이기종 메모리 구성",
                     "출처: doc-mk/Evaluation/DP1/sim/configs/memories_default.json · system-specs.md SYS-H100. 신규 메모리 4종(ScHBM, CXL-PNM, HBF, SSD-PIM)은 과거 세대가 없어 같은 장치를 쓰고 host link BW만 스케일(ASSUMED)")
    rows = [["메모리", "용량", "외부 BW", "내부 BW", "지연", "GPU 직접 접근", "연산(FP16)", "특성 / 이 평가에서의 쓰임"],
            ["HBM (HBM3)", "640 GiB (8×80)", "26.8 TB/s", "26.8 TB/s", "300 ns", "예", "—", "GPU 온패키지. 활성 KV 위치. KV 풀 ≈ 425 GiB (가중치 131·활성값 20 GiB 제외)"],
            ["Samsung Custom HBM (ScHBM)", "160 GiB", "63 GB/s", "6.7 TB/s", "2 µs", "아니오 (host 2홉)", "197.8 TFLOPS", "내부 BW 대비 외부 BW 비대칭 100배 이상. Attention 연산 offload 가능(QK·Softmax·AV)"],
            ["CXL-PNM (LPDDR5X)", "512 GiB", "63 GB/s", "400 GB/s", "300 ns", "아니오", "3.28 TFLOPS", "대용량 연산형 확장 메모리. Attention offload 가능(연산 낮음)"],
            ["DRAM (DDR5-4800)", "1 TiB", "64 GB/s", "300 GB/s", "200 ns", "아니오", "—", "CPU 주 메모리. 기존 offloading의 1차 하위 tier"],
            ["HBF (고대역 플래시)", "2 TiB", "1 TB/s", "1 TB/s", "5 µs", "예 (UCIe 직결)", "—", "HBM과 SSD 사이 용량 계층. 읽기는 빠르나 쓰기 50 GB/s·수명 100 PB → 왕복이 잦은 KV는 부적합"],
            ["SSD (NVMe PCIe5 x4, SSD-PIM)", "16 TiB", "16 GB/s", "200 GB/s", "60 µs", "아니오", "2 TFLOPS (GEMV 한정)", "장기 보관. 쓰기 증폭 4배·수명 10 PB. 연산은 Vector DB 유사도에만 사용(KV Attention 불가)"]]
    add_table(s, 0.4, 1.2, [2.0, 1.2, 0.95, 0.95, 0.7, 1.2, 1.3, 4.23], rows, [0.3] + [0.5] * 6, fills=[None] * 6, size=8, aligns=["l", "c", "c", "c", "c", "c", "c", "l"])
    tbox(s, 0.4, 4.7, 12.5, 2.4,
         [[("이 덱의 이론 값이 반영한 메모리와 반영하지 않은 메모리", {"bold": True, "color": "C00000", "size": 10})],
          [("• 반영: HBM, DRAM, SSD(링크 16 GB/s×0.8). B0 = HBM만, 경쟁 기술 B1 = HBM→DRAM→SSD 단순 offloading", {"size": 9})],
          [("• 미반영: ScHBM, CXL-PNM, HBF. 이 메모리를 후보 설계가 어떻게 쓰는지(보존 tier, Attention offload, GPU 직접 읽기)와 B1이 이 메모리를 어떻게 쓰는지 정해야 모델에 넣을 수 있다. 따라서 현재 후보 값은 기존 3 tier 기준의 잠정값이다", {"size": 9})],
          [("• SSD는 이 시스템 구성에 포함한다(노드당 16 TiB, ASSUMED). 다만 Capacity는 SSD를 보존 용량에서 제외한다(Q1과 중복 방지)", {"size": 9})]],
         size=9, line_spacing=1.12)
    return s


def mapping_slide(prs):
    s = clone_chrome(prs, 5, "Trade-off 별 ↔ 평가 기준 구간 ↔ 이론 예상값",
                     "DP 후보별 소유자 별이 새 기준(GPU-only 대비)의 어떤 수치 구간에 해당하는지, 이론 예상값은 어디에 놓이는지",
                     "예상값 = theory_model.py 분석 모델(시뮬레이션 아님). 붉은 글씨 = 예상값이 소유자 별 구간보다 낮음, 파란 글씨 = 같거나 높음. 기존 중립 벤치마크 시뮬 값은 qa-criteria-v4.md 참조")
    hdr = ["DP / 안", "Q1 처리량", "Q2 Capacity", "Q3 고부하 지연", "Q4 확장성", "Q5 변경 용이성"]
    names = {("DP1", 1): "DP1 1안 상태 관측+이동 결정 통합", ("DP1", 2): "DP1 2안 관측·결정 분리",
             ("DP2", 1): "DP2 1안 스케줄링 시점 결정", ("DP2", 2): "DP2 2안 사전 계획 결정",
             ("DP3", 1): "DP3 1안 Pipe-and-Filter", ("DP3", 2): "DP3 2안 Shared Repository",
             ("DP4", 1): "DP4 1안 중앙 서버 중재", ("DP4", 2): "DP4 2안 락 기반 직접 접근"}
    rows = [hdr]
    for key in sorted(cd.OWNER):
        row = [[[(names[key], {"bold": True, "size": 8})]]]
        for q in ("Q1", "QC", "Q3", "Q4", "Q5"):
            o = cd.OWNER[key].get(q)
            if q == "QC":
                if key in cd.THEORY and "QC" in cd.THEORY[key]:
                    st = cd.theory_star(*key, "QC")
                    row.append([[("소유자 별 미정", {"size": 8, "bold": True, "color": "7F7F7F"})],
                                [(cd.cell_text(*key, "QC") + " → " + ("★1 이상(경계 미정)" if st >= 1 else "☆"), {"size": 7.2, "color": "2F5597"})]])
                else:
                    row.append([[("소유자 별 미정", {"size": 8, "bold": True, "color": "7F7F7F"})], [("산정 보류", {"size": 7.2, "color": "7F7F7F"})]])
                continue
            if o is None:
                row.append([[("—", {"color": "A6A6A6"})]])
                continue
            line1 = [(cd.fmt_star(o) + " → " + cd.BAND_SHORT[q][o], {"size": 8, "bold": True})]
            st = cd.theory_star(*key, q)
            line2 = [(cd.cell_text(*key, q) + " → " + cd.fmt_star(st), {"size": 7.2, "color": "2F5597" if st >= o else "C00000"})]
            row.append([line1, line2])
        rows.append(row)
    add_table(s, 0.4, 1.2, [2.2, 2.15, 2.15, 2.15, 1.95, 1.93], rows, [0.28] + [0.5] * 8, fills=[None] * 8, size=8, aligns=["l"] * 6)
    tot = hit = 0
    for key in cd.OWNER:
        for q, o in cd.OWNER[key].items():
            st = cd.theory_star(*key, q)
            if st is None:
                continue
            tot += 1
            if abs(st - o) <= 0.5:
                hit += 1
    cs = COMP
    tbox(s, 0.4, 5.6, 12.5, 1.5,
         [[("읽는 법", {"bold": True, "color": NAVY, "size": 10})],
          [("• 굵은 줄: 소유자 별과 그 별이 요구하는 새 기준 구간. 작은 줄: 이론 예상값(B0 → 후보)과 새 기준으로 환산한 별. 경쟁 기술(B1)은 Q1 %s(×%.2f), Q3 %s(r %.2f), Q2 ☆, Q4 ★3." % (cd.fmt_star(cs["Q1"][1]), cs["Q1"][0], cd.fmt_star(cs["Q3"][1]), cs["Q3"][0]), {"size": 9})],
          [("• 소유자 별과 ±0.5 이내로 맞는 칸: %d / %d. 후보 간 순위가 소유자 평가와 같은 곳: DP1 Q2(뚜렷), DP2 Q1·Q2·Q4, DP3 Q2. DP1 Q1·Q3는 같은 방향이나 차이가 1%% 미만." % (hit, tot), {"size": 9})],
          [("• 어긋나는 곳(예상값에서 두 안이 사실상 같음): DP2 Q3, DP3 Q1·Q3, DP4 Q1·Q3·Q4. 소유자는 안 간 차이를 예상했다. 파라미터를 바꾸면 달라지는 값이며 가정은 방법 페이지에 적었다.", {"size": 9})]],
         size=9, line_spacing=1.12)
    return s


# ============================================================ main
def main():
    prs = Presentation(BASE)
    S = list(prs.slides)
    edit_requirements(S[0])
    for dp, idx in DP_SLIDE.items():
        annotate_dp(S[idx], dp)
    build_criteria_slides(prs)
    sc1, sc2 = scenario_slides(prs)
    base_s = baseline_slide(prs)
    meth_s = method_slide(prs)
    sim_s = sim_env_slide(prs)
    res = {dp: dp_result_slide(prs, dp) for dp in ("DP1", "DP2", "DP3", "DP4")}
    map_s = mapping_slide(prs)
    tok_s = token_slide(prs)
    cap_s = capacity_slide(prs)
    sys_s = system_slide(prs)
    order = [S[0], S[1], res["DP1"], S[2], res["DP2"], S[3], res["DP3"], S[4], res["DP4"], base_s, meth_s, sim_s,
             S[5], S[6], cap_s, S[7], S[8], S[9], S[10], tok_s, sc1, sys_s, sc2, map_s]
    lst = prs.slides._sldIdLst
    ids = [sl.slide_id for sl in order]
    els = {int(e.get("id")): e for e in lst}
    for e in list(lst):
        lst.remove(e)
    for i in ids:
        lst.append(els[i])
    prs.save(OUT)
    print("saved", OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
